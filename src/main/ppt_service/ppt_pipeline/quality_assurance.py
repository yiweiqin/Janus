from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

try:
    from PIL import Image
except Exception:
    Image = None

from janus_lab.ppt_renderer import SlideSpec
from ppt_pipeline.content_parsing import (
    _dedupe_semantic_items,
    _meaningful_payload_rows,
    _semantic_text,
    _slide_display_points,
)
from ppt_pipeline.image_routing import _slide_image_list
from ppt_pipeline.office_conversion import _convert_powerpoint_to_pdf_windows
from ppt_pipeline.presentation_assets import _shorten_for_cell

def _compact_specs_for_quality_repair(
    specs: list[SlideSpec],
    *,
    level: int = 1,
    slide_numbers: set[int] | None = None,
) -> list[SlideSpec]:
    repaired: list[SlideSpec] = []
    point_limit = 3 if level <= 1 else 2
    point_len = 26 if level <= 1 else 22
    for index, spec in enumerate(specs, start=1):
        if index == 1 or (slide_numbers is not None and index not in slide_numbers):
            repaired.append(spec)
            continue
        points = _slide_display_points(spec, limit=point_limit, include_title=False)
        compact_points = [_shorten_for_cell(point, point_len) for point in points[:point_limit]]
        message = " • " + " • ".join(compact_points) if compact_points else spec.message
        visual = spec.visual or ""
        if not re.search(r"(图|figure|image|diagram|pipeline|架构|流程|表格|案例|截图)", visual, re.IGNORECASE):
            visual = (visual + "；" if visual else "") + "优先使用图/表/流程图承载证据，减少文本卡片"
        repaired.append(
            SlideSpec(
                # Preserve the semantic claim title. Overflow is repaired by
                # autofit/wrapping in the template binder, never by truncation.
                title=spec.title or f"Slide {index}",
                message=message,
                visual=visual,
                speaker_note=spec.speaker_note,
                time=spec.time,
                content_spec=spec.content_spec,
            )
        )
    return repaired


def _adapt_sparse_specs_for_quality_repair(
    specs: list[SlideSpec],
    *,
    slide_numbers: set[int],
    image_indices: set[int] | None = None,
) -> list[SlideSpec]:
    """Match sparse slide content to a layout whose visible slot count it can fill."""
    image_indices = image_indices or set()
    repaired: list[SlideSpec] = []
    for index, spec in enumerate(specs, start=1):
        if index == 1 or index not in slide_numbers:
            repaired.append(spec)
            continue
        points = _slide_display_points(spec, limit=6, include_title=False)
        if not points and str(spec.message or "").strip():
            points = [str(spec.message).strip()]
        payload = dict(spec.content_spec or {})
        payload.setdefault("title", spec.title or f"Slide {index}")
        if index in image_indices or len(points) <= 2:
            layout_id = "basic_content"
            payload["points"] = points
        elif len(points) == 3:
            layout_id = "summary_takeaways"
            payload["cards"] = [{"title": point} for point in points]
            payload["points"] = points
        else:
            layout_id = "evidence_grid"
            payload["cards"] = [{"title": point} for point in points[:6]]
            payload["points"] = points[:6]
        marker = f"layout_id: {layout_id}"
        visual = re.sub(
            r"\blayout_id\s*[:：=]\s*[a-zA-Z0-9_-]+",
            marker,
            spec.visual or "",
            count=1,
            flags=re.IGNORECASE,
        )
        if not re.search(r"\blayout_id\s*[:：=]", visual, re.IGNORECASE):
            visual = f"{marker}\n{visual}".strip()
        repaired.append(SlideSpec(
            title=spec.title or f"Slide {index}",
            message=spec.message,
            visual=visual,
            speaker_note=spec.speaker_note,
            time=spec.time,
            content_spec=payload,
        ))
    return repaired


def _convert_deck_to_quality_pdf(root: Path, deck_path: Path, qa_dir: Path) -> tuple[Path | None, str | None]:
    qa_dir.mkdir(parents=True, exist_ok=True)
    target = qa_dir / "deck.pdf"
    office_error = ""
    if os.name == "nt":
        try:
            pdf_path, office_error = _convert_powerpoint_to_pdf_windows(root, deck_path, target)
        except Exception as exc:
            pdf_path, office_error = None, f"PowerPoint/WPS COM export failed: {exc}"
        if pdf_path is not None:
            return pdf_path, None
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        if os.name != "nt":
            try:
                pdf_path, office_error = _convert_powerpoint_to_pdf_windows(root, deck_path, target)
            except Exception as exc:
                pdf_path, office_error = None, str(exc)
            if pdf_path is not None:
                return pdf_path, None
        return None, f"未安装 LibreOffice/soffice，且 PowerPoint/WPS 导出失败：{office_error or '不可用'}"
    with tempfile.TemporaryDirectory(prefix="opl-ppt-qa-render-") as tmp:
        tmp_path = Path(tmp)
        out_dir = tmp_path / "out"
        profile_dir = tmp_path / "profile"
        out_dir.mkdir()
        profile_dir.mkdir()
        env = os.environ.copy()
        env["HOME"] = str(tmp_path)
        cmd = [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nolockcheck",
            "--nodefault",
            f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(deck_path),
        ]
        try:
            result = subprocess.run(
                cmd,
                cwd=str(root),
                env=env,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=90,
                check=False,
            )
        except subprocess.TimeoutExpired:
            return None, "PPT 导出 PDF 超时。"
        rendered = out_dir / f"{deck_path.stem}.pdf"
        if result.returncode != 0 or not rendered.is_file():
            message = (result.stderr or result.stdout or "PPT 导出 PDF 失败").strip()
            return None, message[-300:]
        shutil.copyfile(rendered, target)
        return target, None


def _render_quality_page_images(pdf_path: Path, qa_dir: Path) -> tuple[list[Path], str | None]:
    try:
        import fitz  # type: ignore

        pages_dir = qa_dir / "pages"
        pages_dir.mkdir(parents=True, exist_ok=True)
        images: list[Path] = []
        doc = fitz.open(str(pdf_path))
        try:
            for index, page in enumerate(doc, start=1):
                pix = page.get_pixmap(matrix=fitz.Matrix(1.0, 1.0), alpha=False)
                path = pages_dir / f"page-{index:02d}.png"
                pix.save(str(path))
                images.append(path)
        finally:
            doc.close()
        return images, None
    except Exception as exc:
        return [], f"PDF 页面截图失败: {exc}"


def _semantic_min_font_size(shape_name: str, text: str = "") -> float:
    name = str(shape_name or "").lower()
    if name.startswith("tpl-title-"):
        compact = re.sub(r"\s+", " ", str(text or "")).strip()
        return 24.0 if len(compact) <= 38 else 22.0 if len(compact) <= 48 else 20.0
    if any(token in name for token in ("grid-", "table-", "mini-table-")):
        return 12.0
    if any(token in name for token in ("caption", "-note", "protocol", "guardrail", "legend", "rule", "dependency")):
        return 12.0
    semantic_prefixes = (
        "basic-", "motivation-", "challenge-", "pipeline-", "loop-", "bench-", "big-", "bar", "bars-",
        "leader-", "ablation-", "evidence-", "case-", "summary-", "target-", "domain-", "route-", "wp-",
        "eval-", "risk-", "roadmap-", "media-",
    )
    if name.startswith(semantic_prefixes):
        return 14.0
    return 0.0


def _ppt_text_quality_issues(deck_path: Path, specs: list[SlideSpec]) -> list[str]:
    issues: list[str] = []
    try:
        prs = Presentation(str(deck_path))
    except Exception as exc:
        return [f"PPTX 无法重新读取: {exc}"]
    if len(prs.slides) != len(specs):
        issues.append(f"页数不一致：计划 {len(specs)} 页，实际 {len(prs.slides)} 页。")
    body_density_units: list[float] = []
    actual_layouts: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        actual_layouts.append(_actual_slide_library_layout_id(slide) or ("cover" if index == 1 else ""))
        total_chars = 0
        density_units = 0.0
        small_semantic_fonts: list[tuple[str, float, float]] = []
        out_of_bounds: list[str] = []
        empty_semantic_slots: list[str] = []
        emphasis_fill_count = 0
        repeated_text: dict[str, list[str]] = {}
        text_boxes = 0
        table_like_shapes = sum(
            1
            for shape in slide.shapes
            if any(token in str(getattr(shape, "name", "") or "").lower() for token in ("grid-", "table-"))
        )
        has_inserted_image = any(
            str(getattr(shape, "name", "") or "").startswith("janus-")
            and str(getattr(shape, "name", "") or "").endswith("-image")
            for shape in slide.shapes
        )
        has_full_width_basic_text = any(
            str(getattr(shape, "name", "") or "") == "basic-text-card"
            and float(getattr(shape, "width", 0) or 0) / 914400 >= 10.0
            for shape in slide.shapes
        )
        challenge_card_count = sum(
            1
            for shape in slide.shapes
            if str(getattr(shape, "name", "") or "").startswith("challenge-card")
        )
        semantic_prefixes = (
            "basic-", "motivation-", "challenge-", "pipeline-", "loop-", "bench-", "big-", "bar", "bars-",
            "leader-", "ablation-", "evidence-", "case-", "summary-", "target-", "domain-", "route-", "wp-",
            "eval-", "risk-", "roadmap-", "media-", "janus-",
        )
        empty_slot_pattern = re.compile(
            r"^(?:motivation-(?:current|target|gap)|challenge-card|pipeline-step-|loop-node-|summary-card|"
            r"target-node|domain-node|route-(?:stage|task)-|roadmap-card-)"
        )
        for shape in slide.shapes:
            name = str(getattr(shape, "name", "") or "")
            if index > 1 and name.startswith(semantic_prefixes):
                left = float(getattr(shape, "left", 0) or 0) / 914400
                top = float(getattr(shape, "top", 0) or 0) / 914400
                right = float((getattr(shape, "left", 0) or 0) + (getattr(shape, "width", 0) or 0)) / 914400
                bottom = float((getattr(shape, "top", 0) or 0) + (getattr(shape, "height", 0) or 0)) / 914400
                if left < 0.34 or right > 12.99 or top < 1.04 or bottom > 6.73:
                    out_of_bounds.append(name)
            if empty_slot_pattern.match(name) and getattr(shape, "has_text_frame", False) and not str(getattr(shape, "text", "") or "").strip():
                empty_semantic_slots.append(name)
            try:
                if str(shape.fill.fore_color.rgb) in {"008A9A", "E08737", "C41230"}:
                    emphasis_fill_count += 1
            except Exception:
                pass
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = re.sub(r"\s+", "", shape.text or "")
            if not text:
                continue
            text_boxes += 1
            total_chars += len(text)
            raw_text = str(shape.text or "")
            normalized_text = re.sub(r"\s+", " ", raw_text).strip().lower()
            if len(normalized_text) >= 14 and not str(getattr(shape, "name", "") or "").startswith("tpl-title-"):
                repeated_text.setdefault(normalized_text, []).append(str(getattr(shape, "name", "") or ""))
            density_units += len(re.findall(r"[\u3400-\u9fff]", raw_text))
            density_units += 1.5 * len(re.findall(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*", raw_text))
            shape_sizes: list[float] = []
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        shape_sizes.append(float(run.font.size.pt))
            threshold = _semantic_min_font_size(str(getattr(shape, "name", "") or ""), raw_text)
            if has_inserted_image and str(getattr(shape, "name", "") or "") == "basic-text-card":
                threshold = min(threshold, 12.0)
            if threshold and shape_sizes and min(shape_sizes) < threshold:
                small_semantic_fonts.append((str(shape.name), min(shape_sizes), threshold))
        if index > 1:
            body_density_units.append(density_units)
        density_limit = 230 if table_like_shapes >= 30 else 185 if table_like_shapes >= 6 else 175 if challenge_card_count >= 4 else 155
        if index > 1 and density_units > density_limit:
            issues.append(f"第 {index} 页文字密度偏高（约 {density_units:.0f} 单位），建议拆分或压缩。")
        if index > 1 and small_semantic_fonts:
            details = "、".join(
                f"{name} {size:.1f}pt<{threshold:.0f}pt"
                for name, size, threshold in small_semantic_fonts[:3]
            )
            issues.append(f"第 {index} 页出现过小字号：{details}。")
        if index > 1 and out_of_bounds:
            issues.append(f"第 {index} 页有元素超出正文安全区：{'、'.join(out_of_bounds[:3])}。")
        if index > 1 and empty_semantic_slots:
            issues.append(f"第 {index} 页保留了空语义卡片：{'、'.join(empty_semantic_slots[:3])}。")
        if index > 1 and emphasis_fill_count > 4:
            issues.append(f"第 {index} 页使用了 {emphasis_fill_count} 个强强调色块，视觉重点过多。")
        duplicate_groups = [names for names in repeated_text.values() if len(names) > 1]
        if index > 1 and duplicate_groups:
            issues.append(f"第 {index} 页存在重复长文本：{'、'.join(duplicate_groups[0][:3])}。")
        if (
            index > 1
            and text_boxes <= 2
            and total_chars > 90
            and not has_inserted_image
            and not has_full_width_basic_text
        ):
            issues.append(f"第 {index} 页文本框过少，版式可能偏单调。")
        visible_text = " ".join(
            str(getattr(shape, "text", "") or "")
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        )
        if re.search(
            r"\[[^\]]{1,80}\]|\b[+\-]?X{1,3}(?:\.X+)?%?\b|template prompt|placeholder|"
            r"(?:^|\s)(?:证据|案例|核心案例|核心视觉)(?:\s*\d+)?(?:\s|$)",
            visible_text,
            re.IGNORECASE,
        ):
            issues.append(f"第 {index} 页仍包含模板占位符或示例文本。")
        if re.search(r"(?:^|\s)(?:label|detail|points|title)\s*[：:]", visible_text, re.IGNORECASE):
            issues.append(f"第 {index} 页暴露了结构化字段名（label/detail/points/title）。")
        if actual_layouts[-1] and any(
            getattr(shape, "has_text_frame", False)
            and str(getattr(shape, "text", "") or "").strip()
            and shape.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for shape in slide.shapes
        ):
            issues.append(f"第 {index} 页仍依赖查看器自动缩字，预览与下载可能不一致。")
        required_conclusion_slot = {
            "result_big_numbers": "big-interpretation",
            "results_bars": "bars-note1",
            "summary_takeaways": "summary-next",
        }.get(actual_layouts[-1])
        if required_conclusion_slot and not any(
            str(getattr(shape, "name", "") or "") == required_conclusion_slot
            and str(getattr(shape, "text", "") or "").strip()
            for shape in slide.shapes
        ):
            issues.append(f"第 {index} 页结果/总结页面没有可见的结论区域。")
        exposed_kickers = [
            str(getattr(shape, "text", "") or "").strip()
            for shape in slide.shapes
            if str(getattr(shape, "name", "") or "").startswith("tpl-kicker-label")
            and str(getattr(shape, "text", "") or "").strip().lower()
            in {"内容", "视觉", "text", "visual", "content"}
        ]
        if exposed_kickers:
            issues.append(f"第 {index} 页暴露了内部栏目标签：{'、'.join(exposed_kickers)}。")
        if index > 1 and index - 1 < len(specs):
            expected_title = str((specs[index - 1].content_spec or {}).get("title") or specs[index - 1].title or "").strip()
            actual_title = next((
                str(getattr(shape, "text", "") or "").strip()
                for shape in slide.shapes
                if str(getattr(shape, "name", "") or "").startswith("tpl-title-")
            ), "")
            if expected_title and actual_title and re.sub(r"\s+", " ", actual_title) != re.sub(r"\s+", " ", expected_title):
                issues.append(f"第 {index} 页标题与页面计划不一致，禁止截断语义标题。")
            issues.extend(_semantic_slot_quality_issues(slide, specs[index - 1], index))
    if body_density_units:
        avg_density = sum(body_density_units) / len(body_density_units)
        if avg_density > 125:
            issues.append(f"整套正文平均文字密度偏高（约 {avg_density:.0f} 单位/页），需要更图表化。")
    body_layouts = [layout for layout in actual_layouts[1:] if layout]
    for layout in sorted(set(body_layouts)):
        count = body_layouts.count(layout)
        if count > 2:
            issues.append(f"整套 PPT 重复使用 {layout} 共 {count} 次，应启用替补版式。")
    for index in range(1, len(actual_layouts)):
        if actual_layouts[index] and actual_layouts[index] == actual_layouts[index - 1]:
            issues.append(f"第 {index}、{index + 1} 页连续使用 {actual_layouts[index]}，页面变化不足。")
    body_families = [_quality_layout_family(layout) for layout in actual_layouts[1:] if layout]
    body_densities = [_quality_layout_density(layout) for layout in actual_layouts[1:] if layout]
    for index in range(1, len(body_densities)):
        if body_densities[index - 1] == body_densities[index] == "dense":
            issues.append(f"第 {index + 1}、{index + 2} 页连续为高密度页面，阅读节奏过紧。")
    for index in range(2, len(body_families)):
        if body_families[index - 2] == body_families[index - 1] == body_families[index]:
            issues.append(f"第 {index}–{index + 2} 页连续使用 {body_families[index]} 类页面，整套节奏缺少变化。")
    return issues


_SLIDE_LIBRARY_TITLE_LAYOUTS = {
    "tpl-title-2": "basic_content",
    "tpl-title-basic-mirror": "basic_content_mirror",
    "tpl-title-3": "motivation_compare",
    "tpl-title-4": "challenge_map",
    "tpl-title-5": "method_pipeline",
    "tpl-title-6": "method_loop",
    "tpl-title-7": "benchmark_metrics",
    "tpl-title-8": "result_big_numbers",
    "tpl-title-9": "results_bars",
    "tpl-title-10": "leaderboard_table",
    "tpl-title-11": "ablation_matrix",
    "tpl-title-12": "evidence_grid",
    "tpl-title-13": "case_gallery",
    "tpl-title-14": "summary_takeaways",
    "tpl-title-15": "project_target_map",
    "tpl-title-16": "domain_object_map",
    "tpl-title-17": "technical_route",
    "tpl-title-18": "workpackage_matrix",
    "tpl-title-19": "evaluation_dashboard",
    "tpl-title-20": "risk_action_table",
    "tpl-title-21": "milestone_roadmap",
    "tpl-title-22": "media_showcase",
}


def _actual_slide_library_layout_id(slide) -> str:
    for shape in slide.shapes:
        layout_id = _SLIDE_LIBRARY_TITLE_LAYOUTS.get(str(getattr(shape, "name", "") or ""))
        if layout_id:
            return layout_id
    return ""


def _quality_layout_family(layout_id: str) -> str:
    if layout_id in {"basic_content", "basic_content_mirror", "motivation_compare", "challenge_map", "summary_takeaways"}:
        return "叙事"
    if layout_id in {"evidence_grid", "case_gallery", "media_showcase"}:
        return "视觉"
    if layout_id in {"method_pipeline", "method_loop", "technical_route", "milestone_roadmap"}:
        return "流程"
    if layout_id in {
        "benchmark_metrics", "result_big_numbers", "results_bars", "leaderboard_table", "ablation_matrix",
        "workpackage_matrix", "evaluation_dashboard", "risk_action_table",
    }:
        return "数据"
    return "结构"


def _quality_layout_density(layout_id: str) -> str:
    if layout_id in {"leaderboard_table", "ablation_matrix", "workpackage_matrix", "risk_action_table", "evaluation_dashboard"}:
        return "dense"
    if layout_id in {"evidence_grid", "case_gallery", "media_showcase"}:
        return "visual"
    if layout_id in {"result_big_numbers", "summary_takeaways", "motivation_compare", "milestone_roadmap"}:
        return "sparse"
    return "medium"


def _semantic_slot_quality_issues(slide, spec: SlideSpec, index: int) -> list[str]:
    payload = spec.content_spec or {}
    layout_match = re.search(r"\blayout_id\s*[:：=]\s*([a-zA-Z0-9_-]+)", spec.visual or "", re.IGNORECASE)
    declared_layout_id = layout_match.group(1) if layout_match else ""
    layout_id = _actual_slide_library_layout_id(slide) or declared_layout_id

    def nonempty_named(prefix: str) -> int:
        return sum(
            1
            for shape in slide.shapes
            if str(getattr(shape, "name", "") or "").startswith(prefix)
            and str(getattr(shape, "text", "") or "").strip()
        )

    issues: list[str] = []
    if layout_id == "motivation_compare":
        current = _semantic_text(payload.get("current"))
        target = _semantic_text(payload.get("target"))
        gap = _semantic_text(payload.get("gap"))
        if not current or not target or not gap:
            issues.append(f"第 {index} 页当前态、目标态与关键差距未完整提供。")
        if current and target and re.sub(r"\W+", "", current.lower()) == re.sub(r"\W+", "", target.lower()):
            issues.append(f"第 {index} 页当前态与目标态内容相同。")
    elif layout_id in {"method_pipeline", "technical_route"}:
        steps = payload.get("steps") if isinstance(payload.get("steps"), list) else payload.get("stages")
        step_count = len(_dedupe_semantic_items(steps or [], limit=6))
        if step_count < 3:
            issues.append(f"第 {index} 页方法流程不足 3 个有效步骤。")
        if nonempty_named("pipeline-step-") and nonempty_named("pipeline-step-") < min(step_count, 5):
            issues.append(f"第 {index} 页方法流程步骤未完整绑定。")
    elif layout_id == "leaderboard_table":
        if len(_meaningful_payload_rows(payload)) < 3:
            issues.append(f"第 {index} 页排行榜数据行不足 3 行。")
    elif layout_id == "ablation_matrix":
        if len(_meaningful_payload_rows(payload)) < 3:
            issues.append(f"第 {index} 页消融矩阵数据行不足 3 行。")
    elif layout_id == "summary_takeaways":
        summary_items = payload.get("findings") or payload.get("cards") or payload.get("points") or []
        if len(_dedupe_semantic_items(summary_items if isinstance(summary_items, list) else [], limit=3)) < 3:
            issues.append(f"第 {index} 页总结要点不足 3 项或存在重复。")

    table_prefix = {
        "leaderboard_table": "leaderboard-grid-",
        "ablation_matrix": "ablation-grid-",
        "benchmark_metrics": "bench-table-",
    }.get(layout_id)
    if table_prefix:
        table_cells = [
            shape for shape in slide.shapes
            if str(getattr(shape, "name", "") or "").startswith(table_prefix)
            and getattr(shape, "has_text_frame", False)
        ]
        populated = sum(1 for shape in table_cells if str(getattr(shape, "text", "") or "").strip())
        if table_cells and populated / len(table_cells) < 0.55:
            issues.append(f"第 {index} 页数据表有效填充率过低，存在空表或字段错位。")

    if layout_id == "project_target_map":
        row = payload.get("row")
        has_row = isinstance(row, (list, dict)) and bool(row)
        grid_count = nonempty_named("target-grid-")
        if has_row and grid_count < 5:
            issues.append(f"第 {index} 页目标验收行未完整绑定。")
        if not has_row and any(str(getattr(shape, "name", "") or "").startswith("target-grid-") for shape in slide.shapes):
            issues.append(f"第 {index} 页没有验收数据却保留了空表格。")
    elif layout_id == "domain_object_map":
        expected = min(5, len(payload.get("nodes") or []))
        actual_nodes = nonempty_named("domain-node")
        actual_edges = sum(1 for shape in slide.shapes if str(getattr(shape, "name", "") or "") == "domain-edge")
        max_edges = {0: 0, 1: 0, 2: 1, 3: 2, 4: 4, 5: 6}.get(expected, 6)
        if expected and actual_nodes != expected:
            issues.append(f"第 {index} 页对象节点数量不一致：期望 {expected}，实际 {actual_nodes}。")
        if actual_edges > max_edges:
            issues.append(f"第 {index} 页存在连接到空节点的残余连线。")
    elif layout_id == "challenge_map":
        risk_table_mode = any(
            str(getattr(shape, "name", "") or "").startswith("risk-grid-")
            for shape in slide.shapes
        )
        if risk_table_mode:
            return issues
        expected = len(payload.get("nodes") or payload.get("risks") or [])
        if expected and nonempty_named("challenge-card") < min(expected, 4):
            issues.append(f"第 {index} 页挑战/风险卡片未完整绑定。")
    elif layout_id == "method_loop":
        expected = len(payload.get("steps") or payload.get("stages") or [])
        if expected < 3:
            issues.append(f"第 {index} 页循环节点不足 3 个。")
        if expected and nonempty_named("loop-node-") < min(expected, 4):
            issues.append(f"第 {index} 页循环节点未完整绑定。")
    elif layout_id == "evidence_grid":
        expected = len(payload.get("cards") or payload.get("captions") or payload.get("nodes") or [])
        basic_image_mode = any(
            str(getattr(shape, "name", "") or "") == "janus-basic-image-image"
            for shape in slide.shapes
        )
        if basic_image_mode:
            return issues
        case_mode = any(
            str(getattr(shape, "name", "") or "").startswith(("case-", "janus-case-"))
            for shape in slide.shapes
        )
        if case_mode:
            actual = nonempty_named("case-main") + nonempty_named("case-thumb-")
            actual += sum(
                1
                for shape in slide.shapes
                if str(getattr(shape, "name", "") or "").startswith("janus-case-")
                and str(getattr(shape, "name", "") or "").endswith("-image")
            )
            if expected and actual < min(expected, 5):
                issues.append(f"第 {index} 页证据/案例卡片未完整绑定。")
            return issues
        inserted = sum(
            1
            for shape in slide.shapes
            if str(getattr(shape, "name", "") or "").startswith("janus-evidence-")
            and str(getattr(shape, "name", "") or "").endswith("-image")
        )
        if expected and nonempty_named("evidence-img-") + inserted < min(expected, 6):
            issues.append(f"第 {index} 页证据卡片未完整绑定。")
    elif layout_id == "case_gallery":
        expected = len(payload.get("cards") or payload.get("captions") or [])
        basic_image_mode = any(
            str(getattr(shape, "name", "") or "") == "janus-basic-image-image"
            for shape in slide.shapes
        )
        if basic_image_mode:
            return issues
        actual = nonempty_named("case-main") + nonempty_named("case-thumb-")
        if expected and actual < min(expected, 5):
            issues.append(f"第 {index} 页案例卡片未完整绑定。")
    elif layout_id == "benchmark_metrics":
        evaluation_mode = any(
            str(getattr(shape, "name", "") or "").startswith("eval-")
            for shape in slide.shapes
        )
        if evaluation_mode:
            return issues
        expected = len(payload.get("kpis") or [])
        if expected < 3 and len(_meaningful_payload_rows(payload)) < 3:
            issues.append(f"第 {index} 页评测指标或数据行不足 3 项。")
        if expected and nonempty_named("bench-kpi") < min(expected, 3):
            issues.append(f"第 {index} 页 KPI 卡片未完整绑定。")
    elif layout_id == "evaluation_dashboard":
        rows = payload.get("status_rows") if isinstance(payload.get("status_rows"), list) else []
        header2 = next((
            str(getattr(shape, "text", "") or "").strip()
            for shape in slide.shapes
            if str(getattr(shape, "name", "") or "") == "eval-status-grid-r1c2"
        ), "")
        has_status = any(isinstance(item, dict) and _display_status_value(item) for item in rows)
        if rows and has_status and "status" not in header2.lower() and "状态" not in header2:
            issues.append(f"第 {index} 页评测状态字段与表头不匹配。")
        if rows and not has_status and ("status" in header2.lower() or "状态" in header2):
            issues.append(f"第 {index} 页没有状态数据，却把证据信号放入 Status 列。")
        if not rows and any(str(getattr(shape, "name", "") or "").startswith("eval-status-grid-") for shape in slide.shapes):
            issues.append(f"第 {index} 页没有评测行却保留了空状态表。")
    elif layout_id == "risk_action_table":
        raw = payload.get("risks") if isinstance(payload.get("risks"), list) else payload.get("rows")
        records = raw if isinstance(raw, list) else []
        has_owner = any(isinstance(item, dict) and str(item.get("owner") or item.get("responsible") or "").strip() for item in records)
        headers = " ".join(
            str(getattr(shape, "text", "") or "")
            for shape in slide.shapes
            if re.fullmatch(r"risk-grid-r1c\d", str(getattr(shape, "name", "") or ""))
        )
        if records and not has_owner and re.search(r"owner|负责人", headers, re.IGNORECASE):
            issues.append(f"第 {index} 页没有责任人数据却保留 Owner 列，可能发生字段错位。")
    elif layout_id == "milestone_roadmap":
        expected = min(6, len(payload.get("milestones") or []))
        actual_dots = sum(1 for shape in slide.shapes if str(getattr(shape, "name", "") or "").startswith("roadmap-dot-"))
        if expected and actual_dots != expected:
            issues.append(f"第 {index} 页路线图节点数量不一致：期望 {expected}，实际 {actual_dots}。")
    return issues


def _display_status_value(item: dict[str, Any]) -> str:
    return str(item.get("status") or item.get("state") or "").strip()


def _visual_quality_issues(page_images: list[Path], specs: list[SlideSpec], image_paths: dict[int, Path | list[Path]]) -> list[str]:
    issues: list[str] = []
    if Image is None:
        return issues
    for index, image_path in enumerate(page_images, start=1):
        if index == 1:
            continue
        try:
            with Image.open(image_path) as img:
                rgb = img.convert("RGB")
                w, h = rgb.size
                sample = rgb.resize((max(1, w // 8), max(1, h // 8)))
                pixels = list(sample.getdata())
        except Exception:
            continue
        non_white = sum(1 for r, g, b in pixels if min(r, g, b) < 245)
        ratio = non_white / max(len(pixels), 1)
        if ratio < 0.065:
            issues.append(f"第 {index} 页视觉内容偏少，页面可能过空或过度依赖文字。")
    required_indices = [
        index
        for index, spec in enumerate(specs, start=1)
        if re.search(r"使用附件原图|source figure|生成插图|generate image", spec.visual or "", re.IGNORECASE)
    ]
    missing_required = [index for index in required_indices if index not in image_paths]
    if missing_required:
        issues.append(f"第 {', '.join(map(str, missing_required))} 页明确要求图片，但没有可用图片产物。")
    return issues


def _ppt_image_binding_issues(
    deck_path: Path,
    image_paths: dict[int, Path | list[Path]],
    *,
    minimum_images: int = 1,
) -> list[str]:
    expected = sum(len(_slide_image_list(value)) for value in image_paths.values())
    if expected < minimum_images:
        return [
            f"整份 PPT 没有可用的内容图片；至少需要 {minimum_images} 张附件原图或 gpt-image-2 生成图片，必须重新生成或重新排版。"
        ]
    try:
        prs = Presentation(str(deck_path))
    except Exception as exc:
        return [f"无法检查图片绑定：{exc}"]
    inserted = sum(
        1
        for slide in prs.slides
        for shape in slide.shapes
        if str(getattr(shape, "name", "") or "").startswith("janus-")
        and str(getattr(shape, "name", "") or "").endswith("-image")
    )
    if inserted < minimum_images:
        return [
            f"整份 PPT 虽准备了 {expected} 张内容图片，但实际未成功插入至少 {minimum_images} 张，必须重新分配到兼容图片槽位。"
        ]
    if inserted < expected:
        return [f"已生成/提取 {expected} 张图片，但仅成功插入 {inserted} 张，图片未完整绑定，必须重新分配到兼容图片槽位。"]
    return []


def _quality_report_lines(
    *,
    deck_path: Path,
    specs: list[SlideSpec],
    image_paths: dict[int, Path | list[Path]],
    root: Path,
    attempt: int,
    minimum_images: int = 1,
) -> tuple[list[str], bool, list[str]]:
    qa_dir = deck_path.parent / "qa" / f"attempt-{attempt}"
    lines = [f"## PPT Render QA - Attempt {attempt}"]
    pdf_path, pdf_error = _convert_deck_to_quality_pdf(root, deck_path, qa_dir)
    issues: list[str] = []
    if pdf_error:
        if "未安装 LibreOffice" in pdf_error or "soffice" in pdf_error:
            lines.append(f"- PDF QA skipped: {pdf_error}")
        else:
            issues.append(pdf_error)
    page_images: list[Path] = []
    if pdf_path is not None:
        lines.append(f"- PDF: {pdf_path}")
        page_images, image_error = _render_quality_page_images(pdf_path, qa_dir)
        if image_error:
            issues.append(image_error)
        elif page_images:
            lines.append(f"- Rendered pages: {len(page_images)}")
    issues.extend(_ppt_text_quality_issues(deck_path, specs))
    issues.extend(_ppt_image_binding_issues(deck_path, image_paths, minimum_images=max(0, minimum_images)))
    if page_images:
        issues.extend(_visual_quality_issues(page_images, specs, image_paths))
    if issues:
        lines.append("- Status: needs repair")
        lines.append("- Findings:")
        lines.extend(f"  - {issue}" for issue in issues[:10])
    else:
        lines.append("- Status: passed")
        lines.append("- Findings: no blocking layout issues detected by automated QA.")
    return lines, not issues, issues



__all__ = [
    "_compact_specs_for_quality_repair",
    "_adapt_sparse_specs_for_quality_repair",
    "_convert_deck_to_quality_pdf",
    "_render_quality_page_images",
    "_semantic_min_font_size",
    "_ppt_text_quality_issues",
    "_actual_slide_library_layout_id",
    "_quality_layout_family",
    "_quality_layout_density",
    "_semantic_slot_quality_issues",
    "_display_status_value",
    "_visual_quality_issues",
    "_ppt_image_binding_issues",
    "_quality_report_lines",
]
