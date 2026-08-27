from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from PIL import Image
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ppt_pipeline.content_payloads import *
from ppt_pipeline.slide_library_core import *
from ppt_pipeline.slide_library_policy import *
from ppt_pipeline.slide_library_academic import *

def _bind_project_target(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 8)
    objective = _display(payload.get("objective")) or (points[0] if points else "项目目标")
    root = _shape_named(slide, "target-root")
    if root is not None:
        root.text_frame.margin_left = Inches(0.04)
        root.text_frame.margin_right = Inches(0.04)
        root.text_frame.margin_top = Inches(0.03)
        root.text_frame.margin_bottom = Inches(0.03)
        _set_geometry(root, 0.85, 1.48, 11.63, 1.08)
    _set_shape_text(root, objective, max_size=18, bold=True)
    _set_text_spacing(root, after=0, line_spacing=1.0, middle=True)
    node_items = _structured_list(payload, "nodes")
    nodes = [
        _compact_node_label(_item_parts(item)[0] or _display(item))
        for item in node_items
        if _item_parts(item)[0] or _display(item)
    ] or [_compact_node_label(point) for point in points[1:5]]
    nodes = nodes[:4]
    row = _target_row(payload)
    count = max(1, len(nodes))
    left, right, gap = 0.85, 12.48, 0.32
    node_width = min(2.76, (right - left - gap * (count - 1)) / count)
    total_width = node_width * count + gap * (count - 1)
    start = left + (right - left - total_width) / 2
    node_y = 3.10 if any(row) else 3.36
    node_height = 1.66 if any(row) else 2.02
    node_shapes = _shapes_named(slide, "target-node")
    active_nodes: list[Any] = []
    for index, shape in enumerate(node_shapes):
        if index >= len(nodes):
            _remove_shape(shape)
            continue
        _set_geometry(shape, start + index * (node_width + gap), node_y, node_width, node_height)
        _set_shape_text(shape, nodes[index], max_size=14, bold=True)
        _set_text_spacing(shape, after=0, line_spacing=1.0, middle=True)
        active_nodes.append(shape)
    for index, link in enumerate(_shapes_named(slide, "target-link")):
        if index >= len(active_nodes):
            _remove_shape(link)
            continue
        node = active_nodes[index]
        _set_line_endpoints(
            link,
            _rectangle_edge_point(root, node),
            _rectangle_edge_point(node, root),
        )
    if not any(row):
        _remove_prefix(slide, "target-grid-")
        return
    for col in range(1, 5):
        _set_named(slide, f"target-grid-r2c{col}", row[col - 1] if col - 1 < len(row) else "", max_size=11)
    headers = (
        ["交付物", "指标", "验收", "负责人"]
        if _is_cjk_payload(payload)
        else ["Output", "Indicator", "Acceptance", "Owner"]
    )
    strip_left, strip_width = 0.85, 11.63 / 4
    for col in range(1, 5):
        header = _shape_named(slide, f"target-grid-r1c{col}")
        body = _shape_named(slide, f"target-grid-r2c{col}")
        for shape in (header, body):
            if shape is None:
                continue
            shape.text_frame.margin_left = Inches(0.02)
            shape.text_frame.margin_right = Inches(0.02)
            shape.text_frame.margin_top = Inches(0.01)
            shape.text_frame.margin_bottom = Inches(0.01)
        if header is not None:
            _set_geometry(header, strip_left + (col - 1) * strip_width, 5.61, strip_width, 0.42)
            _set_shape_text(header, headers[col - 1], max_size=11, bold=True)
        if body is not None:
            _set_geometry(body, strip_left + (col - 1) * strip_width, 6.03, strip_width, 0.55)
            _refit_shape_text(body, max_size=12, min_size=10)
    _style_minimal_grid(slide, "target-grid", columns=4, data_rows=1)


def _layout_domain_edges(slide: Any, nodes: list[Any]) -> None:
    pairs_by_count = {
        0: (),
        1: (),
        2: ((0, 1),),
        3: ((0, 1), (0, 2)),
        4: ((0, 1), (0, 2), (1, 3), (2, 3)),
        5: ((0, 1), (1, 2), (0, 3), (1, 3), (1, 4), (2, 4)),
    }
    pairs = pairs_by_count.get(len(nodes), ())
    for index, edge in enumerate(_shapes_named(slide, "domain-edge")):
        if index >= len(pairs):
            _remove_shape(edge)
            continue
        source_index, target_index = pairs[index]
        source, target = nodes[source_index], nodes[target_index]
        _set_line_endpoints(
            edge,
            _rectangle_edge_point(source, target),
            _rectangle_edge_point(target, source),
        )


def _bind_domain(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 8)
    layers = _structured_list(payload, "layers")
    layer_labels = [
        _compact_node_label(_item_parts(item)[0] or _display(item))
        for item in layers
        if _item_parts(item)[0] or _display(item)
    ]
    layer_shape = _shape_named(slide, "domain-layers")
    _set_geometry(layer_shape, 0.75, 1.48, 2.35, 4.97)
    layer_labels = layer_labels or points[:4]
    layer_heading = _localized(payload, "领域层级", "DOMAIN LAYERS")
    layer_text = "\n\n".join([
        layer_heading,
        "\n".join(f"{index:02d}  {label}" for index, label in enumerate(layer_labels, start=1)),
    ])
    _set_shape_text(layer_shape, layer_text, max_size=14)
    if layer_shape is not None:
        layer_shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        layer_shape.text_frame.margin_left = Inches(0.26)
        layer_shape.text_frame.margin_right = Inches(0.20)
        layer_shape.text_frame.margin_top = Inches(0.28)
    _set_text_spacing(layer_shape, after=8, line_spacing=1.08)
    node_items = _structured_list(payload, "nodes")
    nodes = [
        _compact_node_label(_item_parts(item)[0] or _display(item))
        for item in node_items
        if _item_parts(item)[0] or _display(item)
    ] or [_compact_node_label(point) for point in points[3:8]]
    nodes = nodes[:5]
    positions_by_count = {
        1: ((6.54, 3.02),),
        2: ((4.22, 3.02), (8.86, 3.02)),
        3: ((6.54, 1.72), (4.22, 4.02), (8.86, 4.02)),
        4: ((4.22, 1.78), (8.86, 1.78), (4.22, 4.00), (8.86, 4.00)),
        5: ((3.48, 1.78), (6.54, 1.78), (9.60, 1.78), (5.01, 4.00), (8.07, 4.00)),
    }
    node_shapes = _shapes_named(slide, "domain-node")
    active_nodes: list[Any] = []
    for index, shape in enumerate(node_shapes):
        if index >= len(nodes):
            _remove_shape(shape)
            continue
        x, y = positions_by_count[len(nodes)][index]
        _set_geometry(shape, x, y, 2.45, 1.20)
        _set_shape_text(shape, nodes[index], max_size=15, bold=True)
        _set_text_spacing(shape, after=0, line_spacing=1.0, middle=True)
        active_nodes.append(shape)
    _layout_domain_edges(slide, active_nodes)
    legend = _shape_named(slide, "domain-legend")
    _set_geometry(legend, 3.48, 5.75, 8.57, 0.70)
    _set_shape_text(legend, _display(payload.get("legend")) or "对象与关系", max_size=12)
    _set_text_spacing(legend, after=2, line_spacing=1.0, middle=True)


def _bind_route(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 8)
    stages = _structured_list(payload, "stages")[:4]
    if not stages:
        stages = points[:4] or [""]
    count = min(4, max(1, len(stages)))
    left, right, gap = 0.75, 12.58, 0.28
    width = (right - left - gap * (count - 1)) / count
    for index in range(4):
        stage = stages[index] if index < len(stages) else None
        stage_shape = _shape_named(slide, f"route-stage-{index}")
        task_shape = _shape_named(slide, f"route-task-{index}")
        if stage is None:
            _remove_shape(stage_shape)
            _remove_shape(task_shape)
            continue
        if isinstance(stage, dict):
            title = _display(stage.get("title") or stage.get("label")) or f"阶段 {index + 1}"
            task = _display(stage.get("task") or stage.get("content") or stage.get("detail") or stage.get("points"))
        else:
            title = f"阶段 {index + 1}"
            task = _display(stage)
        x = left + index * (width + gap)
        _set_geometry(stage_shape, x, 1.50, width, 0.78)
        _set_shape_text(stage_shape, f"{index + 1:02d}  {title}", max_size=15, bold=True)
        _set_text_spacing(stage_shape, after=0, line_spacing=1.0, middle=True)
        _set_geometry(task_shape, x, 2.56, width, 2.95)
        _set_shape_text(task_shape, task, max_size=14)
        if task_shape is not None:
            task_shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            task_shape.text_frame.margin_top = Inches(0.24)
            task_shape.text_frame.margin_left = Inches(0.20)
            task_shape.text_frame.margin_right = Inches(0.18)
        _set_text_spacing(task_shape, after=6, line_spacing=1.08)
    active_headers = [_shape_named(slide, f"route-stage-{index}") for index in range(count)]
    for index in range(3):
        arrow = _shape_named(slide, f"route-arrow-{index}")
        if index >= count - 1:
            if arrow is not None:
                _remove_shape(arrow)
            continue
        source, target = active_headers[index], active_headers[index + 1]
        _set_line_endpoints(
            arrow,
            _rectangle_edge_point(source, target),
            _rectangle_edge_point(target, source),
        )
    dependency = _shape_named(slide, "route-dependency")
    _set_geometry(dependency, 0.75, 5.91, 11.83, 0.67)
    _set_shape_text(dependency, _display(payload.get("dependency")) or (points[-1] if points else ""), max_size=13)
    _set_text_spacing(dependency, after=2, line_spacing=1.0, middle=True)


def _bind_workpackages(slide, payload: dict[str, Any]) -> None:
    _fill_grid(slide, "wp-grid", payload, 6, 5, ["任务包", "负责人", "任务", "交付物", "指标", "时间"])
    rows = _rows(payload)
    active_rows = min(5, len(rows) or len(_payload_points(payload, 5)) or 1)
    widths = (1.15, 1.35, 2.55, 2.45, 1.85, 1.80)
    row_height = min(0.82, 4.22 / (active_rows + 1))
    for row_index in range(1, active_rows + 2):
        x = 0.75
        for col, width in enumerate(widths, start=1):
            shape = _shape_named(slide, f"wp-grid-r{row_index}c{col}")
            if shape is not None:
                _set_geometry(shape, x, 1.48 + (row_index - 1) * row_height, width, row_height)
                _refit_shape_text(shape, max_size=13 if row_index == 1 else 12, min_size=11)
            x += width
    _style_minimal_grid(slide, "wp-grid", columns=6, data_rows=active_rows, highlight_row=2 if active_rows else None)
    _set_geometry(_shape_named(slide, "wp-note"), 0.75, 5.93, 11.15, 0.64)
    _set_named(slide, "wp-note", _display(payload.get("note")) or "任务、交付、指标和时间形成闭环", max_size=13)
    _set_text_spacing(_shape_named(slide, "wp-note"), after=2, line_spacing=1.0, middle=True)


def _bind_evaluation(slide, payload: dict[str, Any]) -> None:
    _fill_kpis(slide, "eval", payload)
    kpis = _kpis(payload)
    kpi_count = max(1, min(3, len(_structured_list(payload, "kpis")) or len(kpis)))
    primary_kpi = _semantic_primary_index(kpis[:kpi_count], payload)
    for index in range(1, 4):
        base = _shape_named(slide, f"eval-kpi{index}")
        if base is None:
            continue
        x = float(base.left) / 914400
        width = float(base.width) / 914400
        _set_geometry(base, x, 1.48, width, 2.12)
        accent = _shape_named(slide, f"eval-kpi{index}-accent")
        value_shape = _shape_named(slide, f"eval-kpi{index}-value")
        _set_geometry(accent, x, 1.48, 0.07, 2.12)
        _set_geometry(value_shape, x + 0.19, 1.68, max(0.8, width - 0.38), 0.62)
        is_primary = index - 1 == primary_kpi
        _set_shape_colors(base, fill=SEMANTIC_PRIMARY_LIGHT if is_primary else RGBColor(255, 255, 255), line=RGBColor(218, 227, 233))
        _set_shape_colors(accent, fill=SEMANTIC_PRIMARY if is_primary else SEMANTIC_MUTED, line=None)
        _set_shape_text_color(value_shape, SEMANTIC_PRIMARY if is_primary else SEMANTIC_DARK)
    for index in range(1, 4):
        label = _shape_named(slide, f"eval-kpi{index}-label")
        if label is not None:
            label.text_frame.margin_left = Inches(0.02)
            label.text_frame.margin_right = Inches(0.02)
            label.text_frame.margin_top = Inches(0.01)
            label.text_frame.margin_bottom = Inches(0.01)
            _set_geometry(label, float(label.left) / 914400, 2.36, float(label.width) / 914400, 0.56)
            _refit_shape_text(label, max_size=14, min_size=12)
        note = _shape_named(slide, f"eval-kpi{index}-note")
        if note is not None:
            note.text_frame.margin_left = Inches(0.02)
            note.text_frame.margin_right = Inches(0.02)
            note.text_frame.margin_top = Inches(0.01)
            note.text_frame.margin_bottom = Inches(0.01)
            _set_geometry(note, float(note.left) / 914400, 3.00, float(note.width) / 914400, 0.44)
            value = kpis[index - 1]["note"] if index - 1 < len(kpis) else note.text
            _set_shape_text(note, _compact_supporting_text(value, limit=40), max_size=11)
        for inner_shape in (
            _shape_named(slide, f"eval-kpi{index}-value"),
            label,
            note,
        ):
            _set_shape_colors(inner_shape, fill=None, line=None)
    scope = _shape_named(slide, "eval-scope")
    _set_geometry(scope, 8.78, 1.48, 3.80, 2.12)
    scope_text = _display(payload.get("scope")) or "\n".join(_payload_points(payload, 4))
    _set_shape_text(scope, f"{_localized(payload, '评测范围', 'EVALUATION SCOPE')}\n{scope_text}", max_size=14)
    _set_shape_colors(scope, fill=SEMANTIC_WARM_LIGHT, line=None)
    _style_first_paragraph(scope, color=SEMANTIC_WARM)
    _set_text_spacing(scope, after=7, line_spacing=1.06, middle=True)
    explicit_bars = _structured_list(payload, "bars")
    if explicit_bars:
        bars = _bars(payload, 3)
        primary_bar = _semantic_primary_index(bars, payload)
        for index, item in enumerate(bars, 1):
            _set_named(slide, f"eval-bar{index}-label", item["label"], max_size=11)
            _set_named(slide, f"eval-bar{index}-value", item["value"], max_size=11)
            _resize_fill(slide, f"eval-bar{index}-fill", f"eval-bar{index}-track", item["value"])
            is_primary = index - 1 == primary_bar
            _set_shape_colors(_shape_named(slide, f"eval-bar{index}-track"), fill=RGBColor(233, 241, 245), line=None)
            _set_shape_colors(_shape_named(slide, f"eval-bar{index}-fill"), fill=SEMANTIC_PRIMARY if is_primary else SEMANTIC_MUTED, line=None)
            if is_primary:
                _set_shape_text_color(_shape_named(slide, f"eval-bar{index}-value"), SEMANTIC_PRIMARY)
                _style_first_paragraph(_shape_named(slide, f"eval-bar{index}-label"), bold=True)
    else:
        # Never invent 35/50/65-style placeholder measurements. When the
        # payload only contains qualitative evidence requirements, remove the
        # chart and let the evidence table use the full lower canvas.
        _remove_prefix(slide, "eval-bar")
    headers, status_rows = _evaluation_status_table(payload)
    if not status_rows:
        _remove_prefix(slide, "eval-status-grid-")
        return
    status_payload = {"headers": headers, "rows": status_rows}
    _fill_grid(slide, "eval-status-grid", status_payload, 3, 3, headers)
    if not explicit_bars:
        has_status_column = "status" in (headers[1].lower() if len(headers) > 1 else "") or "状态" in (headers[1] if len(headers) > 1 else "")
        widths = (2.45, 1.65, 7.15) if has_status_column else (2.35, 6.05, 2.85)
        left = 0.9
        for row_index in range(1, 5):
            x = left
            for col, width in enumerate(widths, start=1):
                shape = _shape_named(slide, f"eval-status-grid-r{row_index}c{col}")
                if shape is not None:
                    _set_geometry(shape, x, 4.0 + (row_index - 1) * 0.68, width, 0.68)
                    _refit_shape_text(shape, max_size=14, min_size=12)
                x += width
    _style_minimal_grid(slide, "eval-status-grid", columns=3, data_rows=min(3, len(status_rows)))
    has_status_column = "status" in (headers[1].lower() if len(headers) > 1 else "") or "状态" in (headers[1] if len(headers) > 1 else "")
    if has_status_column:
        for row_index, row in enumerate(status_rows[:3], start=2):
            status = _display(row[1] if len(row) > 1 else "").lower()
            cell = _shape_named(slide, f"eval-status-grid-r{row_index}c2")
            if any(token in status for token in ("ready", "complete", "completed", "pass", "green", "就绪", "完成", "通过")):
                _set_shape_colors(cell, fill=SEMANTIC_SUCCESS_LIGHT, line=None)
                _set_shape_text_color(cell, RGBColor(52, 133, 91))
            elif any(token in status for token in ("risk", "blocked", "fail", "red", "风险", "阻塞", "失败")):
                _set_shape_colors(cell, fill=SEMANTIC_WARM_LIGHT, line=None)
                _set_shape_text_color(cell, SEMANTIC_WARM)
            else:
                _set_shape_colors(cell, fill=SEMANTIC_PENDING_LIGHT, line=None)


def _bind_risk(slide, payload: dict[str, Any]) -> None:
    records = payload.get("_risk_records") if isinstance(payload.get("_risk_records"), list) else []
    if records:
        has_trigger = any(_display(record.get("trigger")) for record in records if isinstance(record, dict))
        has_owner = any(_display(record.get("owner")) for record in records if isinstance(record, dict))
        has_status = any(_display(record.get("status")) for record in records if isinstance(record, dict))
        total_record_chars = sum(
            len(_display(value))
            for record in records
            if isinstance(record, dict)
            for value in record.values()
        )
        show_trigger = has_trigger and (len(records) <= 3 or total_record_chars <= 520)
        columns: list[tuple[str, str]] = [("Risk", "risk")]
        if show_trigger:
            columns.append(("Trigger", "trigger"))
        columns.extend([("Impact", "impact"), ("Mitigation", "action")])
        if has_owner or has_status:
            columns.append(("Owner / Status", "owner_status"))
        columns = columns[:5]
        rows: list[list[str]] = []
        for record in records[:5]:
            owner_status = " / ".join(
                part for part in [_display(record.get("owner")), _display(record.get("status"))] if part
            )
            values = {**record, "owner_status": owner_status}
            rows.append([_display(values.get(key)) for _header, key in columns])
        if _is_cjk_payload(payload):
            translations = {
                "Risk": "风险",
                "Trigger": "触发条件",
                "Impact": "影响",
                "Mitigation": "应对措施",
                "Owner / Status": "负责人 / 状态",
            }
            headers = [translations.get(header, header) for header, _key in columns]
        else:
            headers = [header for header, _key in columns]
        risk_payload = {"headers": headers, "rows": rows}
    else:
        risk_payload = dict(payload)
        rows = _rows(risk_payload)
        headers = _headers(risk_payload) or ["Risk", "Impact", "Mitigation", "Owner", "Status"]
        columns = [(header, str(index)) for index, header in enumerate(headers[:5])]
    column_count = max(1, min(5, len(columns)))
    _fill_grid(slide, "risk-grid", risk_payload, column_count, 5, headers[:column_count])
    for row_index in range(1, 7):
        for col in range(column_count + 1, 6):
            _remove_named(slide, f"risk-grid-r{row_index}c{col}")
    if column_count == 3:
        widths = (2.55, 4.05, 4.8)
    elif column_count == 4 and columns[-1][1] == "owner_status":
        widths = (2.3, 3.0, 4.95, 1.25)
    elif column_count == 4:
        widths = (2.15, 2.65, 2.65, 4.05)
    else:
        widths = (2.0, 2.2, 2.35, 3.75, 1.2)
    active_rows = min(5, len(rows) or 1)
    row_height = min(0.76, 4.25 / (active_rows + 1))
    left = 0.75
    for row_index in range(1, active_rows + 2):
        x = left
        for col, width in enumerate(widths[:column_count], start=1):
            shape = _shape_named(slide, f"risk-grid-r{row_index}c{col}")
            if shape is not None:
                _set_geometry(shape, x, 1.48 + (row_index - 1) * row_height, width, row_height)
                _refit_shape_text(shape, max_size=13 if row_index == 1 else 12, min_size=11)
            x += width
    risk_highlight = _highlight_row_for_values(rows, ("high", "critical", "严重", "高"), default=None)
    _style_minimal_grid(slide, "risk-grid", columns=column_count, data_rows=active_rows, highlight_row=risk_highlight)
    _set_geometry(_shape_named(slide, "risk-rule"), 0.75, 5.93, sum(widths[:column_count]), 0.64)
    _set_named(
        slide,
        "risk-rule",
        _display(payload.get("rule")) or _localized(
            payload,
            "每项风险都要有触发条件、影响和应对措施",
            "Each risk needs a trigger, impact, and mitigation",
        ),
        max_size=13,
    )
    _set_text_spacing(_shape_named(slide, "risk-rule"), after=2, line_spacing=1.0, middle=True)


def _bind_roadmap(slide, payload: dict[str, Any]) -> None:
    milestones = _structured_list(payload, "milestones")
    points = _payload_points(payload, 7)
    active_count = min(6, len(milestones) if milestones else len(points))
    axis_y = 3.72
    card_width = {1: 3.20, 2: 3.10, 3: 2.75, 4: 2.45, 5: 2.15, 6: 1.85}.get(active_count, 1.85)
    left = 0.85 + card_width / 2
    right = 12.48 - card_width / 2
    centers = (
        [(left + right) / 2]
        if active_count == 1
        else [left + index * (right - left) / (active_count - 1) for index in range(active_count)]
    )
    card_height = 1.30
    for index in range(6):
        item = milestones[index] if index < len(milestones) else (points[index] if index < len(points) else "")
        if isinstance(item, dict):
            value = "\n".join(part for part in [_display(item.get("time")), _display(item.get("title"))] if part)
            status = _display(item.get("status") or item.get("state")).strip().lower()
        else:
            value = _display(item)
            status = ""
        if value:
            center_x = centers[index]
            card = _shape_named(slide, f"roadmap-card-{index}")
            dot = _shape_named(slide, f"roadmap-dot-{index}")
            link = _shape_named(slide, f"roadmap-link-{index}")
            card_y = 1.72 if index % 2 == 0 else 4.30
            _set_geometry(card, center_x - card_width / 2, card_y, card_width, card_height)
            _set_shape_text(card, value, max_size=14, bold=True)
            _set_text_spacing(card, after=2, line_spacing=1.0, middle=True)
            _set_geometry(dot, center_x - 0.13, axis_y - 0.13, 0.26, 0.26)
            if index % 2 == 0:
                start = (float(card.left + card.width / 2), float(card.top + card.height))
                end = (float(dot.left + dot.width / 2), float(dot.top))
            else:
                start = (float(dot.left + dot.width / 2), float(dot.top + dot.height))
                end = (float(card.left + card.width / 2), float(card.top))
            _set_line_endpoints(link, start, end)
            if any(token in status for token in ("complete", "completed", "done", "finish", "已完成", "完成")):
                dot_color = RGBColor(52, 133, 91)
            elif any(token in status for token in ("risk", "blocked", "delay", "at risk", "风险", "阻塞", "延期")):
                dot_color = RGBColor(224, 135, 55)
            elif any(token in status for token in ("current", "active", "progress", "doing", "当前", "进行中")):
                dot_color = RGBColor(47, 111, 174)
            else:
                dot_color = RGBColor(158, 170, 180)
            try:
                dot.fill.solid()
                dot.fill.fore_color.rgb = dot_color
                dot.line.color.rgb = dot_color
            except Exception:
                pass
        else:
            _remove_named(slide, f"roadmap-card-{index}")
            _remove_named(slide, f"roadmap-dot-{index}")
            _remove_named(slide, f"roadmap-link-{index}")
    if active_count:
        axis = _shape_named(slide, "roadmap-axis")
        axis_left = centers[0] if active_count == 1 else centers[0]
        axis_right = centers[-1] if active_count > 1 else centers[0] + 0.01
        _set_geometry(axis, axis_left, axis_y, max(0.01, axis_right - axis_left), 0.01)
    note = _shape_named(slide, "roadmap-note")
    _set_geometry(note, 0.85, 5.98, 11.63, 0.61)
    _set_shape_text(note, _display(payload.get("note")) or (points[-1] if points else ""), max_size=12)
    _set_text_spacing(note, after=2, line_spacing=1.0, middle=True)


def _bind_media(slide, payload: dict[str, Any]) -> None:
    main = _shape_named(slide, "media-showcase-main")
    _set_named(slide, "media-showcase-main", _display(payload.get("visual_label") or payload.get("proof_object")), max_size=18)
    _remove_crosses(slide, "media-showcase-main")
    raw_caption = payload.get("caption") or payload.get("proof_object")
    has_caption = bool(raw_caption or payload.get("source") or payload.get("source_note")) and bool(payload.get("_image_count"))
    if not has_caption:
        _set_geometry(main, 0.60, 1.25, 12.13, 5.45)
        return
    _set_geometry(main, 0.60, 1.25, 12.13, 4.58)
    caption = slide.shapes.add_textbox(Inches(0.75), Inches(5.97), Inches(11.83), Inches(0.61))
    caption.name = "media-caption"
    caption_item = raw_caption if raw_caption else {"source": payload.get("source_note") or payload.get("source")}
    if not isinstance(caption_item, dict) and (payload.get("source") or payload.get("source_note")):
        caption_item = {"caption": caption_item, "source": payload.get("source_note") or payload.get("source")}
    _set_shape_text(caption, _figure_caption(payload, caption_item, 0), max_size=11)
    _style_figure_caption(caption)


_LAYOUT_BINDERS = {
    "basic_content": _bind_basic_content,
    "basic_content_mirror": _bind_basic_content,
    "motivation_compare": _bind_motivation,
    "challenge_map": _bind_challenge,
    "method_pipeline": _bind_pipeline,
    "method_loop": _bind_loop,
    "benchmark_metrics": _bind_benchmark,
    "result_big_numbers": _bind_big_numbers,
    "results_bars": _bind_results_bars,
    "leaderboard_table": _bind_leaderboard,
    "ablation_matrix": _bind_ablation,
    "evidence_grid": _bind_evidence,
    "case_gallery": _bind_case,
    "summary_takeaways": _bind_summary,
    "project_target_map": _bind_project_target,
    "domain_object_map": _bind_domain,
    "technical_route": _bind_route,
    "workpackage_matrix": _bind_workpackages,
    "evaluation_dashboard": _bind_evaluation,
    "risk_action_table": _bind_risk,
    "milestone_roadmap": _bind_roadmap,
    "media_showcase": _bind_media,
}


_IMAGE_SLOTS_BY_LAYOUT = {
    "basic_content": ["basic-image"],
    "basic_content_mirror": ["basic-image"],
    "evidence_grid": [f"evidence-img-{row}-{col}" for row in range(2) for col in range(3)],
    "case_gallery": ["case-main", *[f"case-thumb-{index}" for index in range(4)]],
    "media_showcase": ["media-showcase-main"],
}


def _insert_layout_images(slide, layout_id: str, paths: list[Path]) -> int:
    slot_names = _IMAGE_SLOTS_BY_LAYOUT.get(layout_id) or []
    inserted = 0
    for slot_name, path in zip(slot_names, paths):
        slot = _shape_named(slide, slot_name)
        if slot is None or not path.is_file():
            continue
        left, top, width, height = slot.left, slot.top, slot.width, slot.height
        _remove_shape(slot)
        for suffix in ("-cross-a", "-cross-b"):
            cross = _shape_named(slide, f"{slot_name}{suffix}")
            if cross is not None:
                _remove_shape(cross)
        if _image_visual_mode(path) == "crop":
            picture = slide.shapes.add_picture(str(path), left, top, width, height)
            _crop_picture_to_fill(picture, path, width, height)
        else:
            pic_left, pic_top, pic_width, pic_height = _fit_image(path, left, top, width, height)
            picture = slide.shapes.add_picture(str(path), pic_left, pic_top, pic_width, pic_height)
        picture.name = f"janus-{slot_name}-image"
        inserted += 1
    return inserted


def _crop_picture_to_fill(picture: Any, path: Path, width: int, height: int) -> None:
    try:
        with Image.open(path) as image:
            source_ratio = image.width / max(image.height, 1)
        box_ratio = width / max(height, 1)
        if source_ratio > box_ratio:
            crop = max(0.0, min(0.49, (1.0 - box_ratio / source_ratio) / 2.0))
            picture.crop_left = crop
            picture.crop_right = crop
        elif source_ratio < box_ratio:
            crop = max(0.0, min(0.49, (1.0 - source_ratio / box_ratio) / 2.0))
            picture.crop_top = crop
            picture.crop_bottom = crop
    except Exception:
        pass


def _fit_image(path: Path, left: int, top: int, width: int, height: int) -> tuple[int, int, int, int]:
    try:
        with Image.open(path) as image:
            source_ratio = image.width / max(image.height, 1)
        box_ratio = width / max(height, 1)
        if source_ratio >= box_ratio:
            target_width = width
            target_height = int(width / source_ratio)
            return left, top + (height - target_height) // 2, target_width, target_height
        target_height = height
        target_width = int(height * source_ratio)
        return left + (width - target_width) // 2, top, target_width, target_height
    except Exception:
        return left, top, width, height


def _remove_shape(shape) -> None:
    element = shape.element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _scrub_template_placeholders(slide) -> None:
    for shape in _walk_shapes(slide.shapes):
        if not getattr(shape, "has_text_frame", False):
            continue
        text = _shape_text(shape).strip()
        if not text:
            continue
        is_placeholder = bool(re.search(r"\[[^\]]+\]", text))
        is_fake_metric = bool(re.search(r"\bX{1,3}(?:\.X+)?\b", text, re.IGNORECASE))
        is_guide = bool(re.search(r"(?:template prompt|use this page|use as a project|replace this strip|reading guide|risk page rule)", text, re.IGNORECASE))
        shape_name = str(getattr(shape, "name", "") or "")
        is_visual_slot = shape_name.startswith(("basic-image", "evidence-img-", "case-main", "case-thumb-", "media-showcase-main"))
        is_generic_slot = is_visual_slot and bool(re.fullmatch(
            r"(?:证据|案例|图片|视频|核心案例|核心视觉|核心媒体|演示画面|Evidence|Case|Image|Video)(?:\s*[/／-]?\s*\d+)?",
            text,
            re.IGNORECASE,
        ))
        if is_placeholder or is_fake_metric or is_guide or is_generic_slot:
            cleaned = re.sub(r"\[[^\]]+\]", "", text)
            cleaned = re.sub(r"\b[+\-]?X{1,3}(?:\.X+)?%?\b", "", cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r"(?:template prompt|use this page.*|use as a project.*|replace this strip.*|reading guide.*|risk page rule.*)", "", cleaned, flags=re.IGNORECASE)
            if is_generic_slot:
                cleaned = ""
            cleaned = re.sub(r"\s{2,}", " ", cleaned).strip(" -—:：")
            _set_shape_text(shape, cleaned, max_size=14)


def _remaining_placeholder_texts(slide) -> list[str]:
    result: list[str] = []
    for shape in _walk_shapes(slide.shapes):
        text = re.sub(r"\s+", " ", _shape_text(shape)).strip()
        if not text:
            continue
        shape_name = str(getattr(shape, "name", "") or "")
        is_visual_slot = shape_name.startswith(("basic-image", "evidence-img-", "case-main", "case-thumb-", "media-showcase-main"))
        is_placeholder = bool(re.search(
            r"\[[^\]]+\]|\b[+\-]?X{1,3}(?:\.X+)?%?\b|template prompt|placeholder",
            text,
            re.IGNORECASE,
        ))
        is_generic_visual = is_visual_slot and bool(re.fullmatch(
            r"(?:证据|案例|核心案例|核心视觉|Evidence|Case)(?:\s*\d+)?",
            text,
            re.IGNORECASE,
        ))
        if is_placeholder or is_generic_visual:
            if text not in result:
                result.append(text[:80])
    return result

__all__ = [
    "_LAYOUT_BINDERS",
    "_bind_project_target",
    "_layout_domain_edges",
    "_bind_domain",
    "_bind_route",
    "_bind_workpackages",
    "_bind_evaluation",
    "_bind_risk",
    "_bind_roadmap",
    "_bind_media",
    "_insert_layout_images",
    "_crop_picture_to_fill",
    "_fit_image",
    "_remove_shape",
    "_scrub_template_placeholders",
    "_remaining_placeholder_texts",
]
