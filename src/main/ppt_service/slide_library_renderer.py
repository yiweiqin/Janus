from __future__ import annotations

from copy import deepcopy
import math
import re
from pathlib import Path
from typing import Any, Callable, Iterable

from PIL import Image, ImageFilter, ImageStat
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from ppt_pipeline.content_payloads import (
    _compact_node_label,
    _compact_supporting_text,
    _display,
    _item_parts,
    _message_points,
    _payload_numbers,
    _payload_points,
    _structured_list,
)
from ppt_pipeline.slide_library_core import (
    _scan_slide_library,
    _append_deck_rhythm_warnings,
    _clone_slide,
    _remap_relationship_ids,
    _remove_original_slides,
    _walk_shapes,
    _shape_text,
    _coerce_image_paths,
    _shapes_named,
    _shape_named,
    _sample_text_style,
    _latin_font_for_shape,
    _set_run_typefaces,
    _text_visual_units,
    _shape_fit_font_size,
    _apply_shape_font_size,
    _refit_shape_text,
    _set_shape_text,
    _set_named,
    _set_named_many,
    _remove_named,
    _remove_prefix,
    _remove_crosses,
    _bind_cover,
    _placeholder_type,
)
from ppt_pipeline.slide_library_policy import (
    _append_deck_rhythm_warnings,
    _repair_slide_safe_area,
    _effective_layout_id,
    _layout_family,
    _layout_density,
    _layout_candidates,
    _rhythm_layout_id,
    _normalize_layout_payload,
    _card_body,
    _card_text,
    _remove_occurrences_after,
    _remove_slot,
    _set_geometry,
    _shape_center_and_half_extents,
    _rectangle_edge_point,
    _set_line_endpoints,
    _clip_lines_to_shape_borders,
    _style_text_card,
    _style_pipeline_endpoint,
    _set_shape_colors,
    _set_shape_text_color,
    _truthy_semantic_flag,
    _semantic_primary_index,
    _numeric_value,
    _style_first_paragraph,
    _apply_soft_shadow,
    _apply_layout_depth,
    _set_text_spacing,
    _set_slide_notes,
    _image_information_score,
    _image_visual_mode,
    _source_for_index,
    _figure_caption,
    _style_figure_caption,
    _basic_story,
    _basic_text_is_sparse,
    _configure_basic_text_density,
    _reflow_even_cards,
    _is_cjk_payload,
    _localized,
    _compact_pipeline_endpoint,
)
from ppt_pipeline.slide_library_academic import (
    _bind_basic_content,
    _configure_basic_content_layout,
    _motivation_content,
    _bind_motivation,
    _bind_challenge,
    _bind_pipeline,
    _bind_loop,
    _kpis,
    _fill_kpis,
    _rows,
    _target_row,
    _evaluation_status_table,
    _headers,
    _fill_grid,
    _style_minimal_grid,
    _highlight_row_for_values,
    _evidence_positions,
    _evidence_text_card_positions,
    _reflow_case_thumbnails,
    _bind_benchmark,
    _bind_big_numbers,
    _bars,
    _numeric_ratio,
    _resize_fill,
    _bind_results_bars,
    _bind_leaderboard,
    _bind_ablation,
    _bind_evidence,
    _bind_case,
    _bind_summary,
)
from ppt_pipeline.slide_library_project import (
    _LAYOUT_BINDERS,
    _bind_project_target,
    _layout_domain_edges,
    _bind_domain,
    _bind_route,
    _bind_workpackages,
    _bind_evaluation,
    _bind_risk,
    _bind_roadmap,
    _bind_media,
    _insert_layout_images,
    _crop_picture_to_fill,
    _fit_image,
    _remove_shape,
    _scrub_template_placeholders,
    _remaining_placeholder_texts,
)


KNOWN_LAYOUT_IDS = {
    "basic_content",
    "motivation_compare",
    "challenge_map",
    "method_pipeline",
    "method_loop",
    "benchmark_metrics",
    "result_big_numbers",
    "results_bars",
    "leaderboard_table",
    "ablation_matrix",
    "evidence_grid",
    "case_gallery",
    "summary_takeaways",
    "project_target_map",
    "domain_object_map",
    "technical_route",
    "workpackage_matrix",
    "evaluation_dashboard",
    "risk_action_table",
    "milestone_roadmap",
    "media_showcase",
}

REQUIRED_SHAPES: dict[str, set[str]] = {
    "basic_content": {"basic-text-card", "basic-image", "basic-caption"},
    "motivation_compare": {"motivation-current", "motivation-target", "motivation-gap"},
    "challenge_map": {"challenge-card", "challenge-core"},
    "method_pipeline": {"pipeline-input", "pipeline-step-1", "pipeline-step-5", "pipeline-output"},
    "method_loop": {"loop-center", "loop-node-Plan", "loop-node-Act", "loop-node-Reflect", "loop-node-Observe"},
    "benchmark_metrics": {"bench-kpi1-value", "bench-protocol", "bench-table-r1c1", "bench-table-r5c5"},
    "result_big_numbers": {"big-kpi1-value", "big-kpi3-value", "big-interpretation"},
    "results_bars": {"bar1-label", "bar1-fill", "bar4-value", "bars-note1"},
    "leaderboard_table": {"leaderboard-grid-r1c1", "leaderboard-grid-r7c5", "leader-note1"},
    "ablation_matrix": {"ablation-grid-r1c1", "ablation-grid-r6c6", "ablation-callout"},
    "evidence_grid": {"evidence-img-0-0", "evidence-caption-1-2"},
    "case_gallery": {"case-main", "case-caption-main", "case-thumb-3"},
    "summary_takeaways": {"summary-card1", "summary-card3", "summary-next"},
    "project_target_map": {"target-root", "target-node", "target-grid-r2c4"},
    "domain_object_map": {"domain-layers", "domain-node", "domain-legend"},
    "technical_route": {"route-stage-0", "route-task-3", "route-dependency"},
    "workpackage_matrix": {"wp-grid-r1c1", "wp-grid-r6c6", "wp-note"},
    "evaluation_dashboard": {"eval-kpi1-value", "eval-bar1-fill", "eval-status-grid-r4c3"},
    "risk_action_table": {"risk-grid-r1c1", "risk-grid-r6c5", "risk-rule"},
    "milestone_roadmap": {"roadmap-card-0", "roadmap-card-5", "roadmap-note"},
    "media_showcase": {"media-showcase-main"},
}


def template_slide_library(path: Path) -> dict[str, int]:
    prs = Presentation(str(path))
    return _scan_slide_library(prs)


def validate_slide_library(path: Path, required_layout_ids: Iterable[str] | None = None) -> list[str]:
    errors: list[str] = []
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return [f"模板无法打开：{exc}"]
    width = float(prs.slide_width) / 914400
    height = float(prs.slide_height) / 914400
    if abs(width - 13.333) > 0.08 or abs(height - 7.5) > 0.08:
        errors.append(f"模板尺寸不是 16:9（当前 {width:.3f} x {height:.3f}）")
    library = _scan_slide_library(prs)
    required = set(required_layout_ids or KNOWN_LAYOUT_IDS)
    missing = sorted(required - set(library))
    if missing:
        errors.append(f"模板缺少页面：{', '.join(missing)}")
    for layout_id, slide_index in library.items():
        slide = prs.slides[slide_index]
        names = [shape.name for shape in slide.shapes]
        if not any(name.startswith("tpl-title-") for name in names):
            errors.append(f"{layout_id} 缺少 tpl-title-* 标题 shape")
        missing_shapes = sorted(REQUIRED_SHAPES.get(layout_id, set()) - set(names))
        if missing_shapes:
            errors.append(f"{layout_id} 缺少语义 shape：{', '.join(missing_shapes)}")
    return errors


def render_slide_library_presentation(
    *,
    template_path: Path,
    template_id: str,
    specs: list[Any],
    layout_ids: list[str],
    content_payloads: list[dict[str, Any]],
    image_paths: dict[int, Path | list[Path]] | None = None,
    on_progress: Callable[[dict[str, Any]], None] | None = None,
    progress_phase: str = "render",
    progress_attempt: int = 1,
) -> tuple[Presentation, list[str]]:
    prs = Presentation(str(template_path))
    library = _scan_slide_library(prs)
    original_count = len(prs.slides)
    warnings: list[str] = []
    images = image_paths or {}
    role_counts: dict[str, int] = {}
    recent_roles: list[str] = []

    for index, spec in enumerate(specs, start=1):
        if on_progress is not None:
            on_progress({
                "phase": progress_phase,
                "current_slide": index,
                "total_slides": len(specs),
                "phase_current": index,
                "phase_total": len(specs),
                "slide_title": str(getattr(spec, "title", "") or f"Slide {index}"),
                "attempt": progress_attempt,
                "message": f"正在制作第 {index}/{len(specs)} 页：{getattr(spec, 'title', '') or f'Slide {index}'}",
            })
        payload = content_payloads[index - 1] if index - 1 < len(content_payloads) else {}
        slide_images = _coerce_image_paths(images.get(index))
        if index == 1:
            source_index = 0
            role = "cover"
        else:
            requested_layout = layout_ids[index - 1] if index - 1 < len(layout_ids) else "basic_content"
            base_role = _effective_layout_id(requested_layout, payload, slide_images, library, warnings)
            role = _rhythm_layout_id(
                base_role,
                payload,
                slide_images,
                library,
                role_counts,
                recent_roles,
                warnings,
            )
            source_role = "basic_content" if role == "basic_content_mirror" else role
            source_index = library.get(source_role, library.get("basic_content", 1))
            if source_role not in library:
                warnings.append(f"模板缺少 {source_role}，已回退到 basic_content。")
                role = "basic_content"
                source_index = library.get("basic_content", 1)
            role_counts[role] = role_counts.get(role, 0) + 1
            recent_roles.append(role)
        source = prs.slides[source_index]
        slide = _clone_slide(prs, source)
        if role == "cover":
            _bind_cover(slide, spec, template_id)
        else:
            _bind_content_slide(slide, role, spec, payload, slide_images, warnings)
            _repair_slide_safe_area(slide, role, index, warnings)

    _append_deck_rhythm_warnings(recent_roles, warnings)
    _remove_original_slides(prs, original_count)
    return prs, warnings


def _bind_content_slide(slide, layout_id: str, spec: Any, payload: dict[str, Any], image_paths: list[Path], warnings: list[str]) -> None:
    payload = _normalize_layout_payload(layout_id, payload, spec)
    payload["_image_count"] = len(image_paths)
    title_shape = next((shape for shape in slide.shapes if shape.name.startswith("tpl-title-")), None)
    title = _display(payload.get("title")) or str(getattr(spec, "title", "") or layout_id)
    compact_title = re.sub(r"\s+", " ", title).strip()
    title_size = 28 if len(compact_title) <= 38 else 24 if len(compact_title) <= 48 else 20
    _set_shape_text(title_shape, title, max_size=title_size, bold=True)
    if layout_id == "basic_content_mirror" and title_shape is not None:
        title_shape.name = "tpl-title-basic-mirror"
    slots = payload.get("slots") if isinstance(payload.get("slots"), dict) else {}
    for name, value in slots.items():
        if isinstance(value, list):
            _set_named_many(slide, str(name), [_display(value_item) for value_item in value], max_size=18)
        else:
            _set_named(slide, str(name), _display(value), max_size=18)

    binder = _LAYOUT_BINDERS.get(layout_id, _bind_basic_content)
    binder(slide, payload)
    image_only = False
    if layout_id in {"basic_content", "basic_content_mirror"}:
        image_only = _configure_basic_content_layout(
            slide,
            payload,
            image_paths[0] if image_paths else None,
            mirror=layout_id == "basic_content_mirror",
        )
    if image_paths:
        inserted = _insert_layout_images(slide, layout_id, image_paths)
        if inserted < len(image_paths):
            warnings.append(f"{layout_id} 页面收到 {len(image_paths)} 张图片，仅插入 {inserted} 张。")
    _apply_layout_depth(slide)
    note_parts: list[str] = []
    note_layouts = {
        "domain_object_map",
        "project_target_map",
        "milestone_roadmap",
        "evaluation_dashboard",
        "result_big_numbers",
        "risk_action_table",
    }
    if image_only or layout_id in note_layouts:
        note_parts.append(str(getattr(spec, "speaker_note", "") or "").strip())
    if image_only:
        points = _payload_points(payload, 6)
        if points:
            note_parts.append("Supporting points:\n" + "\n".join(f"- {point}" for point in points))
    if layout_id == "domain_object_map":
        layer_details: list[str] = []
        for item in _structured_list(payload, "layers"):
            label, detail = _item_parts(item)
            if label and detail:
                layer_details.append(f"- {label}: {detail}")
        if layer_details:
            note_parts.append("Layer details:\n" + "\n".join(layer_details))
        node_details: list[str] = []
        for item in _structured_list(payload, "nodes"):
            label, detail = _item_parts(item)
            if label and detail:
                node_details.append(f"- {label}: {detail}")
        if node_details:
            note_parts.append("Node details:\n" + "\n".join(node_details))
    if layout_id == "project_target_map":
        objective = _display(payload.get("objective"))
        if objective:
            note_parts.append("Target objective:\n" + objective)
        target_details: list[str] = []
        for item in _structured_list(payload, "nodes"):
            label, detail = _item_parts(item)
            if label and detail:
                target_details.append(f"- {label}: {detail}")
        if target_details:
            note_parts.append("Target-node details:\n" + "\n".join(target_details))
    if layout_id == "risk_action_table":
        risk_details: list[str] = []
        records = payload.get("_risk_records") if isinstance(payload.get("_risk_records"), list) else []
        for record in records:
            if not isinstance(record, dict):
                continue
            risk = _display(record.get("risk"))
            trigger = _display(record.get("trigger"))
            owner = _display(record.get("owner"))
            status = _display(record.get("status"))
            detail = "; ".join(
                part for part in [
                    f"Trigger: {trigger}" if trigger else "",
                    f"Owner: {owner}" if owner else "",
                    f"Status: {status}" if status else "",
                ] if part
            )
            if risk and detail:
                risk_details.append(f"- {risk}: {detail}")
        if risk_details:
            note_parts.append("Risk details:\n" + "\n".join(risk_details))
    if layout_id == "milestone_roadmap":
        milestone_details: list[str] = []
        for item in _structured_list(payload, "milestones"):
            if not isinstance(item, dict):
                continue
            title = _display(item.get("title") or item.get("label"))
            deliverable = _display(item.get("deliverable") or item.get("detail"))
            if title and deliverable:
                milestone_details.append(f"- {title}: {deliverable}")
        if milestone_details:
            note_parts.append("Milestone deliverables:\n" + "\n".join(milestone_details))
    if layout_id in {"evaluation_dashboard", "result_big_numbers"}:
        kpi_details = [
            f"- {item['value']} / {item['label']}: {item['note']}"
            for item in _kpis(payload)
            if item.get("note")
        ]
        if kpi_details:
            note_parts.append("KPI details:\n" + "\n".join(kpi_details))
    if note_parts:
        _set_slide_notes(slide, "\n\n".join(part for part in note_parts if part))
    _scrub_template_placeholders(slide)
    leftovers = _remaining_placeholder_texts(slide)
    if leftovers:
        warnings.append(f"{layout_id} 页面仍有未绑定占位内容：{' / '.join(leftovers[:3])}")
