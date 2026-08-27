from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

try:
    from PIL import Image
except Exception:
    Image = None

from janus_lab.ppt_renderer import SlideSpec
from ppt_pipeline.image_routing import _source_visual_request_text
from ppt_pipeline.render_primitives import _prefers_english_text
from ppt_pipeline.style_catalog import PPTTemplateDecision, resolve_ppt_template

def _shorten_for_cell(text: str, limit: int = 18) -> str:
    value = re.sub(r"\s+", " ", text or "").replace("…", "").replace("...", "")
    value = value.strip(" -•；;。")
    if len(value) <= limit:
        return value
    if _prefers_english_text(value):
        words = re.findall(r"[A-Za-z][A-Za-z0-9'_-]*|\d+(?:\.\d+)?%?", value)
        if words:
            phrase = words[0]
            soft_limit = max(int(limit * 1.35), 22)
            for word in words[1:]:
                candidate = f"{phrase} {word}"
                if len(candidate) > soft_limit:
                    break
                phrase = candidate
                if len(phrase) >= max(10, int(limit * 0.8)):
                    if len(phrase) >= limit:
                        break
            return phrase
    # Do not render ellipses in final slides. Prefer a complete prefix ending at
    # a natural boundary; fall back to a hard cap without adding punctuation.
    boundary = -1
    for marker in ("。", "；", ";", "，", ",", "、", "：", ":", "和", "与", "/", " "):
        pos = value.rfind(marker, 0, limit + 1)
        if pos > boundary and pos >= max(4, int(limit * 0.45)):
            boundary = pos
    if boundary > 0:
        return value[:boundary].strip(" -•；;。，,、:：/和与")
    return value[:limit].strip(" -•；;。，,、:：/和与")


def _fit_image_to_box(path: Path, x: float, y: float, w: float, h: float) -> tuple[float, float, float, float]:
    if Image is None:
        return x, y, w, h
    with Image.open(path) as img:
        img_w, img_h = img.size
    if img_w <= 0 or img_h <= 0:
        return x, y, w, h
    img_ratio = img_w / img_h
    box_ratio = w / h
    if img_ratio > box_ratio:
        height = w / img_ratio
        return x, y + (h - height) / 2, w, height
    width = h * img_ratio
    return x + (w - width) / 2, y, width, h


def _is_rendered_pdf_page_visual(path: Path) -> bool:
    parts = set(path.parts)
    return (
        re.match(r"page-\d+\.png$", path.name, re.IGNORECASE) is not None
        and ("attachment_assets" in parts or "ppt_source_visuals" in parts or "preview_cache" in parts)
    )


def _source_figure_numbers(spec: SlideSpec | None) -> list[int]:
    if spec is None:
        return []
    numbers: list[int] = []
    for match in re.finditer(r"\bfig(?:ure)?\.?\s*(\d{1,2})\b|图\s*(\d{1,2})\b", _source_visual_request_text(spec), re.IGNORECASE):
        raw = match.group(1) or match.group(2)
        try:
            value = int(raw)
        except (TypeError, ValueError):
            continue
        if value not in numbers:
            numbers.append(value)
    return numbers


def _prefer_deeper_pdf_page_crop(spec: SlideSpec | None) -> bool:
    numbers = _source_figure_numbers(spec)
    return len(numbers) >= 2 and min(numbers) <= 5


def _crop_rendered_pdf_page_visual(img, *, prefer_deeper_crop: bool = False):
    if Image is None:
        return img
    source = img.convert("RGB")
    gray = source.convert("L")
    mask = gray.point(lambda pixel: 255 if pixel < 248 else 0)
    bbox = mask.getbbox()
    if not bbox:
        return img
    x0, y0, x1, y1 = bbox
    pad_x = max(8, int(source.width * 0.025))
    pad_y = max(8, int(source.height * 0.015))
    x0 = max(0, x0 - pad_x)
    x1 = min(source.width, x1 + pad_x)
    y0 = max(0, y0 - pad_y)
    y1 = min(source.height, y1 + pad_y)

    # Rendered paper pages are portrait; PPT needs the dominant figure/table
    # area, not the full article page. Find the first whitespace band after the
    # top proof object and crop there.
    if source.height > source.width * 1.18:
        crop_mask = mask.crop((x0, y0, x1, y1))
        width = max(1, x1 - x0)
        min_band = max(10, int(source.height * 0.012))
        min_y = int((y1 - y0) * 0.18)
        max_y = int((y1 - y0) * 0.72)
        cut_y: int | None = None
        if prefer_deeper_crop:
            content_h = y1 - y0
            band_top = y0 + int(content_h * 0.24)
            band_bottom = y0 + int(content_h * 0.56)
            if band_bottom - band_top >= source.height * 0.16:
                y0, y1 = band_top, band_bottom
                return source.crop((x0, y0, x1, y1))
        else:
            blank_run = 0
            for rel_y in range(min_y, max_y):
                row = crop_mask.crop((0, rel_y, width, rel_y + 1))
                ink = row.getbbox()
                if ink is None:
                    density = 0.0
                else:
                    # Fast enough for page-sized thumbnails; avoids numpy.
                    density = sum(1 for value in row.getdata() if value) / width
                if density < 0.006:
                    blank_run += 1
                    if blank_run >= min_band:
                        cut_y = y0 + rel_y - blank_run // 2
                        break
                else:
                    blank_run = 0
        if cut_y is None:
            cut_y = y0 + int((y1 - y0) * 0.46)
        if cut_y - y0 >= source.height * 0.16:
            y1 = min(y1, cut_y + pad_y)

    return source.crop((x0, y0, x1, y1))


def _prepare_image_for_ppt(path: Path, output_dir: Path, index: int, *, spec: SlideSpec | None = None) -> Path | None:
    if Image is None:
        return path if path.exists() and path.suffix.lower() in {".png", ".jpg", ".jpeg"} else None
    try:
        with Image.open(path) as img:
            img.load()
            if img.mode not in {"RGB", "RGBA"}:
                img = img.convert("RGB")
            if _is_rendered_pdf_page_visual(path):
                img = _crop_rendered_pdf_page_visual(img, prefer_deeper_crop=_prefer_deeper_pdf_page_crop(spec))
            prepared = output_dir / f"slide-{index:02d}-prepared.png"
            img.save(prepared, format="PNG")
            return prepared
    except Exception:
        return None


def _clear_template_slides(prs: Presentation) -> None:
    """Remove example slides from a loaded template while preserving masters/theme."""
    slide_id_list = prs.slides._sldIdLst  # python-pptx private API; stable for deleting slides.
    for slide_id in list(slide_id_list):
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        slide_id_list.remove(slide_id)


def _new_presentation(root: Path, selected_template: str | None) -> tuple[Presentation, PPTTemplateDecision]:
    template = resolve_ppt_template(root, selected_template)
    if template.path:
        try:
            prs = Presentation(str(template.path))
            _clear_template_slides(prs)
            return prs, template
        except Exception:
            pass
    return Presentation(), resolve_ppt_template(root, "none")


TEMPLATE_CHROME_RENDER_TIMEOUT_SECONDS = 60

__all__ = [
    "_clear_template_slides",
    "_crop_rendered_pdf_page_visual",
    "_fit_image_to_box",
    "_is_rendered_pdf_page_visual",
    "_new_presentation",
    "_prefer_deeper_pdf_page_crop",
    "_prepare_image_for_ppt",
    "_shorten_for_cell",
    "_source_figure_numbers",
]
