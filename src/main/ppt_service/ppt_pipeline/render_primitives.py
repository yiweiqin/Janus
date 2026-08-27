from __future__ import annotations

import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

def _rgb(color: str) -> RGBColor:
    value = color.lstrip("#").upper()
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _ppt_len(inches: float):
    return Inches(inches)


def _clean_visible_text(text: str) -> str:
    return re.sub(r"\.{3,}|…+", "", str(text or "")).strip()


def _english_signal(text: str) -> int:
    return len(re.findall(r"[A-Za-z]", str(text or "")))


def _cjk_signal(text: str) -> int:
    return len(re.findall(r"[\u4e00-\u9fff]", str(text or "")))


def _prefers_english_text(text: str) -> bool:
    value = str(text or "")
    english = _english_signal(value)
    cjk = _cjk_signal(value)
    return english >= 24 and english >= cjk * 2


def _prefers_english_spec(spec: "SlideSpec") -> bool:
    return _prefers_english_text(" ".join([spec.title or "", spec.message or ""]))


def _localized_label(spec: "SlideSpec", zh: str, en: str) -> str:
    return en if _prefers_english_spec(spec) else zh


def _add_rect(slide, x: float, y: float, w: float, h: float, fill: str, *, line: str | None = None, radius: bool = False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, _ppt_len(x), _ppt_len(y), _ppt_len(w), _ppt_len(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _apply_soft_shadow(shape, *, opacity: int = 16, blur_pt: float = 5.0, distance_pt: float = 2.0) -> None:
    if shape is None:
        return
    try:
        shape_properties = shape._element.spPr
        for tag in ("a:effectLst", "a:effectDag"):
            existing = shape_properties.find(qn(tag))
            if existing is not None:
                shape_properties.remove(existing)
        effects = OxmlElement("a:effectLst")
        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", str(int(blur_pt * 12700)))
        shadow.set("dist", str(int(distance_pt * 12700)))
        shadow.set("dir", "2700000")
        shadow.set("algn", "ctr")
        shadow.set("rotWithShape", "0")
        color = OxmlElement("a:srgbClr")
        color.set("val", "1E293B")
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(max(0, min(100, opacity)) * 1000))
        color.append(alpha)
        shadow.append(color)
        effects.append(shadow)
        extension = shape_properties.find(qn("a:extLst"))
        if extension is None:
            shape_properties.append(effects)
        else:
            shape_properties.insert(shape_properties.index(extension), effects)
    except Exception:
        pass


def _add_line(slide, x1: float, y1: float, x2: float, y2: float, color: str, *, width: float = 1.2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        _ppt_len(x1),
        _ppt_len(y1),
        _ppt_len(x2),
        _ppt_len(y2),
    )
    line.line.color.rgb = _rgb(color)
    line.line.width = Pt(width)
    return line


def _add_circle(slide, x: float, y: float, size: float, fill: str, *, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, _ppt_len(x), _ppt_len(y), _ppt_len(size), _ppt_len(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _add_outline_circle(slide, x: float, y: float, size: float, color: str, *, width: float = 1.1):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, _ppt_len(x), _ppt_len(y), _ppt_len(size), _ppt_len(size))
    shape.fill.background()
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(width)
    return shape


def _add_outline_rect(slide, x: float, y: float, w: float, h: float, color: str, *, width: float = 1.1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _ppt_len(x), _ppt_len(y), _ppt_len(w), _ppt_len(h))
    shape.fill.background()
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(width)
    return shape


def _template_icon_kind(text: str) -> str:
    value = (text or "").lower()
    if any(key in value for key in ["rag", "检索", "retrieve", "top-k", "topk", "search", "exemplar"]):
        return "search"
    if any(key in value for key in ["cohesion", "反思", "反馈", "loop", "迭代", "优化", "gradient"]):
        return "loop"
    if any(key in value for key in ["case", "案例", "提升", "增长", "结果", "contribution"]):
        return "trend"
    if any(key in value for key in ["图像", "视频", "多模态", "image", "video", "tv", "it", "iv"]):
        return "media"
    if any(key in value for key in ["历史", "推荐", "数据", "dataset", "rec", "historical", "benchmark"]):
        return "stack"
    if any(key in value for key in ["用户", "画像", "profile", "preference", "偏好"]):
        return "user"
    if any(key in value for key in ["挑战", "风险", "问题", "缺口", "challenge", "risk"]):
        return "alert"
    return "node"


def _add_template_icon(slide, kind: str, x: float, y: float, size: float, color: str) -> None:
    white = "FFFFFF"
    _add_circle(slide, x, y, size, color)
    cx = x + size / 2
    cy = y + size / 2
    if kind == "search":
        _add_outline_circle(slide, x + size * 0.24, y + size * 0.22, size * 0.34, white, width=1.2)
        _add_line(slide, x + size * 0.55, y + size * 0.55, x + size * 0.72, y + size * 0.72, white, width=1.4)
    elif kind == "loop":
        _add_outline_circle(slide, x + size * 0.24, y + size * 0.23, size * 0.46, white, width=1.2)
        _add_line(slide, x + size * 0.61, y + size * 0.26, x + size * 0.75, y + size * 0.30, white, width=1.2)
        _add_line(slide, x + size * 0.61, y + size * 0.26, x + size * 0.66, y + size * 0.41, white, width=1.2)
    elif kind == "trend":
        _add_line(slide, x + size * 0.22, y + size * 0.67, x + size * 0.42, y + size * 0.52, white, width=1.4)
        _add_line(slide, x + size * 0.42, y + size * 0.52, x + size * 0.58, y + size * 0.58, white, width=1.4)
        _add_line(slide, x + size * 0.58, y + size * 0.58, x + size * 0.76, y + size * 0.34, white, width=1.4)
        _add_line(slide, x + size * 0.68, y + size * 0.34, x + size * 0.76, y + size * 0.34, white, width=1.2)
        _add_line(slide, x + size * 0.76, y + size * 0.34, x + size * 0.76, y + size * 0.43, white, width=1.2)
    elif kind == "media":
        _add_outline_rect(slide, x + size * 0.22, y + size * 0.26, size * 0.54, size * 0.42, white, width=1.1)
        _add_line(slide, x + size * 0.35, y + size * 0.56, x + size * 0.48, y + size * 0.43, white, width=1.2)
        _add_line(slide, x + size * 0.48, y + size * 0.43, x + size * 0.65, y + size * 0.58, white, width=1.2)
    elif kind == "stack":
        _add_rect(slide, x + size * 0.24, y + size * 0.28, size * 0.52, size * 0.09, white)
        _add_rect(slide, x + size * 0.24, y + size * 0.45, size * 0.52, size * 0.09, white)
        _add_rect(slide, x + size * 0.24, y + size * 0.62, size * 0.52, size * 0.09, white)
    elif kind == "user":
        _add_outline_circle(slide, cx - size * 0.12, y + size * 0.22, size * 0.24, white, width=1.1)
        _add_outline_circle(slide, cx - size * 0.25, y + size * 0.50, size * 0.50, white, width=1.1)
    elif kind == "alert":
        _add_line(slide, cx, y + size * 0.22, x + size * 0.76, y + size * 0.70, white, width=1.3)
        _add_line(slide, x + size * 0.76, y + size * 0.70, x + size * 0.24, y + size * 0.70, white, width=1.3)
        _add_line(slide, x + size * 0.24, y + size * 0.70, cx, y + size * 0.22, white, width=1.3)
        _add_line(slide, cx, y + size * 0.38, cx, y + size * 0.54, white, width=1.2)
    else:
        _add_outline_circle(slide, x + size * 0.31, y + size * 0.31, size * 0.38, white, width=1.2)
        _add_line(slide, cx, y + size * 0.20, cx, y + size * 0.30, white, width=1.1)
        _add_line(slide, cx, y + size * 0.70, cx, y + size * 0.80, white, width=1.1)
        _add_line(slide, x + size * 0.20, cy, x + size * 0.30, cy, white, width=1.1)
        _add_line(slide, x + size * 0.70, cy, x + size * 0.80, cy, white, width=1.1)


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    fill: str | None = None,
    margin: float = 0.08,
):
    box = slide.shapes.add_textbox(_ppt_len(x), _ppt_len(y), _ppt_len(w), _ppt_len(h))
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb(fill)
    else:
        box.fill.background()
    box.line.fill.background()
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = _ppt_len(margin)
    frame.margin_right = _ppt_len(margin)
    frame.margin_top = _ppt_len(margin)
    frame.margin_bottom = _ppt_len(margin)
    lines = [line for line in _clean_visible_text(text or " ").split("\n") if line.strip()] or [" "]
    for idx, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = line.strip()
        font = run.font
        font.name = "Microsoft YaHei"
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = _rgb(color)
    return box


def _visual_len(text: str) -> float:
    """Approximate visual width of text in 'full-width character' units.

    CJK characters count as 1.0, latin/digits/spaces as ~0.55. Used to estimate
    how much text fits in a box so we can pick a font size that never overflows.
    """
    total = 0.0
    for ch in text or "":
        total += 1.0 if ord(ch) > 0x2E7F else 0.55
    return total


def _fit_font_size(
    text: str,
    width: float,
    height: float,
    *,
    max_size: int,
    min_size: int = 8,
    margin: float = 0.08,
) -> int:
    """Pick the largest font size (pt) at which `text` fits within width x height.

    Estimates characters-per-line from the usable width and required lines from
    the visual length, then shrinks until the wrapped block fits the box height.
    """
    usable_w = max(width - 2 * margin, 0.4)
    usable_h = max(height - 2 * margin, 0.2)
    units = max(_visual_len(text), 1.0)
    # explicit line breaks force at least that many lines
    forced_lines = (text or " ").count("\n") + 1
    for size in range(max_size, min_size - 1, -1):
        char_w = size / 72.0  # one full-width glyph ~= font size in inches
        line_h = (size / 72.0) * 1.32  # line height with leading
        chars_per_line = max(int(usable_w / char_w), 1)
        needed_lines = max(forced_lines, -(-int(units) // chars_per_line))  # ceil
        if needed_lines * line_h <= usable_h:
            return size
    return min_size


def _add_autofit_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    max_size: int,
    min_size: int = 8,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    fill: str | None = None,
    line_color: str | None = None,
    margin: float = 0.1,
) -> int:
    """Text box that shrinks its font so the content never overflows the box."""
    text = _clean_visible_text(text)
    size = _fit_font_size(text, w, h, max_size=max_size, min_size=min_size, margin=margin)
    box = slide.shapes.add_textbox(_ppt_len(x), _ppt_len(y), _ppt_len(w), _ppt_len(h))
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb(fill)
    else:
        box.fill.background()
    if line_color:
        box.line.color.rgb = _rgb(line_color)
        box.line.width = Pt(0.75)
    else:
        box.line.fill.background()
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = _ppt_len(margin)
    frame.margin_right = _ppt_len(margin)
    frame.margin_top = _ppt_len(min(margin, 0.06))
    frame.margin_bottom = _ppt_len(min(margin, 0.06))
    lines = [ln for ln in str(text or " ").split("\n")] or [" "]
    for idx, ln in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = 1.12
        paragraph.space_after = Pt(2)
        run = paragraph.add_run()
        run.text = ln.strip() or " "
        font = run.font
        font.name = "Microsoft YaHei"
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = _rgb(color)
    return size



__all__ = [
    "_add_autofit_text",
    "_add_circle",
    "_add_line",
    "_add_outline_circle",
    "_add_outline_rect",
    "_add_rect",
    "_add_template_icon",
    "_add_text",
    "_apply_soft_shadow",
    "_clean_visible_text",
    "_fit_font_size",
    "_localized_label",
    "_ppt_len",
    "_prefers_english_spec",
    "_prefers_english_text",
    "_rgb",
    "_template_icon_kind",
    "_visual_len",
]
