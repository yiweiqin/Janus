from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import replace
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

try:
    from PIL import Image
except Exception:
    Image = None

from janus_lab.ppt_renderer import SlideSpec
from slide_library_renderer import KNOWN_LAYOUT_IDS as SLIDE_LIBRARY_LAYOUT_IDS, validate_slide_library
from ppt_pipeline.content_parsing import *
from ppt_pipeline.image_routing import (
    _requested_source_figures,
    _source_path_figure_numbers,
    _source_visual_request_text,
    _spec_layout_id,
)
from ppt_pipeline.office_conversion import *
from ppt_pipeline.presentation_assets import *
from ppt_pipeline.render_primitives import *
from ppt_pipeline.style_catalog import PPTStyleDecision, PPTTemplateDecision, TEMPLATE_RENDER_PROFILES

PPT_LAYOUT_REGISTRY = Path("departments/ppt_department/templates/layouts/layout_registry.json")

def _palette_extras(style: PPTStyleDecision) -> dict[str, str]:
    p = style.palette
    if p["bg"].upper() in {"0B1020", "111827", "0F172A"}:
        return {
            "soft": "172036",
            "soft2": "1E293B",
            "grid": "243047",
            "warm": "F59E0B",
            "good": "A3E635",
            "danger": "FB7185",
        }
    return {
        "soft": "EEF4FF",
        "soft2": "ECFEFF",
        "grid": "D7DEE8",
        "warm": "F59E0B",
        "good": "10B981",
        "danger": "EF4444",
    }


def _is_scut_academic_chrome(style: PPTStyleDecision) -> bool:
    p = style.palette
    return p.get("ink", "").upper() == "3A1F1B" and p.get("accent", "").upper() == "B61918"


def _is_hitsz_academic_chrome(style: PPTStyleDecision) -> bool:
    p = style.palette
    return p.get("ink", "").upper() == "12364A" and p.get("accent", "").upper() == "0B5E7A"


def _academic_ref_colors(style: PPTStyleDecision) -> dict[str, str]:
    if _is_scut_academic_chrome(style):
        return {
            "title": "16608A",
            "line": "B8D7E8",
            "line2": "F2C1C1",
            "fill": "FFFFFF",
            "fill_blue": "F7FBFF",
            "fill_yellow": "FFFCF2",
            "fill_red": "FFF8F5",
            "fill_green": "F6FFFA",
        }
    ex = _palette_extras(style)
    p = style.palette
    return {
        "title": p["accent2"],
        "line": ex["grid"],
        "line2": p["accent"],
        "fill": "FFFFFF",
        "fill_blue": ex["soft"],
        "fill_yellow": "FFF7E6",
        "fill_red": "FFF1F2",
        "fill_green": "ECFDF5",
    }


def _academic_ref_fill(style: PPTStyleDecision, index: int = 0) -> str:
    ref = _academic_ref_colors(style)
    fills = [ref["fill_blue"], ref["fill_yellow"], ref["fill_red"], ref["fill_green"]]
    return fills[index % len(fills)] if _is_scut_academic_chrome(style) else "FFFFFF"


def _add_academic_ref_panel(
    slide,
    style: PPTStyleDecision,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str | None = None,
    accent: str | None = None,
    title: str | None = None,
    radius: bool = True,
) -> None:
    ref = _academic_ref_colors(style)
    panel_fill = fill or (_academic_ref_fill(style, 0) if _is_scut_academic_chrome(style) else "FFFFFF")
    if _is_scut_academic_chrome(style) and h < 2.7:
        radius = False
    panel = _add_rect(slide, x, y, w, h, panel_fill, line=ref["line"], radius=radius)
    if radius and ((w >= 2.8 and h >= 1.15) or (w >= 4.5 and h >= 0.8)):
        _apply_soft_shadow(panel)
    if accent:
        _add_rect(slide, x, y, w, 0.06, accent, radius=False)
    if title:
        _add_text(slide, title, x + 0.18, y + 0.14, max(0.4, w - 0.36), 0.24, size=8, color=ref["title"], bold=True, margin=0)


def _academic_visible_text(style: PPTStyleDecision, text: str, limit: int) -> str:
    value = _clean_visible_text(text)
    if _is_scut_academic_chrome(style):
        return _shorten_for_cell(value, limit)
    return value


def _academic_sparse_points_boost(style: PPTStyleDecision, points: list[str]) -> int:
    if not _is_scut_academic_chrome(style) or not points:
        return 0
    cleaned = [point for point in points if point and not _is_layout_instruction_text(point)]
    if not cleaned:
        return 0
    units = sum(_visual_len(point) for point in cleaned)
    if len(cleaned) <= 1 and units <= 44:
        return 4
    if len(cleaned) <= 2 and units <= 82:
        return 3
    if len(cleaned) <= 3 and units <= 120:
        return 2
    return 0


def _academic_sparse_spec_boost(style: PPTStyleDecision, spec: SlideSpec, points: list[str] | None = None) -> int:
    return _academic_sparse_points_boost(style, points if points is not None else _template_rich_bullets(spec, limit=4))


def _add_academic_plain_list(
    slide,
    style: PPTStyleDecision,
    points: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    start: int = 1,
    max_items: int = 4,
    max_size: int = 11,
) -> float:
    """SCUT academic list: text on white with separators, no card containers."""
    if not points:
        return y
    p = style.palette
    ex = _palette_extras(style)
    ref = _academic_ref_colors(style)
    accents = [p["accent"], p["accent2"], ex["warm"], ex["good"], ex["danger"]]
    visible = points[:max_items]
    boost = _academic_sparse_points_boost(style, visible)
    row_h = min(0.78, h / max(len(visible), 1))
    if _is_scut_academic_chrome(style):
        row_h = min(1.14 if boost >= 3 else 1.04 if boost else 0.96, h / max(len(visible), 1))
    for i, point in enumerate(visible):
        cy = y + i * row_h
        accent = accents[(start + i - 1) % len(accents)]
        _add_line(slide, x, cy + row_h - 0.02, x + w, cy + row_h - 0.02, ref["line"], width=0.65)
        _add_rect(slide, x + 0.08, cy + 0.17, 0.08, max(0.24, row_h - 0.34), accent)
        _add_autofit_text(
            slide,
            _academic_visible_text(style, point, 34),
            x + 0.34,
            cy + 0.07,
            w - 0.4,
            row_h - 0.12,
            max_size=max_size + ((2 + boost) if _is_scut_academic_chrome(style) else 0),
            min_size=(10 if boost >= 2 else 9) if _is_scut_academic_chrome(style) else 7,
            color=p["ink"],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.01,
        )
    return y + len(visible) * row_h


def _add_academic_plain_note(
    slide,
    style: PPTStyleDecision,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    label: str = "说明",
) -> None:
    if not text:
        return
    p = style.palette
    ref = _academic_ref_colors(style)
    _add_line(slide, x, y, x + w, y, ref["line"], width=0.8)
    _add_rect(slide, x, y + 0.13, 0.08, h - 0.2, p["accent"])
    _add_text(slide, label, x + 0.22, y + 0.1, 0.9, 0.26, size=10 if _is_scut_academic_chrome(style) else 8, color=p["accent"], bold=True, margin=0)
    _add_autofit_text(
        slide,
        _academic_visible_text(style, text, 56),
        x + 1.02,
        y + 0.08,
        w - 1.08,
        h - 0.12,
        max_size=12 if _is_scut_academic_chrome(style) else 10,
        min_size=9 if _is_scut_academic_chrome(style) else 7,
        color=p["muted"],
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )


def _template_profile(template: PPTTemplateDecision | None) -> dict[str, Any]:
    if template is None:
        return {}
    return TEMPLATE_RENDER_PROFILES.get(template.template_id, {})


def _style_for_template(style: PPTStyleDecision, template: PPTTemplateDecision) -> PPTStyleDecision:
    profile = _template_profile(template)
    palette = profile.get("palette")
    if not isinstance(palette, dict):
        return style
    merged = {**style.palette, **{key: str(value) for key, value in palette.items()}}
    return replace(style, palette=merged)


def _template_footer_label(style: PPTStyleDecision, template: PPTTemplateDecision | None = None) -> str:
    profile = _template_profile(template)
    footer = profile.get("footer")
    if footer:
        return f"{style.label} · {footer}"
    return style.label


def _add_footer(
    slide,
    style: PPTStyleDecision,
    index: int,
    total: int,
    template: PPTTemplateDecision | None = None,
) -> None:
    muted = style.palette["muted"]
    _add_text(slide, _template_footer_label(style, template), 0.42, 6.95, 5.0, 0.3, size=8, color=muted, margin=0.02)
    _add_text(slide, f"{index} / {total}", 11.6, 6.95, 1.3, 0.3, size=8, color=muted, align=PP_ALIGN.RIGHT, margin=0.02)


def _add_slide_background(slide, style: PPTStyleDecision, *, variant: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_rect(slide, 0, 0, 13.333, 7.5, p["bg"])
    profile = _template_profile(template)
    if template and template.template_id == "hitsz":
        _add_rect(slide, 0, 0, 13.333, 0.42, p["accent"])
        _add_rect(slide, 0.42, 6.78, 12.45, 0.04, p["accent"])
    elif template and template.template_id == "scut":
        _add_rect(slide, 0, 0, 13.333, 0.20, p["accent"])
        _add_rect(slide, 0.42, 6.78, 10.8, 0.04, p["accent"])
        _add_rect(slide, 11.35, 6.72, 1.25, 0.12, p["accent2"])
    else:
        # thin accent header strip + a single restrained side accent — keep decoration minimal
        _add_rect(slide, 0, 0, 13.333, 0.16, p["accent"])
        _add_rect(slide, 0.0, 0.16, 4.4, 0.05, p["accent2"])
    _add_rect(slide, 0.36, 6.86, 12.6, 0.02, ex["grid"])


def _add_title_system(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, *, kicker: str = "", template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    if kicker:
        _add_text(slide, kicker.upper(), 0.42, 0.30, 4.0, 0.24, size=8, color=p["accent"], bold=True, margin=0)
    _add_autofit_text(
        slide, spec.title or f"Slide {index}", 0.40, 0.56, 11.3, 0.74,
        max_size=26, min_size=15, color=p["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0.02,
    )
    _add_rect(slide, 0.42, 1.34, 0.85, 0.06, p["accent"])
    _add_rect(slide, 1.34, 1.34, 0.32, 0.06, p["accent2"])
    _add_footer(slide, style, index, total, template=template)


def _split_points(text: str, *, limit: int = 3) -> list[str]:
    value = re.sub(r"\s+", " ", text or "").strip()
    if not value:
        return []
    parts = re.split(r"[；;。]\s*|\s*/\s*|\n+|(?:\s+[1-9][.)、])", value)
    points = [p.strip(" -•·，,") for p in parts if p.strip(" -•·，,")]
    if len(points) <= 1:
        chunks = re.split(r"[，,]\s*", value)
        points = [p.strip() for p in chunks if p.strip()]
    if len(points) <= limit:
        return points
    head = points[: limit - 1]
    tail = "；".join(points[limit - 1 :])
    return [*head, tail]


def _template_source_slide_index(template: PPTTemplateDecision, role: str) -> int | None:
    profile = _template_profile(template)
    source_slides = profile.get("source_slides")
    if isinstance(source_slides, dict):
        value = source_slides.get(role)
        try:
            return int(value)
        except (TypeError, ValueError):
            return None
    return _template_layout_index(template, role)


def _template_chrome_pdf_cache_path(root: Path, path: Path) -> Path:
    stat = path.stat()
    digest = hashlib.sha256(
        f"{path.resolve()}:{stat.st_mtime_ns}:{stat.st_size}".encode("utf-8", errors="replace")
    ).hexdigest()
    cache_dir = root / "data" / "preview_cache" / "ppt_template_chrome"
    cache_dir.mkdir(parents=True, exist_ok=True)
    return cache_dir / f"{digest}.pdf"


def _convert_template_to_pdf(root: Path, path: Path) -> Path:
    cached = _template_chrome_pdf_cache_path(root, path)
    if cached.is_file() and cached.stat().st_size > 0:
        return cached

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pdf_path, error = _convert_powerpoint_to_pdf_windows(
            root,
            path,
            cached,
            timeout_seconds=TEMPLATE_CHROME_RENDER_TIMEOUT_SECONDS,
        )
        if pdf_path is not None:
            return pdf_path
        raise RuntimeError(error or "Office renderer is not installed")

    with tempfile.TemporaryDirectory(prefix="opl-ppt-template-render-") as tmp:
        tmp_path = Path(tmp)
        out_dir = tmp_path / "out"
        profile_dir = tmp_path / "profile"
        out_dir.mkdir()
        profile_dir.mkdir()
        env = os.environ.copy()
        env["HOME"] = str(tmp_path)
        cmd = [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nolockcheck",
            "--nodefault",
            f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(path),
        ]
        result = subprocess.run(
            cmd,
            cwd=str(root),
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=TEMPLATE_CHROME_RENDER_TIMEOUT_SECONDS,
            check=False,
        )
        rendered = out_dir / f"{path.stem}.pdf"
        if result.returncode != 0 or not rendered.is_file():
            message = (result.stderr or result.stdout or "Template render failed").strip()
            pdf_path, error = _convert_powerpoint_to_pdf_windows(
                root,
                path,
                cached,
                timeout_seconds=TEMPLATE_CHROME_RENDER_TIMEOUT_SECONDS,
            )
            if pdf_path is not None:
                return pdf_path
            raise RuntimeError(message[-300:])
        shutil.copyfile(rendered, cached)
    return cached


def _render_template_chrome_backgrounds(
    root: Path,
    template: PPTTemplateDecision,
    roles: list[str],
    output_dir: Path,
) -> dict[str, Path]:
    if not template.path:
        return {}
    role_to_page = {
        role: _template_source_slide_index(template, role)
        for role in sorted(set(roles))
    }
    role_to_page = {role: page for role, page in role_to_page.items() if page is not None}
    if not role_to_page:
        return {}

    try:
        import fitz  # type: ignore

        pdf_path = _convert_template_to_pdf(root, template.path)
        chrome_dir = output_dir / "template_chrome"
        chrome_dir.mkdir(parents=True, exist_ok=True)
        backgrounds: dict[str, Path] = {}
        doc = fitz.open(str(pdf_path))
        try:
            for role, page_index in role_to_page.items():
                if page_index < 0 or page_index >= doc.page_count:
                    continue
                destination = chrome_dir / f"{template.template_id}-{role}.png"
                if not destination.is_file() or destination.stat().st_size <= 0:
                    page = doc.load_page(page_index)
                    scale = max(1.0, 1600 / max(float(page.rect.width), 1.0))
                    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                    pix.save(str(destination))
                backgrounds[role] = destination
        finally:
            doc.close()
        return backgrounds
    except Exception:
        return {}


def _template_source_placeholder_boxes(template: PPTTemplateDecision, role: str) -> dict[int, tuple[float, float, float, float]]:
    page_index = _template_source_slide_index(template, role)
    if template.path is None or page_index is None:
        return {}
    try:
        prs = Presentation(str(template.path))
        if page_index < 0 or page_index >= len(prs.slides):
            return {}
        slide = prs.slides[page_index]
        boxes: dict[int, tuple[float, float, float, float]] = {}
        for shape in slide.placeholders:
            try:
                boxes[int(shape.placeholder_format.idx)] = _shape_box_inches(shape)
            except Exception:
                continue
        return boxes
    except Exception:
        return {}


def _box_or_default(
    boxes: dict[int, tuple[float, float, float, float]],
    idx: int,
    fallback: tuple[float, float, float, float],
) -> tuple[float, float, float, float]:
    return boxes.get(idx) or fallback


def _blank_layout(prs: Presentation):
    if len(prs.slide_layouts) > 6:
        return prs.slide_layouts[6]
    if len(prs.slide_layouts) > 0:
        return prs.slide_layouts[-1]
    return Presentation().slide_layouts[6]


def _template_layout_index(template: PPTTemplateDecision, role: str) -> int | None:
    profile = _template_profile(template)
    layouts = profile.get("layouts")
    if not isinstance(layouts, dict):
        return None
    value = layouts.get(role)
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _template_slide_layout(prs: Presentation, template: PPTTemplateDecision, role: str):
    index = _template_layout_index(template, role)
    if index is None or index < 0 or index >= len(prs.slide_layouts):
        return None
    return prs.slide_layouts[index]


def _uses_template_layouts(template: PPTTemplateDecision) -> bool:
    return bool(template.path and _template_profile(template).get("layouts"))


def _slide_role(index: int, total: int, spec: SlideSpec, template: PPTTemplateDecision) -> str:
    title_text = spec.title or ""
    text = f"{title_text} {spec.message or ''} {spec.visual or ''}"
    if index == 1:
        return "cover"
    if re.search(r"(目录|agenda|outline)", text, re.IGNORECASE) and _template_layout_index(template, "directory") is not None:
        return "directory"
    if (
        _template_layout_index(template, "directory") is not None
        and re.search(r"(章节|第\s*[一二三四五六七八九十\d]+\s*(?:章|节|部分)|chapter|section|part\s*\d+)", title_text, re.IGNORECASE)
    ):
        return "directory"
    if (
        index == total
        and _template_layout_index(template, "ending") is not None
        and re.search(r"(谢谢|thanks|q&a|qa|讨论|致谢|结束)", text, re.IGNORECASE)
    ):
        return "ending"
    return "content"


def _placeholder_by_idx(slide, idx: int):
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.idx == idx:
                return shape
        except Exception:
            continue
    return None


def _placeholder_type_name(shape) -> str:
    try:
        return str(shape.placeholder_format.type).upper()
    except Exception:
        return ""


def _first_text_placeholder(slide, *, exclude: set[int] | None = None):
    exclude = exclude or set()
    for shape in slide.placeholders:
        try:
            idx = shape.placeholder_format.idx
        except Exception:
            continue
        if idx in exclude or not getattr(shape, "has_text_frame", False):
            continue
        return shape
    return None


def _shape_box_inches(shape) -> tuple[float, float, float, float]:
    return (
        float(shape.left) / 914400,
        float(shape.top) / 914400,
        float(shape.width) / 914400,
        float(shape.height) / 914400,
    )


def _set_text_frame(
    shape,
    lines: list[str],
    *,
    size: int,
    color: str,
    bold_first: bool = False,
    bullet: bool = False,
    align=PP_ALIGN.LEFT,
) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    clean_lines = [_clean_visible_text(line) for line in lines if _clean_visible_text(line)] or [" "]
    for idx, line in enumerate(clean_lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.level = 0
        paragraph.space_after = Pt(5)
        if bullet:
            paragraph.text = ""
        run = paragraph.add_run()
        run.text = line
        font = run.font
        font.name = "Microsoft YaHei"
        font.size = Pt(size)
        font.bold = bool(bold_first and idx == 0)
        font.color.rgb = _rgb(color)


def _set_title_placeholder(slide, spec: SlideSpec, style: PPTStyleDecision) -> None:
    title_shape = _placeholder_by_idx(slide, 0) or _first_text_placeholder(slide)
    if title_shape is not None:
        _set_text_frame(title_shape, [spec.title or "PPT Draft"], size=30, color=style.palette["ink"], bold_first=True)


def _set_footer_placeholders(slide, index: int, total: int, template: PPTTemplateDecision) -> None:
    date_shape = _placeholder_by_idx(slide, 10)
    if date_shape is not None and "DATE" in _placeholder_type_name(date_shape):
        _set_text_frame(date_shape, [datetime.now().strftime("%Y / %m / %d")], size=9, color="666666")
    footer_shape = _placeholder_by_idx(slide, 11)
    if footer_shape is not None and "FOOTER" in _placeholder_type_name(footer_shape):
        _set_text_frame(footer_shape, [_template_profile(template).get("footer", "") or ""], size=9, color="666666", align=PP_ALIGN.CENTER)
    number_shape = _placeholder_by_idx(slide, 12)
    if number_shape is not None and "SLIDE_NUMBER" in _placeholder_type_name(number_shape):
        _set_text_frame(number_shape, [str(index)], size=9, color="666666", align=PP_ALIGN.RIGHT)


def _add_template_chrome_background(slide, background_path: Path) -> None:
    slide.shapes.add_picture(
        str(background_path),
        _ppt_len(0),
        _ppt_len(0),
        width=_ppt_len(13.333),
        height=_ppt_len(7.5),
    )


def _render_template_chrome_cover(
    slide,
    background_path: Path,
    spec: SlideSpec,
    style: PPTStyleDecision,
    template: PPTTemplateDecision,
) -> None:
    _add_template_chrome_background(slide, background_path)
    boxes = _template_source_placeholder_boxes(template, "cover")
    title_box = _box_or_default(boxes, 0, (1.2, 2.05, 10.9, 1.25))
    subtitle_box = _box_or_default(boxes, 1, (1.5, 3.45, 10.3, 0.8))
    title_size = 34 if template.template_id == "hitsz" else 32
    subtitle_size = 17 if template.template_id == "hitsz" else 16
    _add_autofit_text(
        slide,
        spec.title or "PPT Draft",
        *title_box,
        max_size=title_size,
        min_size=18,
        color=style.palette["ink"],
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.04,
    )
    subtitle = (_meaningful_bullets(spec.message or "", limit=1) or [""])[0]
    if subtitle.strip():
        _add_autofit_text(
            slide,
            subtitle,
            *subtitle_box,
            max_size=subtitle_size,
            min_size=10,
            color=style.palette["muted"],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.04,
        )


def _template_chrome_boxes(
    template: PPTTemplateDecision,
    role: str,
) -> tuple[tuple[float, float, float, float], tuple[float, float, float, float]]:
    boxes = _template_source_placeholder_boxes(template, role)
    title_box = _box_or_default(boxes, 0, (0.68, 0.28, 12.0, 0.77))
    content_box = _box_or_default(boxes, 1, _box_or_default(boxes, 10, (0.68, 1.23, 12.0, 5.59)))
    if template.template_id == "scut":
        title_box = (title_box[0], title_box[1] + 0.01, min(title_box[2], 9.55), title_box[3])
        content_box = (content_box[0], max(content_box[1] - 0.06, 1.18), content_box[2], min(content_box[3] + 0.24, 5.58))
    elif template.template_id == "hitsz":
        title_box = (title_box[0], title_box[1] + 0.02, min(title_box[2], 9.35), title_box[3])
        content_box = (content_box[0], max(content_box[1] - 0.08, 1.38), content_box[2], min(content_box[3] + 0.3, 5.42))
    return title_box, content_box


def _add_template_chrome_title(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    title_box: tuple[float, float, float, float],
) -> None:
    title_color = _academic_ref_colors(style)["title"] if _is_scut_academic_chrome(style) else style.palette["ink"]
    _add_autofit_text(
        slide,
        spec.title or "Slide",
        *title_box,
        max_size=27 if _is_scut_academic_chrome(style) else 30,
        min_size=14 if _is_scut_academic_chrome(style) else 15,
        color=title_color,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.03,
    )


def _template_slide_bullets(spec: SlideSpec, *, limit: int = 5) -> list[str]:
    return _slide_display_points(spec, limit=limit)


def _template_rich_bullets(spec: SlideSpec, *, limit: int = 5) -> list[str]:
    visible = _template_slide_bullets(spec, limit=limit + 3)
    candidates = _dedupe_visible_points(visible)
    substantive = [point for point in candidates if not _is_low_substance_bullet(point)]
    points = substantive[:limit]
    if len(points) < min(limit, 3):
        for candidate in candidates:
            if candidate not in points:
                points.append(candidate)
            if len(points) >= limit:
                break
    return points[:limit]


def _template_concept_terms(spec: SlideSpec, *, limit: int = 5) -> list[str]:
    # Visible concept labels must come from audience-facing content. Use the
    # cleaned display-point pipeline so renderer instructions never become
    # short labels.
    display_points = _slide_display_points(spec, limit=limit + 2, include_title=False)
    raw = "；".join(display_points) or spec.message or ""
    pieces = re.split(r"[、,，;；。:：/()\[\]（）【】\s]+|和|与|及|以及", raw)
    stop = {
        "这一页",
        "交代",
        "研究",
        "gap",
        "不是",
        "而是",
        "要让",
        "真正",
        "保持",
        "说明",
        "介绍",
        "通过",
        "用户",
        "内容",
        "vs",
        "rq",
    }
    terms: list[str] = []
    if _prefers_english_spec(spec):
        for point in display_points:
            words = [
                word for word in re.findall(r"[A-Za-z][A-Za-z0-9'_-]*|\d+(?:\.\d+)?%?", point)
                if word.lower() not in {"the", "and", "with", "into", "from", "that", "this", "can", "may", "need", "needs"}
            ]
            if not words:
                continue
            phrase = " ".join(words[:2])
            if phrase and phrase not in terms and not _is_layout_instruction_text(phrase):
                terms.append(phrase)
            if len(terms) >= limit:
                return terms[:limit]
    for piece in pieces:
        value = piece.strip(" -•·")
        if (
            len(value) < 2
            or value.lower() in stop
            or re.fullmatch(r"rq\d+(?:-\d+)?", value, re.IGNORECASE)
            or re.fullmatch(r"模块[一二三四五六七八九十\d]+", value)
            or any(marker in value for marker in ["这一页", "交代", "讲解", "说明"])
            or _is_layout_instruction_text(value)
        ):
            continue
        value = re.sub(r"^(现有|当前|主要|核心)", "", value).strip() or value
        short = _shorten_for_cell(value, 12)
        if short not in terms:
            terms.append(short)
        if len(terms) >= limit:
            break
    if not terms and spec.title:
        title = _shorten_for_cell(spec.title, 14)
        if title not in terms and not re.search(r"rq\d+|模块[一二三四五六七八九十\d]+", title, re.IGNORECASE):
            terms.insert(0, title)
    return terms[:limit]


def _image_overlay_labels(spec: SlideSpec, *, limit: int = 4) -> list[str]:
    if _prefers_english_spec(spec):
        raw = f"{spec.title or ''} {spec.message or ''}".lower()
        if any(key in raw for key in ["trade-off", "tradeoff", "cost", "risk", "balance"]):
            labels = ["Coverage", "Precision", "Cost", "Trust"]
        elif any(key in raw for key in ["pipeline", "workflow", "process", "stage"]):
            labels = ["Input", "Rank", "Rerank", "Feedback"]
        elif any(key in raw for key in ["evaluation", "metric", "offline", "online"]):
            labels = ["Metric", "Signal", "Result", "Action"]
        else:
            labels = ["Context", "Signal", "Model", "Action"]
        return labels[:limit]
    labels: list[str] = []
    for point in _slide_display_points(spec, limit=limit, include_title=False):
        for chunk in re.split(r"[，,；;。:：/、]|和|与|→|->", point):
            label = _shorten_for_cell(chunk, 12)
            if 2 <= len(label) <= 14 and label not in labels and not _is_layout_instruction_text(label):
                labels.append(label)
            if len(labels) >= limit:
                return labels
    for term in _template_concept_terms(spec, limit=limit + 2):
        label = _shorten_for_cell(term, 12)
        if 2 <= len(label) <= 14 and label not in labels and not _is_layout_instruction_text(label):
            labels.append(label)
        if len(labels) >= limit:
            break
    return labels[:limit]


def _add_image_overlay_labels(
    slide,
    style: PPTStyleDecision,
    spec: SlideSpec,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    max_labels: int = 4,
) -> None:
    labels = _image_overlay_labels(spec, limit=max_labels)
    if not labels:
        return
    p = style.palette
    ex = _palette_extras(style)
    ref = _academic_ref_colors(style)
    colors = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
    chip_h = 0.42
    chip_gap = 0.1
    chip_w = min(1.95, max(1.18, (w - 0.32 - chip_gap * (len(labels) - 1)) / max(len(labels), 1)))
    start_x = x + 0.16
    chip_y = y + h - chip_h - 0.16
    for i, label in enumerate(labels):
        cx = start_x + i * (chip_w + chip_gap)
        if cx + chip_w > x + w - 0.08:
            break
        _add_rect(slide, cx, chip_y, chip_w, chip_h, "FFFFFF", line=ref["line"], radius=False)
        _add_rect(slide, cx, chip_y, 0.06, chip_h, colors[i % len(colors)])
        _add_autofit_text(
            slide,
            label,
            cx + 0.1,
            chip_y + 0.06,
            chip_w - 0.16,
            chip_h - 0.1,
            max_size=10,
            min_size=8,
            color=p["ink"],
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.01,
        )


def _is_generated_slide_image(path: Path | None) -> bool:
    return bool(path and "imagegen" in {part.lower() for part in path.parts})


def _load_layout_registry(root: Path) -> dict[str, Any]:
    path = root / PPT_LAYOUT_REGISTRY
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _layout_family_for_style(style: PPTStyleDecision) -> str:
    if style.style_id == "general":
        return "general"
    if style.style_id == "major_project":
        return "major_project"
    return "academic_report"


def _allowed_layout_ids(root: Path, family: str) -> set[str]:
    ids = _family_layout_ids(root, family)
    # An explicit layout_id may intentionally cross style families because both
    # school templates contain the complete page schema.
    ids.update(SLIDE_LIBRARY_LAYOUT_IDS)
    return ids


def _family_layout_ids(root: Path, family: str) -> set[str]:
    registry = _load_layout_registry(root)
    families = registry.get("families") if isinstance(registry.get("families"), dict) else {}
    family_data = families.get(family) if isinstance(families, dict) else None
    layouts = family_data.get("layouts") if isinstance(family_data, dict) else []
    ids = {str(item.get("id", "")).strip() for item in layouts if isinstance(item, dict)}
    ids.discard("")
    return ids


LAYOUT_ID_ALIASES = {
    "title_research": "motivation_compare",
    "problem_contrast": "motivation_compare",
    "challenge_flow": "challenge_map",
    "task_benchmark": "benchmark_metrics",
    "architecture_large": "method_pipeline",
    "method_diagram": "method_pipeline",
    "loop_process": "method_loop",
    "generation_pipeline": "method_pipeline",
    "results_bars": "results_bars",
    "leaderboard_table": "leaderboard_table",
    "evidence_grid": "evidence_grid",
    "conclusion_discussion": "summary_takeaways",
    "key_claim": "challenge_map",
    "claim": "challenge_map",
    "comparison_split": "ablation_matrix",
    "two_column": "ablation_matrix",
    "summary_cards": "summary_takeaways",
    "summary": "summary_takeaways",
    "cover_paper_title": "motivation_compare",
    "problem_contrast_diagram": "motivation_compare",
    "two_lane_challenge_map": "challenge_map",
    "architecture_full_width": "method_pipeline",
    "method_flow_equation": "method_pipeline",
    "iterative_loop_diagram": "method_loop",
    "generation_control_pipeline": "method_pipeline",
    "dataset_metric_matrix": "benchmark_metrics",
    "result_bar_table_combo": "results_bars",
    "leaderboard_focus": "leaderboard_table",
    "ablation_sensitivity_dashboard": "evidence_grid",
    "case_limit_nextwork": "case_gallery",
    "academic_cover_visual": "motivation_compare",
    "compare_diagram_2col": "motivation_compare",
    "problem_challenge_map": "challenge_map",
    "task_benchmark_formula": "method_pipeline",
    "full_figure_architecture": "case_gallery",
    "editable_hypergraph_flow": "method_pipeline",
    "editable_loop_diagram": "method_loop",
    "dual_module_pipeline": "method_pipeline",
    "dataset_table_metric_grid": "benchmark_metrics",
    "result_bars_with_table": "results_bars",
    "leaderboard_heatmap": "leaderboard_table",
    "human_eval_triptych": "evidence_grid",
    "dual_line_cost_chart": "results_bars",
    "case_gallery_before_after": "case_gallery",
    "risk_takeaway_matrix": "summary_takeaways",
    "target_map": "project_target_map",
    "route_map": "technical_route",
    "evaluation_system": "evaluation_dashboard",
    "risk_table": "risk_action_table",
    "roadmap": "milestone_roadmap",
}


def _normalize_layout_id(value: str, allowed: set[str]) -> str | None:
    candidate = (value or "").strip()
    if not candidate:
        return None
    if candidate in allowed:
        return candidate
    alias = LAYOUT_ID_ALIASES.get(candidate)
    if alias and alias in allowed:
        return alias
    return None


def _default_layout_sequence(root: Path, family: str) -> list[str]:
    registry = _load_layout_registry(root)
    families = registry.get("families") if isinstance(registry.get("families"), dict) else {}
    family_data = families.get(family) if isinstance(families, dict) else None
    seq = family_data.get("default_sequence") if isinstance(family_data, dict) else []
    return [str(item).strip() for item in seq if str(item).strip()]


def _explicit_layout_id(spec: SlideSpec, allowed: set[str]) -> str | None:
    raw = "\n".join([spec.visual or "", spec.message or ""])
    patterns = [
        r"\blayout_id\s*[:：=]\s*([a-zA-Z0-9_-]+)",
        r"\blayout\s*[:：=]\s*([a-zA-Z0-9_-]+)",
        r"版式模板\s*[:：=]\s*([a-zA-Z0-9_-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, raw, re.IGNORECASE)
        if match:
            normalized = _normalize_layout_id(match.group(1), allowed)
            if normalized:
                return normalized
    return None


def _clean_layout_markers(text: str) -> str:
    return re.sub(r"(?:^|\n)\s*(?:layout_id|layout|版式模板)\s*[:：=]\s*[a-zA-Z0-9_-]+\s*", "\n", text or "", flags=re.IGNORECASE).strip()


def _select_template_layout_id(
    *,
    root: Path,
    spec: SlideSpec,
    style: PPTStyleDecision,
    index: int,
    total: int,
) -> str:
    family = _layout_family_for_style(style)
    explicit_allowed = _allowed_layout_ids(root, family)
    explicit = _explicit_layout_id(spec, explicit_allowed)
    if explicit:
        return explicit
    allowed = _family_layout_ids(root, family) or explicit_allowed
    raw = f"{spec.title or ''} {spec.message or ''} {spec.visual or ''}"
    numbery = sum(1 for point in _template_rich_bullets(spec, limit=5) if _extract_highlight(point))
    if index == total:
        return "summary_takeaways"
    if re.search(r"(TailorBench|数据集|基准|评测体系|evaluation|dashboard|benchmark)", raw, re.IGNORECASE):
        return "benchmark_metrics" if "benchmark_metrics" in allowed else "evaluation_dashboard"
    if re.search(r"(结果|实验|指标|Recall|NDCG|CTR|CVR)", raw, re.IGNORECASE) and numbery >= 2:
        return "result_big_numbers"
    if re.search(r"(背景|动机|为什么|痛点|必要|价值|供给|需求)", raw, re.IGNORECASE):
        return "motivation_compare" if "motivation_compare" in allowed else "project_target_map"
    if re.search(r"(反馈|反思|迭代|优化|闭环|gradient|reflection|loop|rank)", raw, re.IGNORECASE):
        return "method_loop" if "method_loop" in allowed else "technical_route"
    if re.search(r"(流程|workflow|步骤|链路|路线|架构|框架|系统|模块|pipeline|route)", raw, re.IGNORECASE):
        return "method_pipeline" if "method_pipeline" in allowed else "technical_route"
    if re.search(r"(全幅|满版|单张大图|视频封面|演示画面|full[- ]?bleed|media showcase)", raw, re.IGNORECASE):
        return "media_showcase"
    if re.search(r"(案例|demo|截图|画廊|样例|case|gallery|使用附件原图|source figure|论文图|生成插图)", raw, re.IGNORECASE):
        return "case_gallery" if "case_gallery" in allowed else "motivation_compare"
    if re.search(r"(消融|ablation|对比|挑战|限制|问题|风险|risk|短板|局限)", raw, re.IGNORECASE):
        return "ablation_matrix" if "ablation_matrix" in allowed else "risk_action_table"
    sequence = _default_layout_sequence(root, family)
    if sequence:
        body_index = max(0, index - 2)
        return sequence[body_index % len(sequence)]
    return "method_pipeline"


def _slide_library_content_payload(spec: SlideSpec) -> dict[str, Any]:
    payload = dict(spec.content_spec or {})
    payload.setdefault("message", spec.message or "")
    payload.setdefault("visual", _clean_layout_markers(spec.visual or ""))
    payload.setdefault("speaker_note", spec.speaker_note or "")
    if not isinstance(payload.get("points"), list) or not payload.get("points"):
        payload["points"] = _slide_display_points(spec, limit=10, include_title=False)
    if not isinstance(payload.get("numbers"), list) or not payload.get("numbers"):
        payload["numbers"] = _prominent_numeric_highlights(
            f"{spec.title or ''} {spec.message or ''} {spec.visual or ''}",
            limit=8,
        )
    if not payload.get("proof_object"):
        payload["proof_object"] = "；".join(_evidence_visible_points(spec, limit=3))
    return payload



__all__ = [
    "_palette_extras",
    "_is_scut_academic_chrome",
    "_is_hitsz_academic_chrome",
    "_academic_ref_colors",
    "_academic_ref_fill",
    "_add_academic_ref_panel",
    "_academic_visible_text",
    "_academic_sparse_points_boost",
    "_academic_sparse_spec_boost",
    "_add_academic_plain_list",
    "_add_academic_plain_note",
    "_template_profile",
    "_style_for_template",
    "_template_footer_label",
    "_add_footer",
    "_add_slide_background",
    "_add_title_system",
    "_split_points",
    "_template_source_slide_index",
    "_template_chrome_pdf_cache_path",
    "_convert_template_to_pdf",
    "_render_template_chrome_backgrounds",
    "_template_source_placeholder_boxes",
    "_box_or_default",
    "_blank_layout",
    "_template_layout_index",
    "_template_slide_layout",
    "_uses_template_layouts",
    "_slide_role",
    "_placeholder_by_idx",
    "_placeholder_type_name",
    "_first_text_placeholder",
    "_shape_box_inches",
    "_set_text_frame",
    "_set_title_placeholder",
    "_set_footer_placeholders",
    "_add_template_chrome_background",
    "_render_template_chrome_cover",
    "_template_chrome_boxes",
    "_add_template_chrome_title",
    "_template_slide_bullets",
    "_template_rich_bullets",
    "_template_concept_terms",
    "_image_overlay_labels",
    "_add_image_overlay_labels",
    "_is_generated_slide_image",
    "_load_layout_registry",
    "_layout_family_for_style",
    "_allowed_layout_ids",
    "_family_layout_ids",
    "_normalize_layout_id",
    "_default_layout_sequence",
    "_explicit_layout_id",
    "_clean_layout_markers",
    "_select_template_layout_id",
    "_slide_library_content_payload",
]
