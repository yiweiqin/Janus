from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

from PIL import Image, ImageFilter, ImageStat
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from ppt_pipeline.content_payloads import *
from ppt_pipeline.slide_library_core import *

_IMAGE_SLOTS_BY_LAYOUT = {
    "basic_content": ["basic-image"],
    "basic_content_mirror": ["basic-image"],
    "evidence_grid": [f"evidence-img-{row}-{col}" for row in range(2) for col in range(3)],
    "case_gallery": ["case-main", *[f"case-thumb-{index}" for index in range(4)]],
    "media_showcase": ["media-showcase-main"],
}



def _rows(payload: dict[str, Any]) -> list[list[str]]:
    table = payload.get("table") if isinstance(payload.get("table"), dict) else {}
    raw = table.get("rows") if isinstance(table.get("rows"), list) else payload.get("rows")
    result: list[list[str]] = []
    if isinstance(raw, list):
        for row in raw:
            if isinstance(row, dict):
                result.append([_display(value) for value in row.values()])
            elif isinstance(row, list):
                result.append([_display(value) for value in row])
            else:
                result.append([_display(row)])
    return result
def _repair_slide_safe_area(slide: Any, layout_id: str, index: int, warnings: list[str]) -> None:
    prefixes = (
        "basic-", "motivation-", "challenge-", "pipeline-", "loop-", "bench-", "big-", "bar", "bars-",
        "leader-", "ablation-", "evidence-", "case-", "summary-", "target-", "domain-", "route-", "wp-",
        "eval-", "risk-", "roadmap-", "media-", "janus-",
    )
    left_limit = Inches(0.35)
    right_limit = Inches(12.98)
    top_limit = Inches(1.05)
    bottom_limit = Inches(6.72)
    repaired: list[str] = []
    for shape in slide.shapes:
        name = str(getattr(shape, "name", "") or "")
        if not name.startswith(prefixes):
            continue
        left = int(shape.left)
        top = int(shape.top)
        width = max(1, int(shape.width))
        height = max(1, int(shape.height))
        new_left = max(left_limit, left)
        new_top = max(top_limit, top)
        new_width = min(width, max(1, right_limit - new_left))
        new_height = min(height, max(1, bottom_limit - new_top))
        if (new_left, new_top, new_width, new_height) == (left, top, width, height):
            continue
        shape.left = new_left
        shape.top = new_top
        shape.width = new_width
        shape.height = new_height
        repaired.append(name)
    if repaired:
        warnings.append(
            f"第 {index} 页 {layout_id} 有 {len(repaired)} 个元素超出正文安全区，已自动收回：{'、'.join(repaired[:3])}。"
        )




def _effective_layout_id(
    requested: str,
    payload: dict[str, Any],
    image_paths: list[Path],
    library: dict[str, int],
    warnings: list[str],
) -> str:
    layout_id = requested if requested in library else "basic_content"
    raw = " ".join(
        part for part in [
            _display(payload.get("title")),
            _display(payload.get("message")),
            _display(payload.get("scope")),
        ] if part
    )
    if (
        layout_id == "domain_object_map"
        and "evidence_grid" in library
        and re.search(
            r"(?:where.+(?:not|does not)|strong[- ]?fit|weak[- ]?fit|product patterns?|application patterns?|interaction modes?|适用|不适用|模式|类型)",
            raw,
            re.IGNORECASE,
        )
    ):
        warnings.append("domain_object_map 内容更接近分类/对比，已智能改用 evidence_grid 卡片版式。")
        return "evidence_grid"
    if len(image_paths) == 1 and layout_id in {"evidence_grid", "case_gallery"} and "basic_content" in library:
        warnings.append(f"{layout_id} 当前只有一张主图，已智能改用 basic_content 单图版式，避免空图片卡片。")
        return "basic_content"
    if image_paths and layout_id not in _IMAGE_SLOTS_BY_LAYOUT:
        fallback = "evidence_grid" if len(image_paths) > 1 or any(
            isinstance(payload.get(key), list) and bool(payload.get(key))
            for key in ("cards", "captions")
        ) else "basic_content"
        if fallback in library:
            warnings.append(f"{layout_id} 没有图片槽位，已智能改用 {fallback} 以保留生成/提取图片。")
            return fallback
    if layout_id == "media_showcase" and not image_paths:
        if isinstance(payload.get("steps"), list) and payload.get("steps"):
            fallback = "method_pipeline"
        elif isinstance(payload.get("cards"), list) and payload.get("cards"):
            fallback = "evidence_grid"
        else:
            fallback = "basic_content"
        if fallback in library:
            warnings.append(f"media_showcase 没有可用图片，已智能改用 {fallback}。")
            return fallback
    if layout_id in {"evidence_grid", "case_gallery"} and not image_paths:
        has_text_cards = any(
            isinstance(payload.get(key), list) and bool(payload.get(key))
            for key in ("cards", "captions", "points")
        )
        if not has_text_cards and "basic_content" in library:
            warnings.append(f"{layout_id} 没有图片或卡片内容，已智能改用 basic_content。")
            return "basic_content"
    return layout_id


def _layout_family(layout_id: str) -> str:
    if layout_id in {"basic_content", "basic_content_mirror", "motivation_compare", "challenge_map", "summary_takeaways"}:
        return "narrative"
    if layout_id in {"evidence_grid", "case_gallery", "media_showcase"}:
        return "visual"
    if layout_id in {"method_pipeline", "method_loop", "technical_route", "milestone_roadmap"}:
        return "process"
    if layout_id in {
        "benchmark_metrics", "result_big_numbers", "results_bars", "leaderboard_table", "ablation_matrix",
        "workpackage_matrix", "evaluation_dashboard", "risk_action_table",
    }:
        return "data"
    return "structure"


def _layout_density(layout_id: str, payload: dict[str, Any], image_count: int) -> str:
    if image_count or layout_id in {"evidence_grid", "case_gallery", "media_showcase"}:
        return "visual"
    if layout_id in {"leaderboard_table", "ablation_matrix", "workpackage_matrix", "risk_action_table", "evaluation_dashboard"}:
        return "dense"
    if layout_id in {"result_big_numbers", "summary_takeaways", "motivation_compare", "milestone_roadmap"}:
        return "sparse"
    raw = " ".join(
        _display(payload.get(key))
        for key in ("message", "points", "cards", "rows", "nodes", "stages", "kpis")
    )
    units = len(re.findall(r"[\u3400-\u9fff]", raw)) + len(re.findall(r"[A-Za-z0-9]+", raw)) * 1.4
    return "dense" if units >= 170 else "medium"


def _layout_candidates(
    role: str,
    payload: dict[str, Any],
    image_paths: list[Path],
    library: dict[str, int],
) -> list[str]:
    candidates = [role]
    if role == "basic_content":
        if image_paths:
            candidates.append("basic_content_mirror")
        else:
            item_count = max(
                len(_structured_list(payload, "points")),
                len(_structured_list(payload, "cards")),
                len(_payload_points(payload, 4)),
            )
            if 1 <= item_count <= 3:
                candidates.append("summary_takeaways")
    elif role == "method_pipeline":
        candidates.append("technical_route")
    elif role == "technical_route":
        candidates.append("method_pipeline")
    elif role == "challenge_map" and _structured_list(payload, "risks"):
        candidates.append("risk_action_table")
    elif role == "evidence_grid":
        candidates.append("case_gallery")
    elif role == "case_gallery" and len(image_paths) >= 2:
        candidates.append("evidence_grid")
    elif role == "summary_takeaways":
        has_closing = bool(payload.get("conclusion") or payload.get("next_step") or payload.get("recommendation"))
        has_cards = bool(_structured_list(payload, "cards") or _structured_list(payload, "findings") or _payload_points(payload, 4))
        if has_cards and not has_closing:
            candidates.append("evidence_grid")
    elif role == "benchmark_metrics" and _structured_list(payload, "kpis"):
        if not _rows(payload):
            candidates.append("result_big_numbers")
        if _structured_list(payload, "status_rows"):
            candidates.append("evaluation_dashboard")
    elif role == "result_big_numbers" and _structured_list(payload, "kpis"):
        candidates.append("benchmark_metrics")
    available: list[str] = []
    for candidate in candidates:
        source = "basic_content" if candidate == "basic_content_mirror" else candidate
        if source in library and candidate not in available:
            available.append(candidate)
    return available or [role]


def _rhythm_layout_id(
    role: str,
    payload: dict[str, Any],
    image_paths: list[Path],
    library: dict[str, int],
    role_counts: dict[str, int],
    recent_roles: list[str],
    warnings: list[str],
) -> str:
    candidates = _layout_candidates(role, payload, image_paths, library)
    last_role = recent_roles[-1] if recent_roles else ""
    last_family = _layout_family(last_role) if last_role else ""
    last_density = _layout_density(last_role, {}, 0) if last_role else ""
    recent_families = [_layout_family(item) for item in recent_roles[-2:]]

    def score(candidate: str) -> float:
        value = candidates.index(candidate) * 1.25
        value += role_counts.get(candidate, 0) * 1.8
        if candidate == last_role:
            value += 9.0
        family = _layout_family(candidate)
        density = _layout_density(candidate, payload, len(image_paths))
        if last_family and family == last_family:
            value += 2.4
        if len(recent_families) == 2 and recent_families[0] == recent_families[1] == family:
            value += 5.0
        if last_density == density == "dense":
            value += 7.0
        if last_density == density == "visual":
            value += 2.0
        if candidate == "basic_content_mirror" and last_role in {"basic_content", "basic_content_mirror"}:
            value -= 3.5
        if candidate == "summary_takeaways" and last_density == "dense":
            value -= 2.0
        return value

    selected = min(candidates, key=score)
    if selected != role:
        warnings.append(
            f"为改善整套 PPT 页面节奏并避免重复使用 {role}，已改用 {selected} 替补版式。"
        )
    return selected


def _normalize_layout_payload(layout_id: str, payload: dict[str, Any], spec: Any) -> dict[str, Any]:
    normalized = dict(payload or {})
    normalized.setdefault("title", str(getattr(spec, "title", "") or ""))
    normalized.setdefault("message", str(normalized.get("message") or getattr(spec, "message", "") or ""))

    if layout_id == "method_loop" and not isinstance(normalized.get("steps"), list):
        if isinstance(normalized.get("stages"), list):
            normalized["steps"] = normalized["stages"]

    if layout_id == "technical_route" and isinstance(normalized.get("steps"), list):
        normalized.setdefault("stages", normalized["steps"])
        normalized.setdefault("dependency", _display(normalized.get("output")))

    if layout_id == "risk_action_table":
        table_payload = normalized.get("table") if isinstance(normalized.get("table"), dict) else {}
        raw_records = normalized.get("risks")
        if not isinstance(raw_records, list):
            raw_records = table_payload.get("rows") if isinstance(table_payload.get("rows"), list) else normalized.get("rows")
        if isinstance(raw_records, list) and any(isinstance(item, dict) for item in raw_records):
            records: list[dict[str, str]] = []
            for item in raw_records:
                if not isinstance(item, dict):
                    records.append({"risk": _display(item), "trigger": "", "impact": "", "action": "", "owner": "", "status": ""})
                    continue
                records.append({
                    "risk": _display(item.get("risk") or item.get("title") or item.get("label")),
                    "trigger": _display(item.get("trigger") or item.get("condition") or item.get("cause")),
                    "impact": _display(item.get("impact") or item.get("consequence")),
                    "action": _display(item.get("mitigation") or item.get("action") or item.get("response")),
                    "owner": _display(item.get("owner") or item.get("responsible")),
                    "status": _display(item.get("status") or item.get("state")),
                })
            normalized["_risk_records"] = records
            normalized.pop("risks", None)
            normalized.pop("rows", None)

    if layout_id in {"evidence_grid", "case_gallery"}:
        cards = normalized.get("cards")
        if not isinstance(cards, list) and isinstance(normalized.get("nodes"), list):
            cards = []
            for item in normalized["nodes"]:
                title, detail = _item_parts(item)
                if title or detail:
                    cards.append({"title": title or detail, "points": [detail] if title and detail else []})
            normalized["cards"] = cards
        if isinstance(cards, list) and cards and not isinstance(normalized.get("captions"), list):
            normalized["captions"] = [_card_body(item) for item in cards]

    if layout_id == "benchmark_metrics":
        raw_kpis = normalized.get("kpis")
        if isinstance(raw_kpis, list):
            table_payload = payload.get("table") if isinstance(payload.get("table"), dict) else {}
            has_explicit_rows = isinstance(payload.get("rows"), list) or isinstance(table_payload.get("rows"), list)
            if not has_explicit_rows:
                normalized["_benchmark_rows_derived"] = True
            kpis: list[dict[str, Any]] = []
            rows: list[list[str]] = []
            for item in raw_kpis:
                if not isinstance(item, dict):
                    kpis.append({"value": _display(item), "label": "", "note": ""})
                    continue
                group = _display(item.get("group") or item.get("title"))
                metrics = item.get("metrics") if isinstance(item.get("metrics"), list) else []
                metric_values = [_display(metric) for metric in metrics if _display(metric)]
                metric_text = " / ".join(metric_values)
                # The detailed metric list is already preserved in the table.
                # Keep the KPI card to one representative metric so its label
                # remains presentation-sized instead of shrinking to footnote text.
                metric_label = metric_values[0] if metric_values else ""
                kpis.append({
                    "value": _display(item.get("value")) or group,
                    "label": _display(item.get("label")) or metric_label,
                    "note": _display(item.get("note") or item.get("description")),
                })
                if group or metric_text:
                    rows.append([group, metric_text, "", "", ""])
            normalized["kpis"] = kpis
            if rows and not isinstance(normalized.get("rows"), list):
                normalized["rows"] = rows
            if len(kpis) > 3:
                extra = "\n".join(
                    " — ".join(part for part in [_display(item.get("value")), _display(item.get("label"))] if part)
                    for item in kpis[3:]
                )
                protocol = _display(normalized.get("protocol"))
                normalized["protocol"] = "\n".join(part for part in [protocol, extra] if part)

    return normalized




def _card_body(item: Any) -> str:
    title, detail = _item_parts(item)
    if title and detail:
        return f"{title}\n{detail}"
    return title or detail


def _card_text(item: Any) -> str:
    title, detail = _item_parts(item)
    if not title:
        return detail
    if not detail:
        return title
    details = [line.strip() for line in detail.splitlines() if line.strip()]
    return "\n".join([title, *(f"• {line}" for line in details)])


def _remove_occurrences_after(slide, name: str, keep: int) -> None:
    for shape in list(_shapes_named(slide, name))[keep:]:
        _remove_shape(shape)


def _remove_slot(slide, name: str) -> None:
    _remove_named(slide, name)
    _remove_crosses(slide, name)


def _set_geometry(shape: Any, x: float, y: float, width: float, height: float) -> None:
    if shape is None:
        return
    shape.left = Inches(x)
    shape.top = Inches(y)
    shape.width = Inches(width)
    shape.height = Inches(height)


def _shape_center_and_half_extents(shape: Any) -> tuple[float, float, float, float]:
    return (
        float(shape.left + shape.width / 2),
        float(shape.top + shape.height / 2),
        float(shape.width / 2),
        float(shape.height / 2),
    )


def _rectangle_edge_point(source: Any, target: Any) -> tuple[float, float]:
    """Return where the source-to-target center ray meets source's border."""
    source_x, source_y, half_width, half_height = _shape_center_and_half_extents(source)
    target_x, target_y, _, _ = _shape_center_and_half_extents(target)
    delta_x = target_x - source_x
    delta_y = target_y - source_y
    if not delta_x and not delta_y:
        return source_x, source_y
    scale_x = half_width / abs(delta_x) if delta_x else math.inf
    scale_y = half_height / abs(delta_y) if delta_y else math.inf
    scale = min(scale_x, scale_y)
    return source_x + delta_x * scale, source_y + delta_y * scale


def _set_line_endpoints(line: Any, start: tuple[float, float], end: tuple[float, float]) -> None:
    if line is None:
        return
    start_x, start_y = start
    end_x, end_y = end
    line.left = int(min(start_x, end_x))
    line.top = int(min(start_y, end_y))
    line.width = max(1, int(abs(end_x - start_x)))
    line.height = max(1, int(abs(end_y - start_y)))
    transform = line._element.spPr.xfrm
    for attribute, enabled in (("flipH", end_x < start_x), ("flipV", end_y < start_y)):
        if enabled:
            transform.set(attribute, "1")
        else:
            transform.attrib.pop(attribute, None)


def _clip_lines_to_shape_borders(lines: list[Any], sources: list[Any], target: Any) -> None:
    if target is None:
        return
    for line, source in zip(lines, sources):
        _set_line_endpoints(
            line,
            _rectangle_edge_point(source, target),
            _rectangle_edge_point(target, source),
        )


def _style_text_card(shape: Any) -> None:
    if shape is None:
        return
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(241, 247, 250)
        shape.line.color.rgb = RGBColor(176, 199, 214)
        shape.line.dash_style = MSO_LINE_DASH_STYLE.SOLID
        shape.line.width = Pt(1.0)
        _apply_soft_shadow(shape)
    except Exception:
        pass


def _style_pipeline_endpoint(shape: Any) -> None:
    if shape is None:
        return
    try:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(247, 250, 252)
        shape.line.color.rgb = RGBColor(176, 199, 214)
        shape.line.dash_style = MSO_LINE_DASH_STYLE.SOLID
        shape.line.width = Pt(0.8)
    except Exception:
        pass


SEMANTIC_DARK = RGBColor(15, 58, 86)
SEMANTIC_PRIMARY = RGBColor(0, 138, 154)
SEMANTIC_PRIMARY_LIGHT = RGBColor(239, 248, 249)
SEMANTIC_MUTED = RGBColor(143, 176, 195)
SEMANTIC_MUTED_LIGHT = RGBColor(247, 250, 252)
SEMANTIC_WARM = RGBColor(224, 135, 55)
SEMANTIC_WARM_LIGHT = RGBColor(255, 247, 228)
SEMANTIC_SUCCESS_LIGHT = RGBColor(234, 247, 239)
SEMANTIC_PENDING_LIGHT = RGBColor(242, 245, 247)


def _set_shape_colors(shape: Any, *, fill: RGBColor | None = None, line: RGBColor | None = None) -> None:
    if shape is None:
        return
    try:
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line
            shape.line.width = Pt(0.8)
    except Exception:
        pass


def _set_shape_text_color(shape: Any, color: RGBColor) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            try:
                run.font.color.rgb = color
            except Exception:
                pass


def _truthy_semantic_flag(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "primary", "highlight", "best", "重点", "主要", "是"}


def _semantic_primary_index(items: list[dict[str, Any]], payload: dict[str, Any] | None = None) -> int:
    if not items:
        return 0
    payload = payload or {}
    requested = payload.get("primary_index")
    if isinstance(requested, int) and 0 <= requested < len(items):
        return requested
    requested_text = _display(payload.get("highlight") or payload.get("primary") or payload.get("focus")).lower()
    if requested_text:
        for index, item in enumerate(items):
            if requested_text in " ".join(_display(item.get(key)) for key in ("label", "title", "value", "note")).lower():
                return index
    for index, item in enumerate(items):
        if any(_truthy_semantic_flag(item.get(key)) for key in ("primary", "highlight", "best", "is_primary")):
            return index
    primary_keywords = (
        "ours", "our method", "proposed", "current", "primary", "best", "recommended", "core", "key",
        "本方案", "我们", "当前", "主要", "核心", "关键", "最佳", "推荐",
    )
    for index, item in enumerate(items):
        text = " ".join(_display(item.get(key)) for key in ("label", "title", "note", "status", "state")).lower()
        if any(keyword in text for keyword in primary_keywords):
            return index
    return 0


def _numeric_value(value: Any) -> float | None:
    match = re.search(r"[+\-]?\d+(?:\.\d+)?", _display(value).replace(",", ""))
    return float(match.group(0)) if match else None


def _style_first_paragraph(shape: Any, *, color: RGBColor | None = None, bold: bool = True) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False) or not shape.text_frame.paragraphs:
        return
    for run in shape.text_frame.paragraphs[0].runs:
        run.font.bold = bold
        if color is not None:
            try:
                run.font.color.rgb = color
            except Exception:
                pass


_DEPTH_CARD_PREFIXES = (
    "basic-text-card",
    "challenge-core",
    "loop-center",
    "summary-next",
    "target-root",
)


def _apply_soft_shadow(shape: Any, *, opacity: int = 8, blur_pt: float = 4.0, distance_pt: float = 1.5) -> None:
    if shape is None:
        return
    try:
        shape_properties = shape._element.spPr
        for tag in ("a:effectLst", "a:effectDag"):
            existing = shape_properties.find(qn(tag))
            if existing is not None:
                shape_properties.remove(existing)
        effects = OxmlElement("a:effectLst")
        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", str(int(blur_pt * 12700)))
        shadow.set("dist", str(int(distance_pt * 12700)))
        shadow.set("dir", "2700000")
        shadow.set("algn", "ctr")
        shadow.set("rotWithShape", "0")
        color = OxmlElement("a:srgbClr")
        color.set("val", "1E293B")
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(max(0, min(100, opacity)) * 1000))
        color.append(alpha)
        shadow.append(color)
        effects.append(shadow)
        extension = shape_properties.find(qn("a:extLst"))
        if extension is None:
            shape_properties.append(effects)
        else:
            shape_properties.insert(shape_properties.index(extension), effects)
    except Exception:
        pass


def _apply_layout_depth(slide: Any) -> None:
    for shape in _walk_shapes(slide.shapes):
        name = str(getattr(shape, "name", "") or "")
        if not name.startswith(_DEPTH_CARD_PREFIXES):
            continue
        width = float(getattr(shape, "width", 0) or 0) / 914400
        height = float(getattr(shape, "height", 0) or 0) / 914400
        if width < 1.4 or height < 0.62:
            continue
        _apply_soft_shadow(shape)


def _set_text_spacing(shape: Any, *, after: float, line_spacing: float = 1.1, middle: bool = False) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    if middle:
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in shape.text_frame.paragraphs:
        paragraph.space_after = Pt(after)
        paragraph.line_spacing = line_spacing
    _refit_shape_text(shape)


def _set_slide_notes(slide: Any, text: str) -> None:
    if not str(text or "").strip():
        return
    try:
        frame = slide.notes_slide.notes_text_frame
        frame.clear()
        frame.paragraphs[0].text = str(text).strip()
    except Exception:
        pass


def _image_information_score(path: Path) -> tuple[float, float]:
    try:
        with Image.open(path) as source:
            rgb = source.convert("RGB")
            aspect = rgb.width / max(rgb.height, 1)
            sample = rgb.resize((240, 135))
            gray = sample.convert("L")
            edges = gray.filter(ImageFilter.FIND_EDGES)
            edge_hist = edges.histogram()
            gray_hist = gray.histogram()
            total = max(1, sample.width * sample.height)
            edge_ratio = sum(edge_hist[29:]) / total
            nonwhite_ratio = sum(gray_hist[:245]) / total
            entropy = min(1.0, gray.entropy() / 8.0)
            score = min(1.0, edge_ratio * 2.0 + nonwhite_ratio * 0.35 + entropy * 0.25)
            return score, aspect
    except Exception:
        return 0.0, 1.0


def _image_visual_mode(path: Path) -> str:
    """Choose contain for charts/screenshots and crop-to-fill for photos."""
    try:
        with Image.open(path) as source:
            sample = source.convert("RGB")
            sample.thumbnail((320, 240))
            gray = sample.convert("L")
            hsv = sample.convert("HSV")
            total = max(1, sample.width * sample.height)
            gray_hist = gray.histogram()
            white_ratio = sum(gray_hist[242:]) / total
            edge_hist = gray.filter(ImageFilter.FIND_EDGES).histogram()
            edge_ratio = sum(edge_hist[32:]) / total
            saturation = ImageStat.Stat(hsv.getchannel("S")).mean[0] / 255.0
            aspect = source.width / max(source.height, 1)
        # White-canvas figures, UI screenshots, tables, and extreme aspect
        # ratios lose labels quickly when cropped. Preserve their full extent.
        if aspect >= 2.35 or aspect <= 0.58:
            return "contain"
        if white_ratio >= 0.34 and edge_ratio >= 0.025:
            return "contain"
        if saturation <= 0.13 and edge_ratio >= 0.08:
            return "contain"
    except Exception:
        return "contain"
    return "crop"


def _source_for_index(payload: dict[str, Any], index: int) -> str:
    sources = payload.get("sources")
    if isinstance(sources, list) and index < len(sources):
        return _display(sources[index])
    return ""


def _figure_caption(payload: dict[str, Any], item: Any, index: int) -> str:
    source = _source_for_index(payload, index)
    label = ""
    conclusion = ""
    if isinstance(item, dict):
        label = _display(item.get("number") or item.get("figure") or item.get("label"))
        conclusion = _display(
            item.get("conclusion")
            or item.get("caption")
            or item.get("title")
            or item.get("text")
            or item.get("description")
        )
        source = _display(item.get("source") or item.get("source_note")) or source
    else:
        conclusion = _display(item)
    already_numbered = bool(re.match(r"^(?:图|表|figure|fig\.?|table)\s*[A-Za-z0-9一二三四五六七八九十-]+", conclusion, re.IGNORECASE))
    if not label and not already_numbered:
        label = _localized(payload, f"图 {index + 1}", f"Figure {index + 1}")
    first_line = "  ".join(part for part in [label, conclusion] if part)
    if source:
        source_prefix = _localized(payload, "来源：", "Source: ")
        source_line = source if re.match(r"^(?:来源|source)\s*[：:]", source, re.IGNORECASE) else source_prefix + source
        return "\n".join(part for part in [first_line, source_line] if part)
    return first_line


def _style_figure_caption(shape: Any) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.text_frame.margin_left = Inches(0.04)
    shape.text_frame.margin_right = Inches(0.04)
    shape.text_frame.margin_top = Inches(0.02)
    shape.text_frame.margin_bottom = Inches(0.02)
    try:
        shape.fill.background()
        shape.line.fill.background()
    except Exception:
        pass
    for index, paragraph in enumerate(shape.text_frame.paragraphs):
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(1 if index else 2)
        paragraph.line_spacing = 1.0
        for run in paragraph.runs:
            if index == 0:
                run.font.bold = True
                run.font.size = Pt(11.5)
            else:
                run.font.bold = False
                run.font.size = Pt(9.5)
                try:
                    run.font.color.rgb = RGBColor(91, 105, 116)
                except Exception:
                    pass


def _basic_story(payload: dict[str, Any]) -> tuple[str, bool]:
    points = _payload_points(payload, 6)
    lead = _display(payload.get("conclusion") or payload.get("key_message") or payload.get("takeaway"))
    if lead:
        points = [point for point in points if point.strip() != lead.strip()]
        lines = [lead, *(f"• {point}" for point in points)]
        return "\n".join(lines), True
    if len(points) == 1:
        return points[0], False
    return "\n".join(f"• {point}" for point in points), False


def _basic_text_is_sparse(shape: Any, *, narrow: bool) -> bool:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return False
    text = re.sub(r"\s+", " ", _shape_text(shape)).strip()
    paragraphs = [paragraph for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]
    return len(paragraphs) <= 3 and len(text) <= (112 if narrow else 165)


def _configure_basic_text_density(shape: Any, payload: dict[str, Any], *, narrow: bool) -> None:
    if shape is None or not getattr(shape, "has_text_frame", False):
        return
    paragraphs = [paragraph for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()]
    sparse = _basic_text_is_sparse(shape, narrow=narrow)
    max_size = 18 if narrow and sparse else 15 if narrow else 24 if sparse else 19
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE if sparse else MSO_ANCHOR.TOP
    shape.text_frame.margin_left = Inches(0.28 if narrow else 0.42)
    shape.text_frame.margin_right = Inches(0.24 if narrow else 0.42)
    shape.text_frame.margin_top = Inches(0.24 if narrow else 0.38)
    shape.text_frame.margin_bottom = Inches(0.22 if narrow else 0.32)
    has_lead = bool(_display(payload.get("conclusion") or payload.get("key_message") or payload.get("takeaway")))
    for index, paragraph in enumerate(paragraphs):
        paragraph.space_after = Pt(16 if has_lead and index == 0 else 9 if narrow else 13)
        paragraph.line_spacing = 1.08 if narrow else 1.14
        if has_lead and index == 0:
            for run in paragraph.runs:
                run.font.bold = True
    _refit_shape_text(shape, max_size=max_size, min_size=13 if narrow else 15)


def _reflow_even_cards(shapes: list[Any], count: int, *, left: float, right: float, y: float, height: float, max_width: float = 5.0) -> None:
    if count <= 0:
        return
    gap = 0.36 if count > 1 else 0.0
    available = right - left
    width = min(max_width, (available - gap * (count - 1)) / count)
    total = width * count + gap * (count - 1)
    start = left + (available - total) / 2
    for index, shape in enumerate(shapes[:count]):
        _set_geometry(shape, start + index * (width + gap), y, width, height)


def _is_cjk_payload(payload: dict[str, Any]) -> bool:
    raw = " ".join(_display(payload.get(key)) for key in ("title", "message", "objective", "conclusion"))
    return bool(re.search(r"[\u3400-\u9fff]", raw))


def _localized(payload: dict[str, Any], zh: str, en: str) -> str:
    return zh if _is_cjk_payload(payload) else en


def _compact_pipeline_endpoint(payload: dict[str, Any], value: Any, *, output: bool) -> str:
    text = _display(value)
    if len(text) <= 48:
        return text
    if output:
        return _localized(payload, "受约束且可验证的输出", "Grounded, validated output")
    return _localized(payload, "用户上下文、证据与约束", "User context, evidence, and constraints")



__all__ = [
    "_append_deck_rhythm_warnings",
    "SEMANTIC_DARK",
    "SEMANTIC_PRIMARY",
    "SEMANTIC_PRIMARY_LIGHT",
    "SEMANTIC_MUTED",
    "SEMANTIC_MUTED_LIGHT",
    "SEMANTIC_WARM",
    "SEMANTIC_WARM_LIGHT",
    "SEMANTIC_SUCCESS_LIGHT",
    "SEMANTIC_PENDING_LIGHT",
    "_repair_slide_safe_area",
    "_effective_layout_id",
    "_layout_family",
    "_layout_density",
    "_layout_candidates",
    "_rhythm_layout_id",
    "_normalize_layout_payload",
    "_card_body",
    "_card_text",
    "_remove_occurrences_after",
    "_remove_slot",
    "_set_geometry",
    "_shape_center_and_half_extents",
    "_rectangle_edge_point",
    "_set_line_endpoints",
    "_clip_lines_to_shape_borders",
    "_style_text_card",
    "_style_pipeline_endpoint",
    "_set_shape_colors",
    "_set_shape_text_color",
    "_truthy_semantic_flag",
    "_semantic_primary_index",
    "_numeric_value",
    "_style_first_paragraph",
    "_apply_soft_shadow",
    "_apply_layout_depth",
    "_set_text_spacing",
    "_set_slide_notes",
    "_image_information_score",
    "_image_visual_mode",
    "_source_for_index",
    "_figure_caption",
    "_style_figure_caption",
    "_basic_story",
    "_basic_text_is_sparse",
    "_configure_basic_text_density",
    "_reflow_even_cards",
    "_is_cjk_payload",
    "_localized",
    "_compact_pipeline_endpoint",
]


def _append_deck_rhythm_warnings(roles: list[str], warnings: list[str]) -> None:
    families = [_layout_family(role) for role in roles]
    densities = [_layout_density(role, {}, 0) for role in roles]
    for index in range(1, len(roles)):
        if densities[index - 1] == densities[index] == "dense":
            warnings.append(
                f"第 {index + 1}、{index + 2} 页仍连续使用高密度页面；当前内容缺少安全的语义等价替补版式。"
            )
    for index in range(2, len(roles)):
        if families[index - 2] == families[index - 1] == families[index]:
            warnings.append(
                f"第 {index}–{index + 2} 页连续属于 {families[index]} 类型，建议下一轮内容规划中插入不同密度页面。"
            )
