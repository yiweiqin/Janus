from __future__ import annotations

import base64
import binascii
import json
import hashlib
import os
import re
import secrets
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import tomllib
import urllib.error
import urllib.request
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, replace
from datetime import datetime, timezone
from html import escape as xml_escape, unescape as html_unescape
from pathlib import Path
from typing import Any, Callable

WORKSPACE = Path(__file__).resolve().parent.parent

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

try:
    from PIL import Image
except Exception:
    Image = None

sys.path.insert(0, str(WORKSPACE / "venv" / "lib" / "python3.12" / "site-packages"))

from janus_lab import db as janus_db
from janus_lab.agent_defs import load_organization
from janus_lab.codex_runner import (
    CodexUnavailable,
    DEFAULT_CODEX_SANDBOX,
    load_auth_env,
    resolve_codex_binary,
    shared_auth_path,
    shared_config_path,
    write_repo_templates,
)
from janus_lab.evolution import (
    build_chat_prompt,
    ensure_agent_skill,
    ensure_hr_assets,
    ensure_memory,
    read_memory,
    read_skill,
)
from janus_lab.paths import hr_memory_file, legacy_db_path, memory_file, state_root, tmp_dir
from janus_lab.ppt_renderer import SlideSpec, parse_slide_specs as _base_parse_slide_specs
from slide_library_renderer import (
    KNOWN_LAYOUT_IDS as SLIDE_LIBRARY_LAYOUT_IDS,
    render_slide_library_presentation,
    validate_slide_library,
)
from ppt_pipeline.image_routing import (
    IMAGE_CAPABLE_LAYOUT_IDS,
    _auto_generated_image_eligible,
    _explicit_generated_image_requested,
    _explicit_source_visual_requested,
    _first_slide_image,
    _generated_image_brief,
    _map_source_visuals_to_slides,
    _merge_slide_image_maps,
    _needs_generated_image,
    _requested_source_figures,
    _requested_source_page,
    _selected_image_slide_indices,
    _slide_image_composition,
    _slide_image_list,
    _slide_image_palette,
    _slide_image_prompt,
    _slide_image_role,
    _source_path_figure_numbers,
    _source_path_page_number,
    _source_visual_request_text,
    _source_visual_slide_indices,
    _spec_layout_id,
    _style_image_direction,
    _topic_image_world,
)
from ppt_pipeline.style_catalog import (
    PPTStyleDecision,
    PPTTemplateDecision,
    TEMPLATE_RENDER_PROFILES,
    _deck_artifact_stem,
    _safe_deck_filename_stem,
    ppt_template_prompt_context,
    resolve_ppt_style,
    resolve_ppt_template,
)
from ppt_pipeline.render_primitives import (
    _add_autofit_text,
    _add_circle,
    _add_line,
    _add_outline_circle,
    _add_outline_rect,
    _add_rect,
    _add_template_icon,
    _add_text,
    _apply_soft_shadow,
    _clean_visible_text,
    _fit_font_size,
    _localized_label,
    _ppt_len,
    _prefers_english_spec,
    _prefers_english_text,
    _rgb,
    _template_icon_kind,
    _visual_len,
)
from ppt_pipeline.content_parsing import (
    _fallback_outline_from_message,
    _outline_titles_by_number,
    _parse_slide_ordinal,
    _parse_bullets,
    _is_reference_fragment_text,
    _is_intermediate_artifact_text,
    _strip_visible_source_anchor,
    _is_layout_instruction_text,
    _meaningful_bullets,
    _is_low_substance_bullet,
    _dedupe_visible_points,
    _evidence_visible_points,
    _slide_display_points,
    _clean_slide_cell,
    _normalize_slide_table_key,
    _slide_table_field,
    _is_markdown_separator_row,
    _split_markdown_row,
    _markdown_table_blocks,
    _parse_slide_table_block,
    _parse_multilingual_markdown_table,
    _generic_slide_title_number,
    _is_generic_slide_title,
    _infer_title_from_spec,
    _repair_slide_titles,
    _looks_like_empty_generic_deck,
    _requested_slide_count,
    _merge_specs,
    _fit_specs_to_requested_count,
    _extract_machine_deck_spec,
    _apply_machine_deck_spec,
    _semantic_text,
    _dedupe_semantic_items,
    _sequence_parts,
    _title_sequence_parts,
    _domain_sequence_defaults,
    _payload_rows,
    _meaningful_payload_rows,
    _replace_spec_layout,
    _fallback_from_incomplete_table,
    _repair_semantic_content_specs,
    parse_slide_specs,
)
from ppt_pipeline.presentation_assets import (
    _clear_template_slides,
    _crop_rendered_pdf_page_visual,
    _fit_image_to_box,
    _is_rendered_pdf_page_visual,
    _new_presentation,
    _prefer_deeper_pdf_page_crop,
    _prepare_image_for_ppt,
    _shorten_for_cell,
    _source_figure_numbers,
)
from ppt_pipeline.office_conversion import (
    _convert_powerpoint_to_pdf_powershell_windows,
    _convert_powerpoint_to_pdf_windows,
    _truthy_env,
)
from ppt_pipeline.quality_assurance import (
    _compact_specs_for_quality_repair,
    _adapt_sparse_specs_for_quality_repair,
    _convert_deck_to_quality_pdf,
    _render_quality_page_images,
    _semantic_min_font_size,
    _ppt_text_quality_issues,
    _actual_slide_library_layout_id,
    _quality_layout_family,
    _quality_layout_density,
    _semantic_slot_quality_issues,
    _display_status_value,
    _visual_quality_issues,
    _ppt_image_binding_issues,
    _quality_report_lines,
)
from ppt_pipeline.image_generation import (
    _ppt_imagegen_env,
    _ppt_openai_api_base,
    _ppt_image_bytes_from_response,
    _ppt_generate_image_direct,
    _valid_generated_image,
    _ppt_imagegen_cache_path,
    _copy_generated_image,
    _store_generated_image_cache,
    _generate_slide_image_job,
    _compact_exception,
    _enabled_imagegen,
    _enabled_research,
    _imagegen_script_path,
    _generate_slide_images as _generate_slide_images_impl,
)
import ppt_pipeline.image_generation as _image_generation_module
from ppt_pipeline.template_core import (
    _palette_extras,
    _is_scut_academic_chrome,
    _is_hitsz_academic_chrome,
    _academic_ref_colors,
    _academic_ref_fill,
    _add_academic_ref_panel,
    _academic_visible_text,
    _academic_sparse_points_boost,
    _academic_sparse_spec_boost,
    _add_academic_plain_list,
    _add_academic_plain_note,
    _template_profile,
    _style_for_template,
    _template_footer_label,
    _add_slide_background,
    _add_title_system,
    _split_points,
    _template_source_slide_index,
    _template_chrome_pdf_cache_path,
    _convert_template_to_pdf,
    _render_template_chrome_backgrounds,
    _template_source_placeholder_boxes,
    _box_or_default,
    _blank_layout,
    _template_layout_index,
    _template_slide_layout,
    _uses_template_layouts,
    _slide_role,
    _placeholder_by_idx,
    _placeholder_type_name,
    _first_text_placeholder,
    _shape_box_inches,
    _set_text_frame,
    _set_title_placeholder,
    _set_footer_placeholders,
    _add_template_chrome_background,
    _render_template_chrome_cover,
    _template_chrome_boxes,
    _add_template_chrome_title,
    _template_slide_bullets,
    _template_rich_bullets,
    _template_concept_terms,
    _image_overlay_labels,
    _add_image_overlay_labels,
    _is_generated_slide_image,
    _load_layout_registry,
    _layout_family_for_style,
    _allowed_layout_ids,
    _family_layout_ids,
    _normalize_layout_id,
    _default_layout_sequence,
    _explicit_layout_id,
    _clean_layout_markers,
    _select_template_layout_id,
    _slide_library_content_payload,
)
from ppt_pipeline.template_feature_layouts import (
    _render_template_chrome_image_layout,
    _render_template_chrome_claim_layout,
    _render_template_chrome_claim_band_layout,
    _sparse_support_points,
    _render_template_chrome_statement_sidebar_layout,
    _render_template_chrome_quote_rule_layout,
    _render_template_chrome_compact_mosaic_layout,
    _render_template_chrome_sparse_variant_layout,
    _hitsz_visible_text,
    _hitsz_points,
    _hitsz_support_items,
)
from ppt_pipeline.template_hitsz_layouts import (
    _render_hitsz_focus_strip_layout,
    _render_hitsz_axis_layout,
    _render_hitsz_grid_layout,
    _render_hitsz_metric_band_layout,
    _render_hitsz_claim_rule_layout,
    _render_hitsz_two_column_rule_layout,
    _render_hitsz_takeaway_list_layout,
    _render_hitsz_template_chrome_content,
)
from ppt_pipeline.template_data_layouts import (
    _render_template_chrome_two_column_plain_layout,
    _render_template_chrome_metric_focus_layout,
    _render_template_chrome_ladder_layout,
    _render_template_chrome_claim_variant_layout,
    _render_template_chrome_process_layout,
    _render_template_chrome_matrix_layout,
    _render_template_chrome_bigstat_layout,
    _numeric_highlights,
    _prominent_numeric_highlights,
    _render_template_chrome_results_bars_layout,
    _render_template_chrome_leaderboard_layout,
    _render_template_chrome_evidence_grid_layout,
    _render_template_chrome_motivation_layout,
)
from ppt_pipeline.template_card_layouts import (
    _template_card_variant,
    _render_template_chrome_triptych_cards_layout,
    _render_template_chrome_ladder_cards_layout,
    _render_template_chrome_mosaic_cards_layout,
    _render_template_chrome_loop_layout,
    _render_template_chrome_gallery_layout,
    _render_template_chrome_comparison_layout,
    _render_template_chrome_cards_layout,
    _render_template_chrome_content,
    _render_template_chrome_directory,
    _render_template_chrome_ending,
    _render_template_chrome_slide,
)



PPT_LAYOUT_REGISTRY = Path("departments/ppt_department/templates/layouts/layout_registry.json")
PPTProgressCallback = Callable[[dict[str, Any]], None]


def _generate_slide_images(**kwargs):
    original = _image_generation_module._ppt_generate_image_direct
    _image_generation_module._ppt_generate_image_direct = _ppt_generate_image_direct
    try:
        return _generate_slide_images_impl(**kwargs)
    finally:
        _image_generation_module._ppt_generate_image_direct = original


def _report_ppt_progress(on_progress: PPTProgressCallback | None, payload: dict[str, Any]) -> None:
    if on_progress is None:
        return
    try:
        on_progress(payload)
    except Exception:
        pass


def _delete_path(path: Path) -> None:
    if path.is_dir():
        shutil.rmtree(path)
    elif path.exists():
        path.unlink()


def _is_default_memory(path: Path) -> bool:
    if not path.exists():
        return True
    text = path.read_text(encoding="utf-8", errors="replace")
    return "No approved evolved memory yet." in text or "No approved HR memory yet." in text


def _move_memory_if_useful(src: Path, dst: Path) -> None:
    if src.exists() and _is_default_memory(dst):
        dst.parent.mkdir(parents=True, exist_ok=True)
        if dst.exists():
            dst.unlink()
        shutil.move(str(src), str(dst))


def _migrate_legacy_state_layout(root: Path, organization) -> None:
    state = state_root(root)
    legacy_memory = state / "memory" / "agents"
    for agent_id, agent in organization.agents.items():
        _move_memory_if_useful(
            legacy_memory / agent_id / "MEMORY.md",
            memory_file(agent.department_id, agent.id, root),
        )

    legacy_agents = state / "agents"
    for agent_id, agent in organization.agents.items():
        _move_memory_if_useful(
            legacy_agents / agent_id / "memory" / "MEMORY.md",
            memory_file(agent.department_id, agent.id, root),
        )

    for leftover in ("codex_homes", "agents", "leaders", "memory"):
        path = state / leftover
        if path.exists():
            _delete_path(path)
    old_db = legacy_db_path(root)
    if old_db.exists() and old_db != root / "data" / "janus.db":
        _delete_path(old_db)
    state.mkdir(parents=True, exist_ok=True)
    tmp_dir(root).mkdir(parents=True, exist_ok=True)
    for child in state.iterdir():
        if child.name != "tmp":
            _delete_path(child)


def _ensure_hr_memory(department_id: str, hr_id: str, root: Path) -> Path:
    path = hr_memory_file(department_id, root)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_text(f"# HR Memory: {hr_id}\n\nNo approved HR memory yet.\n", encoding="utf-8")
    return path


def _ensure_workspace_assets(root: Path, conn, organization) -> None:
    janus_db.seed_agents(conn, organization.agents)
    write_repo_templates(root)
    _migrate_legacy_state_layout(root, organization)
    for department in organization.departments.values():
        dept_agents = [agent for agent in organization.agents.values() if agent.department_id == department.id]
        ensure_hr_assets(department, dept_agents, root)
        _ensure_hr_memory(department.id, department.hr.id, root)
    for agent in organization.agents.values():
        ensure_memory(agent, root)
        ensure_agent_skill(agent, root)


@dataclass(frozen=True)
class PPTDebugResult:
    deck_path: Path
    notes_path: Path
    prompt: str
    answer: str
    session_id: str
    style_id: str
    style_label: str
    style_created: bool
    image_paths: list[Path]
    imagegen_errors: list[str]
    codex_error: str | None = None
    research_error: str | None = None
    research_pack: str | None = None


def build_ppt_debug_message(message: str, style: PPTStyleDecision, evidence_pack: str | None = None) -> str:
    evidence_block = ""
    if evidence_pack and evidence_pack.strip():
        evidence_block = (
            "\n\n联网检索素材（由 PPT Designer 内部研究阶段检索，带来源标注）：\n"
            "请把下面这些资料有机融入 slide plan，用来充实内容、增强说服力和饱满度；"
            "在 speaker_note 中保留关键来源，标注为 low confidence 或可能过时的条目要谨慎使用，不要编造未提供的数据。\n"
            f"{evidence_pack.strip()}\n"
        )
    return (
        "用户已进入 PPT 制作功能。请把下面的用户指令作为 PowerPoint 制作需求，"
        "必须生成可渲染为 .pptx 的 slide plan 和 speaker notes，不要只返回普通聊天回答。"
        "如果素材不足，也要基于合理假设生成一份草稿 PPT，不要停留在追问缺失输入。\n\n"
        "PPT 部门运行规则：\n"
        "- leader agent 负责治理；单一 PPT Designer 根据 style_id 负责研究、策略、视觉、版式、图像提示和质检。\n"
        "- 每页必须有单一主张、可编辑文本、视觉建议、讲稿备注和时间建议。\n"
        "- 需要图片时，在 visual 列用 `生成插图：...` 写出可执行的 gpt-image-2 支持图提示词；不要把图片做成不可编辑的整页截图。\n"
        "- `生成插图` 后必须写清具体主体/场景、动作或关系、视觉媒介、构图或镜头、材质与光线、情绪基调；禁止只写“科技感背景、抽象概念图、蓝色未来感、信息流视觉”等可套用到任何主题的泛化提示。\n"
        "- 图片风格必须随主题和所选风格显著变化：学术风偏出版级科学编辑插画，重大项目风偏真实工程/工业叙事，通用风由主题决定摄影、拼贴、绘画或 3D 静物媒介；不要所有主题都画成蓝色网络、发光大脑或多分区信息图。\n"
        "- 同一份 PPT 的生成图应共享一种媒介和色彩逻辑，但每张图要更换主体、景别和构图；生成图内不要放文字、标签、数字、图表或 UI，这些由 PowerPoint 可编辑元素承载。\n"
        "- 流程图、架构图、矩阵、表格、时间线、指标对比这类精确信息应写成 `可编辑流程图/可编辑架构图/可编辑矩阵...`，由 renderer 用 PowerPoint 形状绘制，不交给图片模型。\n"
        "- 如果用户提供了论文 PDF、PPTX、Word 或其他附件，且附件视觉上下文中包含框架图、方法图、系统图、实验图或截图，优先在 visual 列写 `使用附件原图：...`，让 renderer 插入源文档截图/提取图；不要重复要求 gpt-image-2 生成同类图片。\n"
        "- 参考 ppt_ref/ 中成熟 PPT 的质量方向：强封面、稳定标题/页脚、丰富视觉锚点、紧凑 callout、流程/矩阵/卡片/大图混排，避免整套 PPT 都是同一种简陋版式。\n"
        "- 输出内容要让后端 renderer 能做出可编辑且可靠打开的 PowerPoint：精确文字留给 PPT 文本框，生成图只做辅助视觉。\n\n"
        "【页面级 layout_id——必须使用】\n"
        "- 学术汇报风可选 layout_id：basic_content、motivation_compare、challenge_map、method_pipeline、method_loop、benchmark_metrics、result_big_numbers、results_bars、leaderboard_table、ablation_matrix、evidence_grid、case_gallery、media_showcase、summary_takeaways。\n"
        "- 重大项目风可选 layout_id：basic_content、project_target_map、domain_object_map、technical_route、workpackage_matrix、evaluation_dashboard、result_big_numbers、risk_action_table、milestone_roadmap、media_showcase、summary_takeaways。\n"
        "- 每页选择一个最匹配的 layout_id；renderer 会从学校多功能模板复制对应功能页，而不是重新绘制一套卡片。\n"
        "- 整套 PPT 都要保持版式变化：优先每种 layout_id 只用一次；10-12 页通常至少使用 6 种 layout_id，同一版式最多出现 2 次且不要相邻。\n"
        "- 先输出逐页标题列表，每页一行且只展示页码和标题/主题；详细 Markdown 页面表必须放入 ```janus-slide-plan 隐藏块。\n"
        "- 隐藏页面表之后必须附带一个 ```janus-deck-spec 机器块，内容为：{\"schema_version\":\"janus-multifunction-v1\",\"slides\":[{\"layout_id\":\"...\",\"content_spec\":{...}}]}。slides 与表格逐页对应。\n"
        "- content_spec 只能使用语义字段，例如 points/current/target/gap/input/output/steps/kpis/bars/headers/rows/table/captions/cards/nodes/layers/stages/risks/milestones/objective/protocol/scope/conclusion/next_step；不要写模板页码或 PowerPoint shape 名称。\n"
        "- 严格字段结构：challenge_map 使用 1-4 个 nodes{label,detail} 或 risks{risk,impact,mitigation}；method_loop 使用 3-4 个 steps/stages{label,detail}；benchmark_metrics 使用 kpis{value,label,note} 或 {group,metrics}。\n"
        "- evidence_grid/case_gallery 没有图片时必须提供实质 cards{title,points}；media_showcase 必须有真实附件图或生成图，否则改用流程页、文本证据页或 basic_content。\n"
        "- 卡片数量是上限而不是目标，只输出真实内容需要的卡片；renderer 会自动删除并重排多余卡片。content_spec.title 必须与表格 title 完全一致。\n"
        "- proof_object 列写本页主证据对象，例如：可编辑方法流程图、论文 Figure 2 原图、三平台指标卡、消融矩阵、生成概念插图。\n\n"
        "【内容充实度要求——这是质量关键，务必严格遵守】\n"
        "- message 列不要写一句空泛主张；除封面/尾页外，每页至少写 2 个、推荐 3 个完整独立要点，每个要点用 ` • ` 分隔（首点前也加 ` • `）。\n"
        "- message 列是会直接显示在 PPT 页面上的观众可见文字，只能写实质内容、论点、发现、机制解释或结论。\n"
        "- speaker_note 写 1-2 句演讲讲稿或证据补充，但不会显示在页面上；页面正文必须完整写在 message 列。\n"
        "- 整份 PPT 的页面标题、message、speaker_note 必须使用同一种主语言：用户明确要求中文就全中文，明确要求英文就全英文；未明确时跟随用户输入的主要语言。除专有名词、模型名、指标名和论文题名外，不要中英混排。\n"
        "- 严禁在 message 列写排版施工说明，例如“左侧...右侧...”“中间放...”“两个挑战卡片”“放公式简化版”“使用附件原图”“生成插图”“可编辑流程图”。这些只能写在 visual 列。\n"
        "- 每个要点都必须包含【具体对象 + 关系/机制/变化 + 影响/结论】中的至少两项；禁止只写“背景约束、总体目标、实施路径、风险控制、下一步动作、方法模块、验证证据”这类标签。\n"
        "- 每个要点都应是一句能独立站住的完整短句，尽量带具体数字、对比、案例、变量名或机制动词，而不是被截断的半句话。\n"
        "- 如果没有可靠数字，就写清楚因果链或判断依据，例如“高质量 UGC 生产受技能、工具和时间成本约束，导致长尾兴趣供给滞后”，不要写成“高质量 UGC 生产成本”。\n"
        "- 充分使用上面联网检索到的真实数据（带数字与来源），把它们分配到对应页的要点里，让每页饱满、有说服力，避免空洞和大面积留白。\n"
        "- 不要在 message 里写超过 4 个要点；宁可精炼也不要堆砌碎片。每个要点控制在 ~40 字以内，便于版面排布不超框。\n"
        "- visual 列只写一句简洁的视觉/版式建议（如：左侧要点卡 + 右侧数据高亮 / 可编辑四步流程 / 可编辑分层架构图 / 对比矩阵 / 使用附件原图：论文方法框架图 / 生成插图：个性化推荐信息流概念视觉），供 renderer 选择版式。\n"
        "- visual 列不会作为页面正文显示；不要把重要结论只写在 visual 列。\n"
        "- 8 页左右的 deck 至少安排 3 种不同页面结构，并在 1-2 页使用 `生成插图：...` 作为辅助视觉；其余页优先使用可编辑图形系统。\n"
        "- speaker_note 列写完整讲稿，包含关键数字与来源，供演讲者使用（不显示在页面上）。\n\n"
        "输出格式：\n"
        "- 用 Markdown 表格输出，表头严格为 layout_id、title、message、proof_object、visual、speaker_note、time。\n"
        "- 表格每行对应一页。message 列内用 ` • ` 分隔要点；不要在单元格里使用会破坏表格的竖线或换行。\n\n"
        f"已解析 PPT 风格：{style.label}（style_id={style.style_id}，matched_by={style.matched_by}）。\n"
        f"风格提示：{style.prompt}\n"
        f"用户指令：{message.strip()}"
        f"{evidence_block}"
    )


def build_research_scout_message(message: str, style: PPTStyleDecision) -> str:
    return (
        "你正在为一份 PowerPoint 做联网资料检索。请使用 web_search 工具真正上网搜索，"
        "找出可以放进这份 PPT、让它更充实、更有说服力、更饱满、更美观的素材。\n\n"
        "检索目标：\n"
        "- 支撑论点的事实与统计数据（带具体数字）\n"
        "- 让观点具体可感的真实案例或例子\n"
        "- 有说服力的金句、权威观点或引文\n"
        "- 增强饱满度的对比、时间线、规模/市场数字\n"
        "- 可用于做图表或视觉锚点的数据线索（给数据，不画图）\n\n"
        "检索要求：\n"
        "- 优先权威来源（论文、官方/政府网站、统计机构、一手文档），少用泛泛博客并标注。\n"
        "- 关注时效，优先较新的数字，过时的要注明。\n"
        "- 不要编造来源、数字、案例或引文；无法核实就说明。\n\n"
        "输出为带来源标注的 Evidence Pack，按页/章节分组，每条包含：\n"
        "Slide/Section、Item（事实/数字/案例/引文）、Why it fits（一句话：增强 substance/persuasion/fullness/visual 中的哪一项）、"
        "Source（标题 — URL（日期））、Confidence（high/medium/low，并注明过时或冲突）。\n"
        "把最权威、最有力的条目放在前面。只检索供给素材，不要写最终 slide 文案、视觉风格或图片提示词。\n\n"
        f"PPT 风格：{style.label}（{style.prompt}）\n"
        f"用户指令：{message.strip()}"
    )


def run_research_scout(
    *,
    organization,
    message: str,
    style: PPTStyleDecision,
    root: Path,
    dry_run: bool = False,
) -> tuple[str | None, str | None]:
    """Run the PPT research scout agent online and return (evidence_pack, error).

    Failure is non-fatal: returns (None, error) so the deck still generates.
    """
    scout = organization.agents.get("ppt")
    if scout is None or scout.department_id != "ppt_department":
        return None, "ppt agent is not configured"

    ensure_memory(scout, root)
    ensure_agent_skill(scout, root)

    scout_message = build_research_scout_message(message, style)
    scout_prompt = build_chat_prompt(
        agent=scout,
        memory=read_memory(scout, root),
        skill=read_skill(scout, root),
        history=[],
        user_message=scout_message,
        show_process=False,
    )
    try:
        pack = run_codex_exec_utf8(
            prompt=scout_prompt,
            agent_id=scout.id,
            root=root,
            dry_run=dry_run,
        )
    except CodexUnavailable as exc:
        return None, str(exc)
    return pack, None


def run_codex_exec_utf8(
    *,
    prompt: str,
    agent_id: str,
    root: Path,
    role: str = "agent",
    sandbox: str = DEFAULT_CODEX_SANDBOX,
    timeout_seconds: int = 900,
    dry_run: bool = False,
) -> str:
    if dry_run:
        return "[dry-run prompt]\n\n" + prompt

    codex_bin = resolve_codex_binary()
    if codex_bin is None:
        raise CodexUnavailable(
            "Codex CLI was not found. Install Codex or set JANUS_CODEX_BIN to the executable path."
        )

    tmp = tmp_dir(root)
    tmp.mkdir(parents=True, exist_ok=True)
    output_path = tmp / f"codex-last-message-{uuid.uuid4()}.txt"
    temp_parent = tmp / "codex"
    temp_parent.mkdir(parents=True, exist_ok=True)
    temp_home = tempfile.mkdtemp(prefix=f"{role}-{agent_id}-", dir=temp_parent)
    cmd = [
        codex_bin,
        "exec",
        "--cd",
        str(root),
        "--sandbox",
        sandbox,
        "--skip-git-repo-check",
        "--ephemeral",
        "--color",
        "never",
        "--output-last-message",
        str(output_path),
        "-",
    ]
    env = os.environ.copy()
    env["CODEX_HOME"] = temp_home
    env.update(load_auth_env(root))
    try:
        write_repo_templates(root)
        if shared_config_path(root).exists():
            shutil.copy2(shared_config_path(root), Path(temp_home) / "config.toml")
        if shared_auth_path(root).exists():
            shutil.copy2(shared_auth_path(root), Path(temp_home) / "auth.json")
        completed = subprocess.run(
            cmd,
            input=prompt.encode("utf-8"),
            capture_output=True,
            cwd=str(root),
            env=env,
            timeout=timeout_seconds,
            check=False,
        )
    except OSError as exc:
        raise CodexUnavailable(
            f"Codex CLI could not be launched ({exc}). Set JANUS_CODEX_BIN to a runnable CLI binary."
        ) from exc
    finally:
        shutil.rmtree(temp_home, ignore_errors=True)

    if output_path.exists():
        text = output_path.read_text(encoding="utf-8", errors="replace").strip()
    else:
        text = completed.stdout.decode("utf-8", errors="replace").strip()
    if completed.returncode != 0:
        stderr = completed.stderr.decode("utf-8", errors="replace").strip()
        raise CodexUnavailable(f"Codex exec failed with code {completed.returncode}: {stderr or text}")
    return text or completed.stdout.decode("utf-8", errors="replace").strip()


def _requested_slide_count(message: str) -> int:
    match = re.search(r"(\d{1,2})\s*(?:页|张|p|P|slides?|Slides?)", message)
    if match:
        return max(1, min(30, int(match.group(1))))
    return 6


def _fallback_slide_plan(message: str, style: PPTStyleDecision, error: str) -> str:
    count = _requested_slide_count(message)
    topic = message.strip().replace("|", "/")[:80] or "PPT 草稿"
    rows = [
        "| title | message | visual | speaker_note | time |",
        "|---|---|---|---|---|",
    ]
    outline = _fallback_outline_from_message(message, count)
    if outline:
        for index, (title, points, visual) in enumerate(outline, start=1):
            message_cell = " • " + " • ".join(points[:4]) if points else f" • 围绕“{title}”提炼核心内容"
            visual_cell = visual or "可编辑要点卡片 + 简洁图形锚点"
            rows.append(
                "| "
                f"{title.replace('|', '/')} | "
                f"{message_cell.replace('|', '/')} | "
                f"{visual_cell.replace('|', '/')} | "
                f"根据用户原始 prompt 的第 {index} 页要求展开讲解，保留关键信息并避免大段堆砌。| "
                "1 min |"
            )
        rows.append("## Debug note")
        rows.append("Codex 调用失败，已从用户逐页需求中提取本地 fallback 草稿以便继续渲染 PPT。")
        rows.append(error.replace("\n", " ")[:500])
        return "\n".join(rows)

    section_titles = [
        "标题与目标",
        "背景与痛点",
        "核心框架",
        "流程设计",
        "关键模块",
        "案例与证据",
        "效果与价值",
        "风险与边界",
        "下一步计划",
        "总结",
    ]
    visual_patterns = [
        "强封面：大标题、细长强调线、四个能力标签",
        "左侧洞察卡片 + 右侧场景图/概念图",
        "三卡片框架：输入、处理、输出",
        "横向流程图：阶段、动作、产物",
        "模块矩阵：角色、职责、接口",
        "大图锚点 + 关键证据 callout",
        "价值阶梯：效率、质量、可控性",
        "风险雷达：数据、模型、渲染、协作",
        "路线图：短期、中期、长期",
        "总结页：一句话结论 + 三个行动项",
    ]
    for index in range(1, count + 1):
        title = section_titles[index - 1] if index <= len(section_titles) else f"补充页 {index}"
        visual = visual_patterns[index - 1] if index <= len(visual_patterns) else "可编辑图形系统 + 辅助视觉"
        rows.append(
            "| "
            f"{title} | "
            f"围绕“{topic}”展开第 {index} 页内容；当前采用 {style.label}，用一句明确结论驱动本页，不堆砌长段文字。| "
            f"{visual}；使用 {style.label} 的色彩体系，图片只做视觉辅助，文字保持可编辑。| "
            "讲清本页在整套 PPT 中的作用，并提示后续可替换为真实素材。| "
            "1 min |"
        )
    rows.append("")
    rows.append("## Debug note")
    rows.append("Codex 调用失败，已生成本地 fallback 草稿以便测试 PPT 渲染链路。")
    rows.append(error.replace("\n", " ")[:500])
    return "\n".join(rows)


def _render_template_cover(slide, spec: SlideSpec, style: PPTStyleDecision, template: PPTTemplateDecision, total: int) -> None:
    _set_title_placeholder(slide, spec, style)
    subtitle_shape = _placeholder_by_idx(slide, 1)
    subtitle = (_meaningful_bullets(spec.message or "", limit=1) or [""])[0]
    if subtitle_shape is not None:
        _set_text_frame(subtitle_shape, [subtitle], size=18, color=style.palette["muted"], align=PP_ALIGN.CENTER)
    _set_footer_placeholders(slide, 1, total, template)


def _render_template_content(slide, spec: SlideSpec, style: PPTStyleDecision, template: PPTTemplateDecision, index: int, total: int) -> None:
    _set_title_placeholder(slide, spec, style)
    _set_footer_placeholders(slide, index, total, template)
    content = _placeholder_by_idx(slide, 1) or _placeholder_by_idx(slide, 2) or _first_text_placeholder(slide, exclude={0, 10, 11, 12})
    bullets = _meaningful_bullets(spec.message or "", limit=5)
    if not bullets:
        bullets = [spec.title or "本页核心观点"]
    if content is not None:
        _set_text_frame(content, bullets, size=18 if len(bullets) <= 3 else 15, color=style.palette["ink"], bullet=True)
        x, y, w, h = _shape_box_inches(content)
    else:
        x, y, w, h = 0.9, 1.8, 11.4, 4.8
        _add_autofit_text(
            slide, "\n".join(bullets), x, y, w, h,
            max_size=18, min_size=10, color=style.palette["ink"], anchor=MSO_ANCHOR.TOP,
            margin=0.1,
        )


def _render_template_directory(slide, spec: SlideSpec, style: PPTStyleDecision, template: PPTTemplateDecision, index: int, total: int) -> None:
    _set_footer_placeholders(slide, index, total, template)
    title_shape = _placeholder_by_idx(slide, 0) or _first_text_placeholder(slide)
    title = re.sub(r"^(目录|章节|chapter|section|agenda|outline)\s*[:：\-—]*\s*", "", spec.title or "", flags=re.IGNORECASE).strip()
    title = title or "汇报结构"
    if title_shape is not None:
        _set_text_frame(title_shape, [title], size=32, color=style.palette["ink"], bold_first=True, align=PP_ALIGN.CENTER)
    points = _meaningful_bullets(spec.message or "", limit=3)
    body_shape = _placeholder_by_idx(slide, 1) or _placeholder_by_idx(slide, 10)
    if body_shape is not None and points:
        _set_text_frame(body_shape, points, size=15, color=style.palette["muted"], align=PP_ALIGN.CENTER)


def _render_template_ending(slide, spec: SlideSpec, style: PPTStyleDecision, template: PPTTemplateDecision, index: int, total: int) -> None:
    _set_footer_placeholders(slide, index, total, template)
    title_shape = _placeholder_by_idx(slide, 0) or _first_text_placeholder(slide)
    title = spec.title if re.search(r"(谢谢|thanks|q&a|qa|讨论)", spec.title or "", re.IGNORECASE) else "谢谢观看"
    if title_shape is not None:
        _set_text_frame(title_shape, [title], size=30, color=style.palette["ink"], bold_first=True, align=PP_ALIGN.CENTER)
    body_shape = _placeholder_by_idx(slide, 10) or _placeholder_by_idx(slide, 1)
    points = _meaningful_bullets(spec.message or "", limit=3)
    if body_shape is not None and points:
        _set_text_frame(body_shape, points, size=14, color=style.palette["muted"], align=PP_ALIGN.CENTER)


def _render_template_layout_slide(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    template: PPTTemplateDecision,
    *,
    index: int,
    total: int,
    role: str,
) -> None:
    if role == "cover":
        _render_template_cover(slide, spec, style, template, total)
    elif role == "directory":
        _render_template_directory(slide, spec, style, template, index, total)
    elif role == "ending":
        _render_template_ending(slide, spec, style, template, index, total)
    else:
        _render_template_content(slide, spec, style, template, index, total)


def _add_footer(slide, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    muted = style.palette["muted"]
    _add_text(slide, _template_footer_label(style, template), 0.42, 6.95, 5.0, 0.3, size=8, color=muted, margin=0.02)
    _add_text(slide, f"{index} / {total}", 11.6, 6.95, 1.3, 0.3, size=8, color=muted, align=PP_ALIGN.RIGHT, margin=0.02)


def _render_cover_slide(slide, spec: SlideSpec, style: PPTStyleDecision, total: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_rect(slide, 0, 0, 13.333, 7.5, p["bg"])
    if template and template.template_id == "hitsz":
        _add_rect(slide, 0, 0, 13.333, 1.05, p["accent"])
        _add_rect(slide, 0.9, 1.55, 0.08, 3.6, p["accent"])
        _add_rect(slide, 0.9, 5.22, 5.2, 0.05, p["accent2"])
    elif template and template.template_id == "scut":
        _add_rect(slide, 0, 0, 13.333, 0.22, p["accent"])
        _add_rect(slide, 0.9, 1.5, 0.1, 3.7, p["accent"])
        _add_rect(slide, 0.9, 5.22, 5.6, 0.07, p["accent2"])
    else:
        _add_rect(slide, 0, 0, 13.333, 0.2, p["accent"])
        _add_rect(slide, 0.9, 1.5, 0.1, 3.7, p["accent2"])
    # restrained geometric motif on the right, not a cluttered field of dots
    _add_rect(slide, 8.9, 1.5, 3.5, 3.5, p["panel"], radius=True)
    _add_rect(slide, 8.9, 1.5, 3.5, 0.12, p["accent"])
    for i in range(3):
        _add_circle(slide, 9.4 + i * 0.95, 2.2, 0.55, [p["accent"], p["accent2"], ex["warm"]][i])
    _add_rect(slide, 9.3, 3.5, 2.7, 0.16, p["accent"])
    _add_rect(slide, 9.3, 3.9, 1.9, 0.16, p["accent2"])
    _add_rect(slide, 9.3, 4.3, 2.3, 0.16, ex["grid"])
    # title + subtitle, auto-fit so long Chinese titles never overflow
    _add_autofit_text(
        slide, spec.title or "PowerPoint Draft", 1.2, 1.6, 7.2, 2.1,
        max_size=40, min_size=20, color=p["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0.05,
    )
    subtitle = (_meaningful_bullets(spec.message or "", limit=1) or [""])[0]
    if subtitle:
        _add_autofit_text(
            slide, subtitle, 1.2, 3.85, 7.0, 1.3,
            max_size=16, min_size=11, color=p["muted"], anchor=MSO_ANCHOR.TOP, margin=0.05,
        )
    _add_rect(slide, 1.2, 5.5, 0.55, 0.55, p["accent"])
    _add_text(slide, f"{_template_footer_label(style, template)} · 共 {total} 页", 1.95, 5.62, 5.5, 0.4, size=13, color=p["ink"], bold=True, margin=0.04)
    _add_text(slide, datetime.now().strftime("%Y / %m / %d"), 1.95, 6.05, 4.0, 0.3, size=10, color=p["muted"], margin=0.04)


def _render_image_slide(slide, spec: SlideSpec, style: PPTStyleDecision, image_path: Path, index: int, total: int, prepared_dir: Path, template: PPTTemplateDecision | None = None) -> list[str]:
    p = style.palette
    warnings: list[str] = []
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="visual evidence", template=template)
    points = _meaningful_bullets(spec.message or "", limit=3)
    if not points:
        points = [spec.title or "本页核心观点"]
    n = len(points)
    top, bottom = 1.72, 6.5
    gap = 0.22
    card_h = (bottom - top - gap * (n - 1)) / n
    for i, point in enumerate(points[:3]):
        y = top + i * (card_h + gap)
        _add_rect(slide, 0.42, y, 5.05, card_h, p["panel"], radius=True)
        _add_rect(slide, 0.42, y, 0.09, card_h, p["accent2"])
        _add_circle(slide, 0.72, y + card_h * 0.5 - 0.09, 0.18, p["accent"])
        _add_autofit_text(
            slide, point, 1.05, y, 4.28, card_h,
            max_size=14, min_size=9, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE, margin=0.08,
        )
    _add_rect(slide, 5.75, 1.72, 7.15, 4.78, p["panel"], radius=True)
    _add_rect(slide, 5.95, 1.92, 6.75, 0.08, p["accent2"])
    prepared = _prepare_image_for_ppt(image_path, prepared_dir, index, spec=spec)
    if prepared:
        x, y, w, h = _fit_image_to_box(prepared, 6.0, 2.1, 6.65, 3.9)
        picture = slide.shapes.add_picture(str(prepared), _ppt_len(x), _ppt_len(y), width=_ppt_len(w), height=_ppt_len(h))
        picture.name = "janus-legacy-content-image"
        if _is_generated_slide_image(image_path):
            _add_image_overlay_labels(slide, style, spec, x, y, w, h, max_labels=4)
    else:
        warnings.append(f"Slide {index}: image could not be validated and was replaced by an editable placeholder: {image_path}")
        _render_bullet_feature_slide(slide, spec, style, index, total, template=template)
        return warnings
    _add_footer(slide, style, index, total, template=template)
    return warnings


def _semantic_terms(spec: SlideSpec, *, limit: int = 6) -> list[str]:
    raw = " ".join([spec.title or "", spec.message or "", spec.visual or ""])
    candidates = re.split(r"[\s,，;；。:：/、|]+", raw)
    stop = {
        "PPT",
        "ppt",
        "PowerPoint",
        "本页",
        "当前",
        "使用",
        "主题",
        "内容",
        "视觉",
        "辅助",
        "可编辑",
        "进行",
        "通过",
        "围绕",
    }
    terms: list[str] = []
    for item in candidates:
        token = item.strip(" -·•()（）[]【】")
        if len(token) < 2 or token in stop:
            continue
        if token not in terms:
            terms.append(token)
    fallback = ["目标", "输入", "智能体", "生成", "渲染", "反馈"]
    for token in fallback:
        if len(terms) >= limit:
            break
        if token not in terms:
            terms.append(token)
    return terms[:limit]


def _render_network_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="semantic map", template=template)
    claim = "；".join(_meaningful_bullets(spec.message or "", limit=3)) or spec.title or "本页核心观点"
    _add_text(slide, claim, 0.55, 1.35, 4.65, 0.78, size=12, color=p["muted"])
    terms = _semantic_terms(spec, limit=6)
    center_x, center_y = 8.85, 3.45
    _add_circle(slide, center_x - 0.62, center_y - 0.62, 1.24, p["accent"], line=p["accent2"])
    _add_text(slide, terms[0], center_x - 0.5, center_y - 0.14, 1.0, 0.26, size=9, color=p["bg"], bold=True, align=PP_ALIGN.CENTER, margin=0)
    positions = [(6.65, 1.82), (10.35, 1.78), (11.0, 4.65), (8.55, 5.45), (6.05, 4.55)]
    colors = [p["accent2"], ex["warm"], ex["good"], p["accent"], ex["danger"]]
    for i, (x, y) in enumerate(positions):
        _add_line(slide, center_x, center_y, x + 0.45, y + 0.45, ex["grid"], width=1.1)
        _add_circle(slide, x, y, 0.9, colors[i % len(colors)], line=ex["grid"])
        _add_text(slide, terms[i + 1], x + 0.08, y + 0.3, 0.74, 0.22, size=7, color=p["bg"], bold=True, align=PP_ALIGN.CENTER, margin=0)
    for i, point in enumerate(_split_points(claim, limit=3)[:3]):
        y = 2.45 + i * 1.02
        _add_rect(slide, 0.62, y, 4.75, 0.72, p["panel"], line=ex["grid"], radius=True)
        _add_rect(slide, 0.62, y, 0.12, 0.72, colors[i % len(colors)])
        _add_text(slide, point, 0.9, y + 0.13, 4.1, 0.34, size=10, color=p["ink"], bold=i == 0)


def _render_process_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="workflow", template=template)
    points = _slide_display_points(spec, limit=4)
    n = min(max(len(points), 1), 4)
    colors = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
    gap = 0.34
    card_w = (12.5 - gap * (n - 1)) / n
    x0, y, card_h = 0.42, 2.7, 2.7
    for i, point in enumerate((points or [spec.title or ""])[:n]):
        x = x0 + i * (card_w + gap)
        _add_rect(slide, x, y, card_w, card_h, p["panel"], radius=True)
        _add_rect(slide, x, y, card_w, 0.09, colors[i % len(colors)])
        _add_circle(slide, x + 0.26, y - 0.34, 0.66, colors[i % len(colors)])
        _add_text(slide, f"{i + 1}", x + 0.26, y - 0.27, 0.66, 0.5, size=18, color=p["bg"], bold=True, align=PP_ALIGN.CENTER, margin=0)
        _add_autofit_text(
            slide, point, x + 0.22, y + 0.45, card_w - 0.44, card_h - 0.7,
            max_size=15, min_size=9, color=p["ink"], anchor=MSO_ANCHOR.TOP, margin=0.06,
        )
        if i < n - 1:
            _add_text(slide, "→", x + card_w + 0.02, y + card_h / 2 - 0.2, gap, 0.4, size=18, color=colors[i % len(colors)], bold=True, align=PP_ALIGN.CENTER, margin=0)


def _render_matrix_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="capability matrix", template=template)
    rows = _slide_display_points(spec, limit=4)
    cols = [
        _localized_label(spec, "对象", "OBJECT"),
        _localized_label(spec, "机制", "MECHANISM"),
        _localized_label(spec, "产出", "OUTPUT"),
    ]
    x0, y0 = 0.82, 1.65
    cell_w, cell_h = 2.75, 0.82
    matrix_claim = "；".join(rows[:2]) or spec.title or ""
    _add_text(slide, matrix_claim, 8.98, 1.65, 3.2, 0.72, size=12, color=p["muted"])
    for c, col in enumerate([_localized_label(spec, "模块", "MODULE"), *cols]):
        _add_rect(slide, x0 + c * cell_w, y0, cell_w - 0.04, 0.58, p["accent"] if c == 0 else p["accent2"], radius=True)
        _add_text(slide, col, x0 + c * cell_w + 0.12, y0 + 0.13, cell_w - 0.28, 0.2, size=9, color=p["bg"], bold=True, align=PP_ALIGN.CENTER, margin=0)
    for r, row in enumerate((rows or [spec.title or ""])[:4]):
        y = y0 + 0.72 + r * cell_h
        _add_rect(slide, x0, y, cell_w - 0.04, cell_h - 0.08, ex["soft"], line=ex["grid"], radius=True)
        _add_text(slide, row, x0 + 0.15, y + 0.18, cell_w - 0.35, 0.28, size=9, color=p["ink"], bold=True)
        for c, col in enumerate(cols, start=1):
            _add_rect(slide, x0 + c * cell_w, y, cell_w - 0.04, cell_h - 0.08, p["panel"], line=ex["grid"], radius=True)
            label = _shorten_for_cell(row, 18) if c == 2 else col
            _add_text(slide, label, x0 + c * cell_w + 0.18, y + 0.19, cell_w - 0.36, 0.24, size=8, color=p["muted"], align=PP_ALIGN.CENTER, margin=0)
    if len(rows) > 1:
        _add_rect(slide, 8.95, 3.0, 3.35, 1.65, p["panel"], line=p["accent"], radius=True)
        _add_text(slide, "核心关系", 9.2, 3.25, 2.75, 0.25, size=11, color=p["ink"], bold=True)
        _add_text(slide, "；".join(rows[:2]), 9.2, 3.78, 2.75, 0.62, size=10, color=p["muted"])


def _render_timeline_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="roadmap", template=template)
    points = _slide_display_points(spec, limit=4)
    n = min(max(len(points), 1), 4)
    _add_line(slide, 1.0, 3.85, 12.4, 3.85, p["accent"], width=2.0)
    colors = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
    step = (12.4 - 1.0) / max(n, 1)
    for i, point in enumerate((points or [spec.title or ""])[:n]):
        x = 1.0 + step * (i + 0.5)
        _add_circle(slide, x - 0.3, 3.55, 0.6, colors[i % len(colors)])
        _add_text(slide, f"T{i + 1}", x - 0.3, 3.62, 0.6, 0.46, size=14, color=p["bg"], bold=True, align=PP_ALIGN.CENTER, margin=0)
        box_y = 1.95 if i % 2 == 0 else 4.45
        box_w = min(step - 0.3, 2.7)
        _add_rect(slide, x - box_w / 2, box_y, box_w, 1.3, p["panel"], radius=True)
        _add_rect(slide, x - box_w / 2, box_y, box_w, 0.08, colors[i % len(colors)])
        _add_autofit_text(
            slide, point, x - box_w / 2 + 0.1, box_y + 0.18, box_w - 0.2, 1.0,
            max_size=13, min_size=9, color=p["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.05,
        )
        _add_line(slide, x, 3.85, x, box_y + (1.3 if i % 2 == 0 else 0), ex["grid"], width=1.0)


def _render_summary_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="takeaways", template=template)
    points = _slide_display_points(spec, limit=3)
    claim = points[0] if points else (spec.title or "")
    _add_autofit_text(
        slide, claim, 0.42, 1.7, 12.5, 1.25,
        max_size=24, min_size=14, color=p["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0.1,
    )
    show = points[:3]
    n = len(show)
    gap = 0.34
    card_w = (12.5 - gap * (n - 1)) / n
    colors = [p["accent"], p["accent2"], ex["good"]]
    x0, y, card_h = 0.42, 3.35, 2.7
    for i, point in enumerate(show):
        x = x0 + i * (card_w + gap)
        _add_rect(slide, x, y, card_w, card_h, p["panel"], radius=True)
        _add_rect(slide, x, y, card_w, 0.09, colors[i % len(colors)])
        _add_circle(slide, x + 0.26, y + 0.28, 0.56, colors[i % len(colors)])
        _add_text(slide, "•", x + 0.26, y + 0.32, 0.56, 0.42, size=20, color=p["bg"], bold=True, align=PP_ALIGN.CENTER, margin=0)
        _add_autofit_text(
            slide, point, x + 0.22, y + 1.0, card_w - 0.44, card_h - 1.2,
            max_size=14, min_size=9, color=p["ink"], anchor=MSO_ANCHOR.TOP, margin=0.06,
        )


def _extract_highlight(text: str) -> str | None:
    """Pull a short headline number/percentage from a bullet, for emphasis."""
    if not text:
        return None
    skip_context = re.compile(r"(?:table|fig(?:ure)?|图|表|page|slide|section|sec\.?|第|r@|n@|recall@|ndcg@)\s*$", re.IGNORECASE)
    candidates: list[tuple[int, str]] = []
    for match in re.finditer(r"\d+(?:[.,]\d+)?\s*(?:%|％|亿|万|倍|分钟|秒|岁|年|美元|mmHg|kg|克)?", text):
        token = match.group().strip()
        if not any(ch.isdigit() for ch in token):
            continue
        before = text[max(0, match.start() - 16):match.start()]
        after = text[match.end():match.end() + 10]
        prev_char = text[match.start() - 1] if match.start() > 0 else ""
        next_char = text[match.end()] if match.end() < len(text) else ""
        if re.match(r"[A-Za-z]", prev_char) or re.match(r"[A-Za-z]", next_char):
            continue
        if skip_context.search(before):
            continue
        if before.endswith("@") or re.search(r"^\s*[.)、:：-]", after):
            continue
        value = token.replace(",", ".")
        score = 0
        if "." in value:
            score += 20
        if "%" in token or "％" in token:
            score += 15
        if re.search(r"(coherence|novelty|aesthetic|hallucination|recall|ndcg|r@|n@|ctr|cvr)", before + after, re.IGNORECASE):
            score += 8
        if re.fullmatch(r"\d{1,2}", token):
            score -= 8
        candidates.append((score, token))
    if candidates:
        candidates.sort(key=lambda item: item[0], reverse=True)
        return candidates[0][1]
    return None


def _render_bullet_feature_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    """Left: vertically distributed bullet cards. Right: data-highlight panel.

    The signature, content-rich layout — each bullet gets its own card sized to
    the available column, and the strongest number is surfaced as a big stat.
    """
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="key points", template=template)
    bullets = _meaningful_bullets(spec.message or "", limit=4)
    if not bullets:
        bullets = [spec.title or "本页核心观点"]

    # Left column: evenly distributed cards, font auto-fit so nothing overflows.
    col_x, col_w = 0.42, 7.05
    top, bottom = 1.72, 6.62
    n = len(bullets)
    gap = 0.22
    card_h = (bottom - top - gap * (n - 1)) / n
    accents = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
    for i, b in enumerate(bullets):
        y = top + i * (card_h + gap)
        _add_rect(slide, col_x, y, col_w, card_h, p["panel"], radius=True)
        _add_rect(slide, col_x, y, 0.09, card_h, accents[i % len(accents)])
        _add_circle(slide, col_x + 0.36, y + card_h * 0.5 - 0.09, 0.18, accents[i % len(accents)])
        _add_autofit_text(
            slide, b, col_x + 0.68, y, col_w - 0.88, card_h,
            max_size=15, min_size=9, color=p["ink"], anchor=MSO_ANCHOR.MIDDLE, margin=0.08,
        )

    # Right column: data-highlight panel anchored by the strongest number.
    rx, rw = 7.78, 5.12
    _add_rect(slide, rx, 1.72, rw, 4.9, p["panel"], radius=True)
    _add_rect(slide, rx, 1.72, rw, 0.1, p["accent"])
    prominent_highlights = _prominent_numeric_highlights("；".join(bullets), limit=1)
    highlight = prominent_highlights[0] if prominent_highlights else next(
        (h for b in bullets if (h := _extract_highlight(b)) and not re.fullmatch(r"\d{1,2}", h)),
        None,
    )
    if highlight:
        _add_text(slide, "DATA HIGHLIGHT", rx + 0.4, 2.1, rw - 0.8, 0.3, size=9, color=p["accent"], bold=True, margin=0)
        _add_autofit_text(
            slide, highlight, rx + 0.4, 2.5, rw - 0.8, 1.5,
            max_size=66, min_size=24, color=p["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0.05,
        )
        context = next((b for b in bullets if highlight in b or _extract_highlight(b) == highlight), bullets[0])
        _add_autofit_text(
            slide, context, rx + 0.4, 4.15, rw - 0.8, 2.1,
            max_size=14, min_size=10, color=p["muted"], anchor=MSO_ANCHOR.TOP, margin=0.08,
        )
    else:
        _add_text(slide, "本页要点", rx + 0.4, 2.1, rw - 0.8, 0.3, size=11, color=p["accent"], bold=True, margin=0)
        _add_autofit_text(
            slide, spec.title or "", rx + 0.4, 2.55, rw - 0.8, 1.4,
            max_size=30, min_size=16, color=p["ink"], bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0.05,
        )
        _add_autofit_text(
            slide, "；".join(bullets[:2]), rx + 0.4, 4.1, rw - 0.8, 2.2,
            max_size=14, min_size=10, color=p["muted"], anchor=MSO_ANCHOR.TOP, margin=0.08,
        )


def _render_bigstat_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    """A row of large statistic tiles — for data-heavy pages."""
    p = style.palette
    ex = _palette_extras(style)
    _add_slide_background(slide, style, variant=index, template=template)
    _add_title_system(slide, spec, style, index, total, kicker="by the numbers", template=template)
    bullets = _meaningful_bullets(spec.message or "", limit=4) or [spec.title or "本页核心观点"]
    stats = [(b, _extract_highlight(b)) for b in bullets]
    if sum(1 for _, h in stats if h) < 2:
        _render_bullet_feature_slide(slide, spec, style, index, total, template=template)
        return
    n = min(len(stats), 4)
    gap = 0.3
    total_w = 12.5
    card_w = (total_w - gap * (n - 1)) / n
    x0, y0, card_h = 0.42, 2.1, 3.0
    accents = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
    for i, (text, hl) in enumerate(stats[:n]):
        x = x0 + i * (card_w + gap)
        _add_rect(slide, x, y0, card_w, card_h, p["panel"], radius=True)
        _add_rect(slide, x, y0, card_w, 0.1, accents[i % len(accents)])
        _add_autofit_text(
            slide, hl or "—", x + 0.2, y0 + 0.35, card_w - 0.4, 1.2,
            max_size=48, min_size=20, color=accents[i % len(accents)], bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.05,
        )
        _add_autofit_text(
            slide, text, x + 0.2, y0 + 1.65, card_w - 0.4, card_h - 1.85,
            max_size=13, min_size=9, color=p["ink"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, margin=0.08,
        )


def _render_infographic_slide(slide, spec: SlideSpec, style: PPTStyleDecision, index: int, total: int, template: PPTTemplateDecision | None = None) -> None:
    visual_text = " ".join([spec.title or "", spec.visual or "", spec.message or ""])
    bullets = _parse_bullets(spec.message or "", limit=4)
    numbery = sum(1 for b in bullets if _extract_highlight(b))
    if index == total:
        _render_summary_slide(slide, spec, style, index, total, template=template)
    elif any(key in visual_text for key in ["流程", "workflow", "步骤", "链路", "阶段", "四步", "三步"]):
        _render_process_slide(slide, spec, style, index, total, template=template)
    elif any(key in visual_text for key in ["路线", "计划", "roadmap", "时间线", "下一步", "阶段计划"]):
        _render_timeline_slide(slide, spec, style, index, total, template=template)
    elif any(key in visual_text for key in ["大数字", "数据", "统计", "占比", "增长"]) and numbery >= 2:
        _render_bigstat_slide(slide, spec, style, index, total, template=template)
    else:
        # Default to the content-rich bullet+highlight layout — the workhorse.
        _render_bullet_feature_slide(slide, spec, style, index, total, template=template)


def _emu(inches: float) -> int:
    return int(inches * 914400)


def _solid_fill(color: str) -> str:
    return f'<a:solidFill><a:srgbClr val="{xml_escape(color.lstrip("#").upper())}"/></a:solidFill>'


def _xml_paragraph(
    text: str,
    *,
    size: int = 2200,
    bold: bool = False,
    color: str = "111827",
) -> str:
    text = xml_escape(text or " ")
    return (
        "<a:p>"
        '<a:pPr marL="0" indent="0"/>'
        f'<a:r><a:rPr lang="zh-CN" sz="{size}" b="{1 if bold else 0}">'
        f'<a:solidFill><a:srgbClr val="{xml_escape(color)}"/></a:solidFill>'
        '<a:latin typeface="Microsoft YaHei"/><a:ea typeface="Microsoft YaHei"/>'
        f"</a:rPr><a:t>{text}</a:t></a:r>"
        f'<a:endParaRPr lang="zh-CN" sz="{size}"/>'
        "</a:p>"
    )


def _xml_textbox(
    *,
    shape_id: int,
    name: str,
    left: float,
    top: float,
    width: float,
    height: float,
    paragraphs: list[str],
    font_size: int,
    color: str,
    bold: bool = False,
    fill: str | None = None,
) -> str:
    paragraph_xml = "".join(
        _xml_paragraph(item, size=font_size, bold=bold, color=color)
        for item in (paragraphs or [" "])
    )
    fill_xml = _solid_fill(fill) if fill else "<a:noFill/>"
    return f"""
<p:sp>
  <p:nvSpPr>
    <p:cNvPr id="{shape_id}" name="{xml_escape(name)}"/>
    <p:cNvSpPr txBox="1"/>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm>
      <a:off x="{_emu(left)}" y="{_emu(top)}"/>
      <a:ext cx="{_emu(width)}" cy="{_emu(height)}"/>
    </a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    {fill_xml}
    <a:ln><a:noFill/></a:ln>
  </p:spPr>
  <p:txBody>
    <a:bodyPr wrap="square" anchor="t" lIns="91440" tIns="45720" rIns="91440" bIns="45720"/>
    <a:lstStyle/>
    {paragraph_xml}
  </p:txBody>
</p:sp>
""".strip()


def _xml_rect(
    *,
    shape_id: int,
    name: str,
    left: float,
    top: float,
    width: float,
    height: float,
    fill: str,
) -> str:
    return f"""
<p:sp>
  <p:nvSpPr><p:cNvPr id="{shape_id}" name="{xml_escape(name)}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{_emu(left)}" y="{_emu(top)}"/><a:ext cx="{_emu(width)}" cy="{_emu(height)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    {_solid_fill(fill)}
    <a:ln><a:noFill/></a:ln>
  </p:spPr>
  <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
</p:sp>
""".strip()


def _xml_picture(
    *,
    shape_id: int,
    rel_id: str,
    left: float,
    top: float,
    width: float,
    height: float,
) -> str:
    return f"""
<p:pic>
  <p:nvPicPr>
    <p:cNvPr id="{shape_id}" name="Generated visual"/>
    <p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>
    <p:nvPr/>
  </p:nvPicPr>
  <p:blipFill>
    <a:blip r:embed="{rel_id}"/>
    <a:stretch><a:fillRect/></a:stretch>
  </p:blipFill>
  <p:spPr>
    <a:xfrm><a:off x="{_emu(left)}" y="{_emu(top)}"/><a:ext cx="{_emu(width)}" cy="{_emu(height)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
</p:pic>
""".strip()


def _slide_xml(spec: SlideSpec, index: int, total: int, style: PPTStyleDecision, has_image: bool) -> str:
    title = spec.title or f"Slide {index}"
    palette = style.palette
    body = _slide_display_points(spec, limit=5) or [" "]
    text_width = 5.75 if has_image else 11.6
    visual_xml = (
        _xml_picture(shape_id=8, rel_id="rId2", left=6.75, top=1.45, width=5.75, height=4.55)
        if has_image
        else _xml_rect(shape_id=8, name="Visual anchor", left=7.3, top=1.55, width=4.75, height=4.25, fill=palette["panel"])
    )
    return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
       xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
       xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:nvGrpSpPr>
        <p:cNvPr id="1" name=""/>
        <p:cNvGrpSpPr/>
        <p:nvPr/>
      </p:nvGrpSpPr>
      <p:grpSpPr>
        <a:xfrm>
          <a:off x="0" y="0"/><a:ext cx="0" cy="0"/>
          <a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/>
        </a:xfrm>
      </p:grpSpPr>
      {_xml_rect(shape_id=2, name="Background", left=0, top=0, width=13.333, height=7.5, fill=palette["bg"])}
      {_xml_rect(shape_id=3, name="Accent bar", left=0, top=0, width=0.16, height=7.5, fill=palette["accent"])}
      {_xml_rect(shape_id=4, name="Top rule", left=0.55, top=1.15, width=11.9, height=0.04, fill=palette["accent2"])}
      {_xml_textbox(shape_id=5, name="Title", left=0.62, top=0.32, width=12.0, height=0.78, paragraphs=[title], font_size=3000, color=palette["ink"], bold=True)}
      {_xml_textbox(shape_id=6, name="Body", left=0.72, top=1.45, width=text_width, height=4.75, paragraphs=body, font_size=1900, color=palette["ink"], fill=palette["panel"])}
      {visual_xml}
      {_xml_textbox(shape_id=7, name="Footer", left=0.72, top=6.62, width=11.8, height=0.3, paragraphs=[f"{style.label} | {index} / {total}"], font_size=950, color=palette["muted"])}
    </p:spTree>
  </p:cSld>
  <p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>
"""


def _write_file(zf: zipfile.ZipFile, path: str, text: str) -> None:
    zf.writestr(path, text.encode("utf-8"))


def _content_type_defaults(image_paths: dict[int, Path | list[Path]]) -> str:
    defaults = [
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
        '<Default Extension="xml" ContentType="application/xml"/>',
    ]
    exts = {
        path.suffix.lower().lstrip(".")
        for value in image_paths.values()
        for path in _slide_image_list(value)
    }
    if "png" in exts:
        defaults.append('<Default Extension="png" ContentType="image/png"/>')
    if "jpg" in exts or "jpeg" in exts:
        defaults.append('<Default Extension="jpg" ContentType="image/jpeg"/>')
        defaults.append('<Default Extension="jpeg" ContentType="image/jpeg"/>')
    return "\n  ".join(defaults)


def _render_debug_pptx(
    *,
    root: Path,
    user_id: str,
    agent_id: str,
    session_id: str,
    user_message: str,
    assistant_answer: str,
    style: PPTStyleDecision,
    image_paths: dict[int, Path | list[Path]] | None = None,
    imagegen_errors: list[str] | None = None,
    selected_template: str | None = None,
    source_image_paths: list[Path] | None = None,
    specs_override: list[SlideSpec] | None = None,
    qa_report_lines: list[str] | None = None,
    on_progress: PPTProgressCallback | None = None,
    progress_phase: str = "render",
    progress_attempt: int = 1,
) -> tuple[Path, Path, list[str], PPTTemplateDecision]:
    specs = specs_override or parse_slide_specs(user_message, assistant_answer)
    if not specs:
        specs = [SlideSpec(title="PPT Draft", message=user_message)]

    source_image_map = _map_source_visuals_to_slides(specs, source_image_paths)
    image_paths = _merge_slide_image_maps(image_paths, source_image_map)
    deck_name = _deck_artifact_stem(user_message, specs)
    output_dir = root / "outputs" / "ppt_department" / f"{session_id}-{deck_name}"
    output_dir.mkdir(parents=True, exist_ok=True)
    deck_path = output_dir / f"{deck_name}.pptx"
    notes_path = output_dir / "speaker_notes.md"

    requested_template = resolve_ppt_template(root, selected_template)
    template = requested_template
    render_style = _style_for_template(style, template)
    prepared_dir = output_dir / "prepared_media"
    prepared_dir.mkdir(parents=True, exist_ok=True)

    image_warnings = list(imagegen_errors or [])
    profile = _template_profile(template)
    library_rendered = False
    if template.path and profile.get("engine") == "slide_library":
        layout_ids = [
            "cover" if index == 1 else _select_template_layout_id(
                root=root,
                spec=spec,
                style=render_style,
                index=index,
                total=len(specs),
            )
            for index, spec in enumerate(specs, start=1)
        ]
        required_layouts = {layout_id for layout_id in layout_ids if layout_id != "cover"}
        validation_errors = validate_slide_library(template.path, required_layouts)
        if validation_errors:
            image_warnings.extend(f"Slide-library validation: {error}" for error in validation_errors)
        else:
            try:
                prs, library_warnings = render_slide_library_presentation(
                    template_path=template.path,
                    template_id=template.template_id,
                    specs=specs,
                    layout_ids=layout_ids,
                    content_payloads=[_slide_library_content_payload(spec) for spec in specs],
                    image_paths=image_paths,
                    on_progress=on_progress,
                    progress_phase=progress_phase,
                    progress_attempt=progress_attempt,
                )
                image_warnings.extend(library_warnings)
                library_rendered = True
            except Exception as exc:
                image_warnings.append(f"Slide-library renderer failed; used legacy fallback: {_compact_exception(exc)}")

    if not library_rendered:
        prs, template = _new_presentation(root, selected_template)
        if template.template_id != requested_template.template_id:
            image_warnings.append(
                f"Requested template {requested_template.label} ({requested_template.template_id}) could not be loaded; "
                f"used {template.label} ({template.template_id}) instead."
            )
        render_style = _style_for_template(style, template)
        blank_layout = _blank_layout(prs)
        slide_roles = [_slide_role(index, len(specs), spec, template) for index, spec in enumerate(specs, start=1)]
        template_backgrounds = _render_template_chrome_backgrounds(root, template, slide_roles, prepared_dir)
        if template.path and slide_roles and not template_backgrounds:
            image_warnings.append(
                "Template chrome backgrounds could not be rendered; install PyMuPDF so selected PPT templates are applied as full-slide backgrounds."
            )
        for index, spec in enumerate(specs, start=1):
            _report_ppt_progress(on_progress, {
                "phase": progress_phase,
                "current_slide": index,
                "total_slides": len(specs),
                "phase_current": index,
                "phase_total": len(specs),
                "slide_title": spec.title or f"Slide {index}",
                "attempt": progress_attempt,
                "message": f"正在制作第 {index}/{len(specs)} 页：{spec.title or f'Slide {index}'}",
            })
            role = slide_roles[index - 1]
            template_background = template_backgrounds.get(role)
            if template_background is not None:
                slide = prs.slides.add_slide(blank_layout)
                _render_template_chrome_slide(
                    slide,
                    template_background,
                    spec,
                    render_style,
                    template,
                    root,
                    role=role,
                    index=index,
                    total=len(specs),
                    image_path=_first_slide_image(image_paths.get(index)) if role != "cover" else None,
                    prepared_dir=prepared_dir,
                )
                continue
            template_layout = _template_slide_layout(prs, template, role) if _uses_template_layouts(template) else None
            slide = prs.slides.add_slide(template_layout or blank_layout)
            if template_layout is not None:
                _render_template_layout_slide(
                    slide,
                    spec,
                    render_style,
                    template,
                    index=index,
                    total=len(specs),
                    role=role,
                )
                continue
            if index == 1:
                _render_cover_slide(slide, spec, render_style, len(specs), template=template)
                continue
            image_path = _first_slide_image(image_paths.get(index))
            if image_path:
                image_warnings.extend(
                    _render_image_slide(slide, spec, render_style, image_path, index, len(specs), prepared_dir, template=template)
                )
            else:
                _render_infographic_slide(slide, spec, render_style, index, len(specs), template=template)

    prs.slide_width = _ppt_len(13.333)
    prs.slide_height = _ppt_len(7.5)

    prs.core_properties.author = "OPL PPT Department"
    prs.core_properties.title = specs[0].title if specs else deck_name
    prs.core_properties.subject = f"{style.label} PowerPoint deck"
    prs.core_properties.keywords = f"OPL,PPT,{style.style_id},{template.template_id}"
    prs.save(deck_path)
    _report_ppt_progress(on_progress, {
        "phase": progress_phase,
        "current_slide": len(specs),
        "total_slides": len(specs),
        "phase_current": len(specs),
        "phase_total": len(specs),
        "slide_title": specs[-1].title if specs else "",
        "attempt": progress_attempt,
        "message": f"{len(specs)}/{len(specs)} 页已完成，正在保存 PPTX",
    })

    notes_lines = [
        f"# Speaker Notes for {deck_name}",
        f"- user: {user_id}",
        f"- agent: {agent_id}",
        f"- session: {session_id}",
        f"- style: {style.label} ({style.style_id})",
        f"- template: {template.label} ({template.template_id})",
        f"- style created this run: {style.created}",
        "",
    ]
    if image_paths:
        notes_lines.append("## Slide Images")
        for index, paths in sorted(image_paths.items()):
            label = "source attachment visual" if index in source_image_map else "generated image"
            for path in paths:
                notes_lines.append(f"- Slide {index}: {path} ({label})")
        notes_lines.append("")
    if image_warnings:
        notes_lines.append("## Imagegen Errors")
        notes_lines.extend(f"- {err}" for err in image_warnings)
        notes_lines.append("")
    if qa_report_lines:
        notes_lines.extend(qa_report_lines)
        notes_lines.append("")
    for index, spec in enumerate(specs, start=1):
        notes_lines.append(f"## Slide {index}: {spec.title or f'Slide {index}'}")
        notes_lines.append(spec.speaker_note or spec.message or "(no note)")
        notes_lines.append("")
    notes_path.write_text("\n".join(notes_lines).rstrip() + "\n", encoding="utf-8")
    return deck_path, notes_path, image_warnings, template


def _ppt_quality_issue_score(issues: list[str]) -> int:
    return sum(
        20 if re.search(r"整份 PPT.*(?:没有可用的内容图片|未成功插入)", issue) else 1
        for issue in issues
    )


def _ppt_missing_required_image(issues: list[str]) -> bool:
    return any(
        re.search(r"整份 PPT.*(?:没有可用的内容图片|未成功插入)", issue)
        for issue in issues
    )


def render_styled_ppt_artifact(
    *,
    root: Path,
    user_id: str,
    agent_id: str,
    session_id: str,
    user_message: str,
    assistant_answer: str,
    selected_style: str | None = None,
    selected_template: str | None = None,
    source_image_paths: list[Path] | None = None,
    pre_generated_image_paths: dict[int, Path] | None = None,
    host_imagegen_errors: list[str] | None = None,
    enable_imagegen: bool | None = None,
    minimum_content_images: int = 1,
    on_progress: PPTProgressCallback | None = None,
) -> tuple[Path, Path, list[str], PPTStyleDecision, PPTTemplateDecision]:
    """Render a normal chat PPT answer using the styled PowerPoint renderer.

    This keeps the formal chat flow on the same visual system as the PPT debug
    generator without re-running Codex or research.
    """
    style = resolve_ppt_style(
        message=user_message,
        root=root,
        selected_style=selected_style,
        selected_agent_id=agent_id,
        save_new_style=True,
    )
    specs = parse_slide_specs(user_message, assistant_answer)
    if not specs:
        raise ValueError("未能从 agent 回答中解析出 PPT 页面表，已停止生成，避免交付空白草稿。")
    source_image_map = _map_source_visuals_to_slides(specs, source_image_paths)
    image_paths: dict[int, Path] = {}
    for raw_index, raw_path in (pre_generated_image_paths or {}).items():
        try:
            index = int(raw_index)
            path = Path(raw_path)
        except (TypeError, ValueError):
            continue
        if index > 0 and path.is_file():
            image_paths[index] = path
    imagegen_errors = list(host_imagegen_errors or [])
    if not image_paths and not source_image_map:
        image_paths, fallback_errors = _generate_slide_images(
            root=root,
            session_id=session_id,
            specs=specs,
            style=style,
            enable_imagegen=enable_imagegen,
            selected_template=selected_template,
            reserved_image_indices=set(source_image_map),
            on_progress=on_progress,
        )
        imagegen_errors.extend(fallback_errors)
    deck_path, notes_path, render_warnings, effective_template = _render_debug_pptx(
        root=root,
        user_id=user_id,
        agent_id=agent_id,
        session_id=session_id,
        user_message=user_message,
        assistant_answer=assistant_answer,
        style=style,
        image_paths=image_paths,
        imagegen_errors=imagegen_errors,
        selected_template=selected_template,
        source_image_paths=source_image_paths,
        on_progress=on_progress,
        progress_phase="render",
        progress_attempt=1,
    )
    _report_ppt_progress(on_progress, {
        "phase": "qa",
        "current_slide": len(specs),
        "total_slides": len(specs),
        "phase_current": 0,
        "phase_total": 1,
        "attempt": 1,
        "message": f"{len(specs)}/{len(specs)} 页已完成，正在检查排版和视觉质量",
    })
    report_lines, passed, issues = _quality_report_lines(
        deck_path=deck_path,
        specs=specs,
        image_paths=_merge_slide_image_maps(image_paths, source_image_map),
        root=root,
        attempt=1,
        minimum_images=minimum_content_images,
    )
    qa_history = [*report_lines]
    active_specs = specs
    final_passed = passed
    final_issues = list(issues)
    best_deck = deck_path.read_bytes()
    best_notes = notes_path.read_bytes()
    best_score = _ppt_quality_issue_score(final_issues)
    best_template = effective_template
    qa_attempt = 1
    _report_ppt_progress(on_progress, {
        "phase": "qa",
        "current_slide": len(specs),
        "total_slides": len(specs),
        "phase_current": 1,
        "phase_total": 1,
        "attempt": 1,
        "message": (
            f"{len(specs)}/{len(specs)} 页排版检查通过"
            if passed
            else f"排版检查发现 {len(issues)} 个待处理问题，准备定向修复"
        ),
    })
    if minimum_content_images > 0 and _ppt_missing_required_image(final_issues):
        _report_ppt_progress(on_progress, {
            "phase": "image",
            "status": "retrying",
            "current_slide": len(specs),
            "total_slides": len(specs),
            "phase_current": 0,
            "phase_total": 1,
            "attempt": 2,
            "message": "自检发现整份 PPT 尚无有效图片，正在重新生成配图并调整图片页版式",
        })
        combined_images = _merge_slide_image_maps(image_paths, source_image_map)
        retry_images: dict[int, Path] = {}
        if not combined_images:
            retry_images, retry_errors = _generate_slide_images(
                root=root,
                session_id=session_id,
                specs=active_specs,
                style=style,
                enable_imagegen=enable_imagegen,
                selected_template=selected_template,
                reserved_image_indices=set(source_image_map),
                minimum_images=1,
                on_progress=on_progress,
            )
            imagegen_errors.extend(retry_errors)
            image_paths.update(retry_images)
            combined_images = _merge_slide_image_maps(image_paths, source_image_map)
        if combined_images:
            image_slide_numbers = {
                index for index, paths in combined_images.items()
                if len(paths) == 1
            }
            candidate_specs = _adapt_sparse_specs_for_quality_repair(
                active_specs,
                slide_numbers=image_slide_numbers,
                image_indices=image_slide_numbers,
            )
            qa_attempt += 1
            deck_path, notes_path, candidate_render_warnings, candidate_template = _render_debug_pptx(
                root=root,
                user_id=user_id,
                agent_id=agent_id,
                session_id=session_id,
                user_message=user_message,
                assistant_answer=assistant_answer,
                style=style,
                image_paths=image_paths,
                imagegen_errors=imagegen_errors,
                selected_template=selected_template,
                source_image_paths=source_image_paths,
                specs_override=candidate_specs,
                qa_report_lines=None,
                on_progress=on_progress,
                progress_phase="repair",
                progress_attempt=qa_attempt,
            )
            next_report, candidate_passed, candidate_issues = _quality_report_lines(
                deck_path=deck_path,
                specs=candidate_specs,
                image_paths=combined_images,
                root=root,
                attempt=qa_attempt,
                minimum_images=minimum_content_images,
            )
            qa_history.extend([
                "",
                "## PPT QA Minimum-image Repair",
                "- 首轮产物未满足至少一张内容图片的硬性要求，已复用现有图片或重新生成，并重排对应页面。",
                "",
                *next_report,
            ])
            candidate_score = _ppt_quality_issue_score(candidate_issues)
            if candidate_score < best_score:
                active_specs = candidate_specs
                final_passed = candidate_passed
                final_issues = list(candidate_issues)
                best_score = candidate_score
                best_deck = deck_path.read_bytes()
                best_notes = notes_path.read_bytes()
                render_warnings = candidate_render_warnings
                best_template = candidate_template
            else:
                deck_path.write_bytes(best_deck)
                notes_path.write_bytes(best_notes)
        else:
            qa_history.extend([
                "",
                "## PPT QA Minimum-image Repair",
                "- 已触发最低图片数量补救，但附件中没有可用视觉素材，且图片生成服务未返回有效图片。",
                "",
            ])
    for repair_level in (1, 2):
        if final_passed:
            break
        text_repairable = [
            issue
            for issue in final_issues
            if re.search(r"文字密度偏高|出现过小字号", issue)
        ]
        semantic_repairable = [
            issue
            for issue in final_issues
            if re.search(
                r"当前态|目标态|关键差距|有效步骤|流程步骤|数据行不足|填充率过低|总结要点|循环节点|KPI 卡片|未完整绑定|空表格|数量不一致",
                issue,
            )
        ]
        sparse_repairable = [
            issue
            for issue in final_issues
            if re.search(r"视觉内容偏少|文本框过少|版式可能偏单调", issue)
        ]
        repair_slides = {
            int(match.group(1))
            for issue in text_repairable
            if (match := re.search(r"第\s*(\d+)\s*页", issue))
        }
        semantic_slides = {
            int(match.group(1))
            for issue in semantic_repairable
            if (match := re.search(r"第\s*(\d+)\s*页", issue))
        }
        sparse_slides = {
            int(match.group(1))
            for issue in sparse_repairable
            if (match := re.search(r"第\s*(\d+)\s*页", issue))
        }
        image_binding_slides = (
            set(_merge_slide_image_maps(image_paths, source_image_map))
            if any(re.search(r"图片未完整绑定|未成功插入", issue) for issue in final_issues)
            else set()
        )
        sparse_slides.update(image_binding_slides)
        if not repair_slides and not semantic_slides and not sparse_slides:
            qa_history.extend([
                "",
                "## PPT QA Repair Decision",
                "- 当前问题属于图片、语义槽位或页面稀疏度问题，不执行全局文字压缩，避免破坏完整标题和内容结构。",
                "",
            ])
            break
        candidate_specs = _repair_semantic_content_specs(user_message, active_specs)
        if sparse_slides:
            candidate_specs = _adapt_sparse_specs_for_quality_repair(
                candidate_specs,
                slide_numbers=sparse_slides,
                image_indices=set(_merge_slide_image_maps(image_paths, source_image_map)),
            )
        if repair_slides:
            candidate_specs = _compact_specs_for_quality_repair(
                candidate_specs,
                level=repair_level,
                slide_numbers=repair_slides,
            )
        qa_history.extend(
            [
                "",
                "## PPT QA Repair Plan",
                f"- 自动自检发现排版风险，执行第 {repair_level} 次修复重渲。",
                f"- 语义结构修复页面：{', '.join(map(str, sorted(semantic_slides))) or '无'}；稀疏版式重配页面：{', '.join(map(str, sorted(sparse_slides))) or '无'}；文字密度修复页面：{', '.join(map(str, sorted(repair_slides))) or '无'}；标题保持完整。",
                "",
            ]
        )
        qa_attempt += 1
        deck_path, notes_path, candidate_render_warnings, candidate_template = _render_debug_pptx(
            root=root,
            user_id=user_id,
            agent_id=agent_id,
            session_id=session_id,
            user_message=user_message,
            assistant_answer=assistant_answer,
            style=style,
            image_paths=image_paths,
            imagegen_errors=imagegen_errors,
            selected_template=selected_template,
            source_image_paths=source_image_paths,
            specs_override=candidate_specs,
            qa_report_lines=None,
            on_progress=on_progress,
            progress_phase="repair",
            progress_attempt=qa_attempt,
        )
        next_report, candidate_passed, candidate_issues = _quality_report_lines(
            deck_path=deck_path,
            specs=candidate_specs,
            image_paths=_merge_slide_image_maps(image_paths, source_image_map),
            root=root,
            attempt=qa_attempt,
            minimum_images=minimum_content_images,
        )
        qa_history.extend(["", *next_report])
        candidate_score = _ppt_quality_issue_score(candidate_issues)
        if candidate_score < best_score:
            active_specs = candidate_specs
            final_passed = candidate_passed
            final_issues = list(candidate_issues)
            best_score = candidate_score
            best_deck = deck_path.read_bytes()
            best_notes = notes_path.read_bytes()
            render_warnings = candidate_render_warnings
            best_template = candidate_template
        else:
            deck_path.write_bytes(best_deck)
            notes_path.write_bytes(best_notes)
            qa_history.extend([
                "",
                "## PPT QA Best-version Guard",
                "- 本轮修复未降低问题数量，已恢复上一版，禁止用更差的最后一次尝试覆盖成品。",
                "",
            ])
            break
    deck_path.write_bytes(best_deck)
    notes_path.write_bytes(best_notes)
    notes_text = notes_path.read_text(encoding="utf-8", errors="replace")
    final_status = [
        "",
        "## PPT QA Final Decision",
        "- Status: passed" if final_passed else "- Status: delivered with remaining warnings",
        "- Delivery rule: the deck is checked after PDF rendering; remaining warnings are listed above for traceability.",
        "",
    ]
    notes_text = notes_text.rstrip() + "\n\n" + "\n".join(qa_history)
    notes_path.write_text(notes_text.rstrip() + "\n" + "\n".join(final_status), encoding="utf-8")
    user_warnings = list(dict.fromkeys([*render_warnings, *final_issues]))
    return deck_path, notes_path, user_warnings, style, best_template


def generate_ppt_debug_artifact(
    *,
    message: str,
    root: Path | None = None,
    user_id: str = "debug",
    session_id: str | None = None,
    dry_run: bool = False,
    fallback_on_codex_error: bool = True,
    selected_style: str | None = None,
    selected_agent_id: str | None = "ppt",
    selected_template: str | None = None,
    save_new_style: bool = True,
    enable_imagegen: bool | None = None,
    enable_research: bool | None = None,
) -> PPTDebugResult:
    message = message.strip()
    if not message:
        raise ValueError("message is required")

    root = Path(root or os.environ.get("JANUS_WORKSPACE", WORKSPACE)).resolve()
    os.environ["JANUS_WORKSPACE"] = str(root)

    organization = load_organization(root)
    agent = organization.agents.get("ppt")
    if agent is None or agent.department_id != "ppt_department":
        raise ValueError("PPT leader runtime agent is not configured")

    style = resolve_ppt_style(
        message=message,
        root=root,
        selected_style=selected_style,
        selected_agent_id=selected_agent_id,
        save_new_style=save_new_style,
    )

    conn = janus_db.connect(root)
    try:
        janus_db.migrate(conn)
        _ensure_workspace_assets(root, conn, organization)
    finally:
        conn.close()

    ensure_memory(agent, root)
    ensure_agent_skill(agent, root)

    # Step 1: run the research scout online to gather a source-cited evidence pack
    # that enriches the deck. Non-fatal: the deck still generates if research fails.
    research_pack: str | None = None
    research_error: str | None = None
    if _enabled_research(enable_research):
        research_pack, research_error = run_research_scout(
            organization=organization,
            message=message,
            style=style,
            root=root,
            dry_run=dry_run,
        )

    # Step 2: hand the evidence pack to the leader and generate the slide plan.
    ppt_message = (
        build_ppt_debug_message(message, style, research_pack)
        + ppt_template_prompt_context(root, selected_template)
    )
    prompt = build_chat_prompt(
        agent=agent,
        memory=read_memory(agent, root),
        skill=read_skill(agent, root),
        history=[],
        user_message=ppt_message,
        show_process=True,
    )
    codex_error = None
    try:
        answer = run_codex_exec_utf8(prompt=prompt, agent_id=agent.id, root=root, dry_run=dry_run)
    except CodexUnavailable as exc:
        if not fallback_on_codex_error:
            raise
        codex_error = str(exc)
        answer = _fallback_slide_plan(message, style, codex_error)

    debug_session_id = session_id or f"debug-{secrets.token_hex(8)}"
    specs = parse_slide_specs(ppt_message, answer) or [SlideSpec(title="PPT Draft", message=message)]
    source_image_map: dict[int, list[Path]] = {}
    image_paths, imagegen_errors = _generate_slide_images(
        root=root,
        session_id=debug_session_id,
        specs=specs,
        style=style,
        enable_imagegen=enable_imagegen,
        selected_template=selected_template,
        reserved_image_indices=set(source_image_map),
    )
    deck_path, notes_path, _render_warnings, _effective_template = _render_debug_pptx(
        root=root,
        user_id=(user_id or "debug").strip() or "debug",
        agent_id=agent.id,
        session_id=debug_session_id,
        user_message=ppt_message,
        assistant_answer=answer,
        style=style,
        image_paths=image_paths,
        imagegen_errors=imagegen_errors,
        selected_template=selected_template,
    )

    return PPTDebugResult(
        deck_path=deck_path,
        notes_path=notes_path,
        prompt=prompt,
        answer=answer,
        session_id=debug_session_id,
        style_id=style.style_id,
        style_label=style.label,
        style_created=style.created,
        image_paths=list(image_paths.values()),
        imagegen_errors=imagegen_errors,
        codex_error=codex_error,
        research_error=research_error,
        research_pack=research_pack,
    )
