from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from janus_lab.ppt_renderer import SlideSpec
from ppt_pipeline.content_parsing import *
from ppt_pipeline.image_routing import _slide_image_list
from ppt_pipeline.render_primitives import *
from ppt_pipeline.style_catalog import PPTStyleDecision, PPTTemplateDecision
from ppt_pipeline.template_core import *

def _render_template_chrome_image_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    image_path: Path,
    content_box: tuple[float, float, float, float],
    prepared_dir: Path,
    index: int,
) -> None:
    p = style.palette
    ex = _palette_extras(style)
    x, y, w, h = content_box
    left_w = min(4.55, w * 0.42)
    right_x = x + left_w + 0.35
    right_w = max(1.0, w - left_w - 0.35)
    bullets = _template_slide_bullets(spec, limit=3)
    boost = _academic_sparse_spec_boost(style, spec, bullets)
    card_h = min(1.15, (h - 0.25 * (len(bullets) - 1)) / max(len(bullets), 1))
    if _is_scut_academic_chrome(style):
        ref = _academic_ref_colors(style)
        if not _is_generated_slide_image(image_path):
            _add_line(slide, x, y + 0.08, x + w, y + 0.08, ref["line"], width=1.0)
            prepared = _prepare_image_for_ppt(image_path, prepared_dir, index, spec=spec)
            if prepared:
                figure_h = min(3.68 if boost >= 3 else 3.42 if boost else 3.18, h * (0.72 if boost >= 3 else 0.67 if boost else 0.62))
                ix, iy, iw, ih = _fit_image_to_box(prepared, x + 0.18, y + 0.28, w - 0.36, figure_h)
                picture = slide.shapes.add_picture(str(prepared), _ppt_len(ix), _ppt_len(iy), width=_ppt_len(iw), height=_ppt_len(ih))
                picture.name = "janus-template-content-image"
                _add_outline_rect(slide, ix, iy, iw, ih, ref["line"], width=0.8)
            note_y = y + min(4.08 if boost >= 3 else 3.92 if boost else 3.78, h - 1.05)
            if bullets:
                _add_academic_plain_list(
                    slide,
                    style,
                    bullets,
                    x + 0.1,
                    note_y,
                    w - 0.2,
                    max(0.9, y + h - note_y - 0.08),
                    start=1,
                    max_items=3,
                    max_size=11 if boost else 10,
                )
            return
        _add_academic_plain_list(
            slide,
            style,
            bullets,
            x,
            y + 0.22,
            left_w,
            min(h - 0.62, 2.7),
            start=1,
            max_items=3,
            max_size=11 if boost else 10,
        )
        _add_line(slide, right_x, y + 0.08, right_x + right_w, y + 0.08, ref["line"], width=1.0)
        prepared = _prepare_image_for_ppt(image_path, prepared_dir, index, spec=spec)
        if prepared:
            ix, iy, iw, ih = _fit_image_to_box(prepared, right_x + 0.08, y + 0.32, right_w - 0.16, h - 0.52)
            picture = slide.shapes.add_picture(str(prepared), _ppt_len(ix), _ppt_len(iy), width=_ppt_len(iw), height=_ppt_len(ih))
            picture.name = "janus-template-content-image"
            _add_outline_rect(slide, ix, iy, iw, ih, ref["line"], width=0.8)
            if _is_generated_slide_image(image_path):
                _add_image_overlay_labels(slide, style, spec, ix, iy, iw, ih, max_labels=4)
        return

    for i, point in enumerate(bullets):
        cy = y + i * (card_h + 0.25)
        _add_academic_ref_panel(slide, style, x, cy, left_w, card_h, fill=_academic_ref_fill(style, i), accent=p["accent"])
        _add_rect(slide, x, cy, 0.08, card_h, p["accent"])
        _add_autofit_text(
            slide,
            point,
            x + 0.24,
            cy + 0.08,
            left_w - 0.38,
            card_h - 0.12,
            max_size=13,
            min_size=8,
            color=p["ink"],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.04,
        )
    _add_academic_ref_panel(slide, style, right_x, y, right_w, h - 0.05, fill="FFFFFF")
    prepared = _prepare_image_for_ppt(image_path, prepared_dir, index, spec=spec)
    if prepared:
        ix, iy, iw, ih = _fit_image_to_box(prepared, right_x + 0.14, y + 0.14, right_w - 0.28, h - 0.34)
        picture = slide.shapes.add_picture(str(prepared), _ppt_len(ix), _ppt_len(iy), width=_ppt_len(iw), height=_ppt_len(ih))
        picture.name = "janus-template-content-image"
        if _is_generated_slide_image(image_path):
            _add_image_overlay_labels(slide, style, spec, ix, iy, iw, ih, max_labels=4)


def _render_template_chrome_claim_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    content_box: tuple[float, float, float, float],
) -> None:
    p = style.palette
    ex = _palette_extras(style)
    x, y, w, h = content_box
    points = _template_rich_bullets(spec, limit=4)
    boost = _academic_sparse_spec_boost(style, spec, points)
    claim = points[0] if points else (spec.title or "")
    left_w = min(4.9, max(4.25, w * 0.4)) if _is_scut_academic_chrome(style) else min(5.65, max(4.65, w * 0.46))
    gap = 0.36 if _is_scut_academic_chrome(style) else 0.42
    right_x = x + left_w + gap
    right_w = max(3.8, w - left_w - gap)
    top = y + 0.18
    usable_h = max(3.8, h - 0.28)

    if _is_scut_academic_chrome(style):
        ref = _academic_ref_colors(style)
        _add_line(slide, x, top + 0.1, x, top + usable_h - 0.1, p["accent"], width=2.0)
        _add_text(slide, _localized_label(spec, "核心判断", "KEY CLAIM"), x + 0.28, top + 0.12, 1.05, 0.26, size=9, color=p["accent"], bold=True, margin=0)
        _add_autofit_text(
            slide,
            claim,
            x + 0.28,
            top + 0.48,
            left_w - 0.5,
            min(1.62 if boost >= 3 else 1.5, usable_h * 0.38),
            max_size=21 + min(boost, 4),
            min_size=13 if boost >= 2 else 12,
            color=p["ink"],
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.02,
        )
        supporting = points[1:4]
        if supporting:
            _add_academic_plain_list(
                slide,
                style,
                supporting,
                x + 0.28,
                top + 1.92,
                left_w - 0.42,
                max(1.35, usable_h - 2.02),
                start=1,
                max_items=3,
                max_size=11 + min(boost, 3),
            )

        _add_line(slide, right_x, top + 0.08, right_x + right_w, top + 0.08, ref["line"], width=1.0)
        _add_text(slide, _localized_label(spec, "关系示意", "RELATION MAP"), right_x, top + 0.2, 1.05, 0.26, size=9, color=ref["title"], bold=True, margin=0)
        terms = _template_concept_terms(spec, limit=5)
        if not terms:
            terms = [_shorten_for_cell(spec.title or "核心关系", 12)]
        center_x = right_x + right_w * 0.5
        center_y = top + usable_h * 0.52
        center_label = terms[0]
        center_d = 1.52 if boost >= 3 else 1.46 if boost else 1.38
        _add_circle(slide, center_x - center_d / 2, center_y - center_d / 2, center_d, p["accent"], line=p["accent2"])
        _add_autofit_text(
            slide,
            center_label,
            center_x - 0.58,
            center_y - 0.22,
            1.16,
            0.44,
            max_size=12 + min(boost, 2),
            min_size=10 if boost >= 2 else 9,
            color=p["bg"],
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.01,
        )
        node_w = min(2.68 if boost else 2.46, right_w * (0.37 if boost else 0.34))
        node_h = 0.82 if boost >= 2 else 0.74
        node_positions = [
            (right_x + 0.36, top + usable_h * 0.23),
            (right_x + right_w - node_w - 0.24, top + usable_h * 0.17),
            (right_x + right_w - node_w - 0.36, top + usable_h * 0.68),
            (right_x + 0.54, top + usable_h * 0.73),
        ]
        colors = [p["accent2"], ex["warm"], ex["good"], ex["danger"]]
        for i, label in enumerate(terms[1:5]):
            nx, ny = node_positions[i]
            _add_line(slide, center_x, center_y, nx + node_w * 0.5, ny + node_h * 0.5, ref["line"], width=0.9)
            _add_rect(slide, nx, ny, node_w, node_h, "FFFFFF", line=colors[i], radius=False)
            _add_autofit_text(
                slide,
                label,
                nx + 0.12,
                ny + 0.1,
                node_w - 0.24,
                node_h - 0.18,
                max_size=12 + min(boost, 2),
                min_size=10 if boost >= 2 else 9,
                color=p["ink"],
                bold=True,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
                margin=0.01,
            )
        if len(terms) == 1:
            _add_academic_plain_note(
                slide,
                style,
                _shorten_for_cell(claim, 38),
                right_x + 0.28,
                top + usable_h - 0.78,
                right_w - 0.56,
                0.62,
                label="注",
            )
        return

    _add_academic_ref_panel(slide, style, x, top, left_w, usable_h, fill=_academic_ref_fill(style, 0), radius=True)
    _add_rect(slide, x, top, 0.09, usable_h, p["accent"])
    _add_text(slide, "01", x + 0.24, top + 0.22, 0.7, 0.28, size=11, color=p["accent"], bold=True, margin=0)
    _add_autofit_text(
        slide,
        claim,
        x + 0.58,
        top + 0.62,
        left_w - 0.9,
        min(1.55, usable_h * 0.36),
        max_size=17,
        min_size=10,
        color=p["ink"],
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.06,
    )
    supporting = points[1:4]
    if supporting:
        item_top = top + 2.35
        item_h = min(0.72, (top + usable_h - item_top - 0.18) / max(len(supporting), 1))
        for i, point in enumerate(supporting):
            iy = item_top + i * (item_h + 0.14)
            _add_circle(slide, x + 0.58, iy + 0.2, 0.18, [p["accent2"], ex["warm"], ex["good"]][i % 3])
            _add_autofit_text(
                slide,
                _academic_visible_text(style, point, 30),
                x + 0.92,
                iy,
                left_w - 1.18,
                item_h,
                max_size=11,
                min_size=8,
                color=p["muted"],
                anchor=MSO_ANCHOR.MIDDLE,
                margin=0.02,
            )

    _add_academic_ref_panel(slide, style, right_x, top, right_w, usable_h, fill="FFFFFF", radius=True)
    terms = _template_concept_terms(spec, limit=5)
    if not terms:
        terms = [_shorten_for_cell(spec.title or "核心关系", 12)]
    center_x = right_x + right_w * 0.5
    center_y = top + usable_h * 0.52
    center_label = terms[0]
    _add_circle(slide, center_x - 0.55, center_y - 0.55, 1.1, p["accent"], line=p["accent2"])
    _add_autofit_text(
        slide,
        center_label,
        center_x - 0.42,
        center_y - 0.18,
        0.84,
        0.36,
        max_size=9,
        min_size=8 if _is_scut_academic_chrome(style) else 6,
        color=p["bg"],
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )
    node_positions = [
        (right_x + right_w * 0.16, top + usable_h * 0.24),
        (right_x + right_w * 0.66, top + usable_h * 0.17),
        (right_x + right_w * 0.75, top + usable_h * 0.70),
        (right_x + right_w * 0.22, top + usable_h * 0.76),
    ]
    colors = [p["accent2"], ex["warm"], ex["good"], ex["danger"]]
    for i, label in enumerate(terms[1:5]):
        nx, ny = node_positions[i]
        _add_line(slide, center_x, center_y, nx + 0.46, ny + 0.28, ex["grid"], width=1.0)
        node_w = min(2.08 if _is_scut_academic_chrome(style) else 1.9, right_w * 0.3)
        node_h = 0.64 if _is_scut_academic_chrome(style) else 0.56
        _add_rect(slide, nx, ny, node_w, node_h, "FFFFFF", line=colors[i], radius=False)
        _add_autofit_text(
            slide,
            label,
            nx + 0.12,
            ny + 0.1,
            node_w - 0.24,
            node_h - 0.18,
            max_size=10 if _is_scut_academic_chrome(style) else 9,
            min_size=8 if _is_scut_academic_chrome(style) else 6,
            color=p["ink"],
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.01,
        )
    if len(terms) == 1:
        _add_academic_ref_panel(slide, style, right_x + 0.55, top + usable_h - 1.05, right_w - 1.1, 0.62, fill=_academic_ref_colors(style)["fill_blue"], radius=True)
        _add_autofit_text(
            slide,
            _shorten_for_cell(claim, 34),
            right_x + 0.75,
            top + usable_h - 0.92,
            right_w - 1.5,
            0.36,
            max_size=10,
            min_size=8 if _is_scut_academic_chrome(style) else 7,
            color=p["muted"],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.02,
        )


def _render_template_chrome_claim_band_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    content_box: tuple[float, float, float, float],
) -> None:
    if not _is_scut_academic_chrome(style):
        _render_template_chrome_claim_layout(slide, spec, style, content_box)
        return
    p = style.palette
    ex = _palette_extras(style)
    ref = _academic_ref_colors(style)
    x, y, w, h = content_box
    points = _template_rich_bullets(spec, limit=5)
    boost = _academic_sparse_spec_boost(style, spec, points)
    claim = _academic_visible_text(style, points[0] if points else (spec.title or ""), 48)
    top = y + 0.18
    _add_line(slide, x, top + 0.08, x + w, top + 0.08, ref["line"], width=1.0)
    _add_rect(slide, x, top + 0.34, 0.13, 1.24, p["accent"])
    _add_text(slide, "KEY CLAIM", x + 0.32, top + 0.28, 1.25, 0.24, size=9, color=p["accent"], bold=True, margin=0)
    _add_autofit_text(
        slide,
        claim,
        x + 0.32,
        top + 0.64,
        w - 0.72,
        1.18 if boost < 3 else 1.34,
        max_size=25 + min(boost, 3),
        min_size=14,
        color=p["ink"],
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )

    evidence = points[1:5]
    evidence = evidence[:4]
    band_top = top + (2.16 if boost < 3 else 2.34)
    band_h = min(1.72, max(1.16, y + h - band_top - 0.38))
    if evidence and band_h > 0.8:
        gap = 0.18
        item_w = (w - gap * (len(evidence) - 1)) / max(len(evidence), 1)
        accents = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
        for i, item in enumerate(evidence):
            cx = x + i * (item_w + gap)
            _add_line(slide, cx, band_top, cx + item_w, band_top, accents[i % len(accents)], width=1.4)
            _add_template_icon(slide, _template_icon_kind(item), cx + 0.02, band_top + 0.26, 0.46, accents[i % len(accents)])
            _add_autofit_text(
                slide,
                _academic_visible_text(style, item, 30),
                cx + 0.62,
                band_top + 0.16,
                item_w - 0.68,
                band_h - 0.2,
                max_size=12 + min(boost, 2),
                min_size=9,
                color=p["ink"],
                anchor=MSO_ANCHOR.MIDDLE,
                margin=0.01,
            )
    note = (points[:1] or [None])[0]
    if note and y + h - 0.72 > band_top + band_h:
        _add_academic_plain_note(slide, style, note, x, y + h - 0.68, w, 0.58, label=_localized_label(spec, "讲解", "NOTE"))


def _sparse_support_points(spec: SlideSpec, primary: str, *, limit: int = 4) -> list[str]:
    points: list[str] = []
    def usable(value: str) -> bool:
        value = _clean_visible_text(value)
        if not value or _is_layout_instruction_text(value) or _is_reference_fragment_text(value):
            return False
        if re.fullmatch(r"[A-Za-z][A-Za-z0-9'_^\-]{0,4}", value):
            return False
        if value in {"叠加为", "用于", "作为", "展示", "说明", "强调"}:
            return False
        return _visual_len(value) >= 4

    for candidate in _template_rich_bullets(spec, limit=limit + 2):
        if usable(candidate) and candidate != primary and candidate not in points:
            points.append(candidate)
    for term in _template_concept_terms(spec, limit=limit + 2):
        if usable(term) and term not in primary and term not in points:
            points.append(term)
        if len(points) >= limit:
            break
    return points[:limit]


def _render_template_chrome_statement_sidebar_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    content_box: tuple[float, float, float, float],
) -> None:
    if not _is_scut_academic_chrome(style):
        _render_template_chrome_claim_layout(slide, spec, style, content_box)
        return
    p = style.palette
    ex = _palette_extras(style)
    ref = _academic_ref_colors(style)
    x, y, w, h = content_box
    points = _template_rich_bullets(spec, limit=5)
    primary = _academic_visible_text(style, points[0] if points else (spec.title or ""), 52)
    support = _sparse_support_points(spec, primary, limit=3)
    boost = _academic_sparse_spec_boost(style, spec, points)
    left_w = min(6.3, w * 0.56)
    right_x = x + left_w + 0.46
    right_w = w - left_w - 0.46
    top = y + 0.24
    _add_line(slide, x, top + 0.05, x + left_w, top + 0.05, ref["line"], width=1.0)
    _add_rect(slide, x, top + 0.34, 0.12, 1.48, p["accent"])
    _add_text(slide, _localized_label(spec, "核心判断", "KEY CLAIM"), x + 0.32, top + 0.26, 1.1, 0.28, size=10, color=p["accent"], bold=True, margin=0)
    _add_autofit_text(
        slide,
        primary,
        x + 0.32,
        top + 0.72,
        left_w - 0.48,
        min(2.05 if boost >= 2 else 1.82, h - 1.24),
        max_size=25 + min(boost, 4),
        min_size=14,
        color=p["ink"],
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )
    _add_line(slide, right_x, top + 0.05, right_x + right_w, top + 0.05, ref["line"], width=1.0)
    _add_text(slide, "支撑信息", right_x, top + 0.26, 1.1, 0.28, size=10, color=p["ink"], bold=True, margin=0)
    row_h = min(0.82 if boost < 3 else 0.96, (h - 0.92) / max(len(support), 1))
    colors = [p["accent"], p["accent2"], ex["warm"]]
    for i, item in enumerate(support or [primary]):
        cy = top + 0.72 + i * row_h
        _add_template_icon(slide, _template_icon_kind(item), right_x, cy + 0.08, 0.42, colors[i % len(colors)])
        _add_line(slide, right_x + 0.56, cy + row_h - 0.02, right_x + right_w, cy + row_h - 0.02, ref["line"], width=0.7)
        _add_autofit_text(
            slide,
            _academic_visible_text(style, item, 34),
            right_x + 0.68,
            cy + 0.02,
            right_w - 0.72,
            row_h - 0.06,
            max_size=12 + min(boost, 2),
            min_size=9,
            color=p["ink"],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.01,
        )


def _render_template_chrome_quote_rule_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    content_box: tuple[float, float, float, float],
) -> None:
    if not _is_scut_academic_chrome(style):
        _render_template_chrome_claim_layout(slide, spec, style, content_box)
        return
    p = style.palette
    ex = _palette_extras(style)
    ref = _academic_ref_colors(style)
    x, y, w, h = content_box
    points = _template_rich_bullets(spec, limit=5)
    primary = _academic_visible_text(style, points[0] if points else (spec.title or ""), 58)
    support = _sparse_support_points(spec, primary, limit=4)
    boost = _academic_sparse_spec_boost(style, spec, points)
    top = y + 0.34
    _add_text(slide, "研究定位", x, top + 0.02, 1.1, 0.26, size=9, color=p["accent"], bold=True, margin=0)
    _add_line(slide, x + 1.12, top + 0.16, x + w, top + 0.16, ref["line"], width=1.0)
    _add_autofit_text(
        slide,
        primary,
        x + 0.12,
        top + 0.66,
        w - 0.24,
        min(1.62 if boost < 3 else 1.86, h - 1.72),
        max_size=27 + min(boost, 4),
        min_size=15,
        color=p["ink"],
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )
    bottom_y = y + h - 1.36
    _add_line(slide, x + 0.3, bottom_y, x + w - 0.3, bottom_y, ref["line"], width=0.9)
    chip_count = min(len(support), 4)
    if chip_count:
        gap = 0.16
        chip_w = (w - 0.6 - gap * (chip_count - 1)) / chip_count
        colors = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
        for i, item in enumerate(support[:chip_count]):
            cx = x + 0.3 + i * (chip_w + gap)
            _add_rect(slide, cx, bottom_y + 0.22, 0.08, 0.54, colors[i % len(colors)])
            _add_autofit_text(
                slide,
                _academic_visible_text(style, item, 24),
                cx + 0.18,
                bottom_y + 0.14,
                chip_w - 0.22,
                0.72,
                max_size=11 + min(boost, 2),
                min_size=8,
                color=p["ink"],
                anchor=MSO_ANCHOR.MIDDLE,
                margin=0.01,
            )


def _render_template_chrome_compact_mosaic_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    content_box: tuple[float, float, float, float],
) -> None:
    if not _is_scut_academic_chrome(style):
        _render_template_chrome_cards_layout(slide, spec, style, content_box)
        return
    p = style.palette
    ex = _palette_extras(style)
    ref = _academic_ref_colors(style)
    x, y, w, h = content_box
    points = _template_rich_bullets(spec, limit=5)
    primary = _academic_visible_text(style, points[0] if points else (spec.title or ""), 40)
    support = _sparse_support_points(spec, primary, limit=4)
    if not support:
        _render_template_chrome_statement_sidebar_layout(slide, spec, style, content_box)
        return
    boost = _academic_sparse_spec_boost(style, spec, points)
    top = y + 0.24
    left_w = min(4.65, w * 0.38)
    grid_x = x + left_w + 0.42
    grid_w = w - left_w - 0.42
    _add_line(slide, x, top + 0.06, x + left_w, top + 0.06, ref["line"], width=1.0)
    _add_text(slide, "主结论", x, top + 0.24, 0.82, 0.26, size=9, color=p["accent"], bold=True, margin=0)
    _add_autofit_text(
        slide,
        primary,
        x,
        top + 0.72,
        left_w,
        min(2.28, h - 1.12),
        max_size=21 + min(boost, 4),
        min_size=13,
        color=p["ink"],
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.01,
    )
    _add_line(slide, grid_x, top + 0.06, grid_x + grid_w, top + 0.06, ref["line"], width=1.0)
    item_count = min(len(support), 4)
    cols = 1 if item_count == 1 else 2
    rows = 1 if item_count <= 2 else 2
    gap_x = 0.28
    gap_y = 0.24
    cell_w = (grid_w - gap_x * (cols - 1)) / cols
    cell_h = min(1.4 if rows == 1 else (1.36 if boost >= 2 else 1.18), (h - 0.8 - gap_y * (rows - 1)) / rows)
    colors = [p["accent"], p["accent2"], ex["warm"], ex["good"]]
    for i, item in enumerate(support[:item_count]):
        col = i % cols
        row = i // cols
        cx = grid_x + col * (cell_w + gap_x)
        cy = top + 0.58 + row * (cell_h + gap_y)
        _add_line(slide, cx, cy, cx + cell_w, cy, colors[i % len(colors)], width=1.3)
        _add_rect(slide, cx + 0.08, cy + 0.18, 0.08, max(0.3, cell_h - 0.36), colors[i % len(colors)])
        _add_autofit_text(
            slide,
            _academic_visible_text(style, item, 28),
            cx + 0.3,
            cy + 0.08,
            cell_w - 0.36,
            cell_h - 0.12,
            max_size=12 + min(boost, 2),
            min_size=8,
            color=p["ink"],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.01,
        )


def _render_template_chrome_sparse_variant_layout(
    slide,
    spec: SlideSpec,
    style: PPTStyleDecision,
    content_box: tuple[float, float, float, float],
    *,
    index: int = 0,
) -> None:
    if not _is_scut_academic_chrome(style):
        _render_template_chrome_claim_layout(slide, spec, style, content_box)
        return
    points = _template_rich_bullets(spec, limit=5)
    raw = f"{spec.title or ''} {spec.message or ''} {spec.visual or ''}"
    numbery = sum(1 for point in points if _extract_highlight(point))
    if numbery >= 1 and any(key in raw for key in ["结果", "实验", "指标", "评估", "性能", "提升", "下降", "Recall", "NDCG", "CTR", "CVR"]):
        _render_template_chrome_metric_focus_layout(slide, spec, style, content_box)
        return
    variant = index % 5
    if variant == 0:
        _render_template_chrome_statement_sidebar_layout(slide, spec, style, content_box)
    elif variant == 1:
        _render_template_chrome_quote_rule_layout(slide, spec, style, content_box)
    elif variant == 2:
        _render_template_chrome_compact_mosaic_layout(slide, spec, style, content_box)
    elif variant == 3:
        _render_template_chrome_metric_focus_layout(slide, spec, style, content_box)
    else:
        _render_template_chrome_claim_band_layout(slide, spec, style, content_box)


def _hitsz_visible_text(text: str, limit: int = 42) -> str:
    return _shorten_for_cell(_clean_visible_text(text), limit)


def _hitsz_points(spec: SlideSpec, *, limit: int = 5) -> list[str]:
    points = _template_rich_bullets(spec, limit=limit)
    if points:
        return points[:limit]
    return [spec.title] if spec.title else []


def _hitsz_support_items(spec: SlideSpec, primary: str, points: list[str], *, limit: int = 4) -> list[str]:
    banned = {"研究对象", "方法模块", "验证证据", "结果含义", "问题定义", "方法抓手", "后续展开", "核心内容"}
    items: list[str] = []

    def add(candidate: str) -> None:
        value = _clean_visible_text(candidate)
        if (
            not value
            or value == primary
            or value in banned
            or value in items
            or _is_layout_instruction_text(value)
            or _is_reference_fragment_text(value)
        ):
            return
        items.append(value)

    for point in points[1:]:
        add(point)
        if len(items) >= limit:
            return items[:limit]
    for point in _sparse_support_points(spec, primary, limit=limit):
        add(point)
        if len(items) >= limit:
            break
    return items[:limit]



__all__ = [
    "_render_template_chrome_image_layout",
    "_render_template_chrome_claim_layout",
    "_render_template_chrome_claim_band_layout",
    "_sparse_support_points",
    "_render_template_chrome_statement_sidebar_layout",
    "_render_template_chrome_quote_rule_layout",
    "_render_template_chrome_compact_mosaic_layout",
    "_render_template_chrome_sparse_variant_layout",
    "_hitsz_visible_text",
    "_hitsz_points",
    "_hitsz_support_items",
]
