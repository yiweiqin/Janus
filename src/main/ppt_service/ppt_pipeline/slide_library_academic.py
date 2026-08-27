from __future__ import annotations

import math
import re
from typing import Any

from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ppt_pipeline.content_payloads import *
from ppt_pipeline.slide_library_core import *
from ppt_pipeline.slide_library_policy import *

def _bind_basic_content(slide, payload: dict[str, Any]) -> None:
    story, _has_lead = _basic_story(payload)
    _set_named(slide, "basic-text-card", story, max_size=19)
    raw_caption = payload.get("caption")
    if raw_caption or payload.get("source_note") or payload.get("source"):
        caption_item = raw_caption if raw_caption else {"source": payload.get("source_note") or payload.get("source")}
        caption = _figure_caption(payload, caption_item, 0)
        _set_named(slide, "basic-caption", caption, max_size=11)
        _style_figure_caption(_shape_named(slide, "basic-caption"))
    else:
        _remove_named(slide, "basic-caption")
    _set_named(slide, "basic-image", "", max_size=14)
    _remove_crosses(slide, "basic-image")
    _remove_prefix(slide, "tpl-kicker-label")


def _configure_basic_content_layout(
    slide,
    payload: dict[str, Any],
    image_path: Path | None,
    *,
    mirror: bool,
) -> bool:
    text = _shape_named(slide, "basic-text-card")
    image = _shape_named(slide, "basic-image")
    caption = _shape_named(slide, "basic-caption")
    if image_path is None:
        if _basic_text_is_sparse(text, narrow=False):
            _set_geometry(text, 0.95, 1.78, 11.43, 3.78)
        else:
            _set_geometry(text, 0.75, 1.45, 11.83, 5.1)
        _configure_basic_text_density(text, payload, narrow=False)
        _remove_slot(slide, "basic-image")
        _remove_named(slide, "basic-caption")
        return False
    information_score, aspect = _image_information_score(image_path)
    if _image_visual_mode(image_path) == "contain" and information_score >= 0.55 and aspect >= 1.35:
        # Dense infographics already carry the explanatory load. Give them the
        # entire body canvas and move supporting prose into PowerPoint notes.
        _remove_named(slide, "basic-text-card")
        _set_geometry(image, 0.75, 1.2, 11.83, 5.35)
        _remove_named(slide, "basic-caption")
        return True
    # Human-made reference decks consistently let one strong image carry the
    # page. Give the visual roughly two thirds of the body canvas.
    image_x, text_x = (0.75, 9.10) if mirror else (4.43, 0.75)
    _set_geometry(text, text_x, 1.45, 3.38, 5.1)
    _configure_basic_text_density(text, payload, narrow=True)
    has_caption = bool(payload.get("caption") or payload.get("source_note") or payload.get("source"))
    image_height = 4.25 if has_caption else 5.1
    _set_geometry(image, image_x, 1.45, 8.15, image_height)
    if caption is not None:
        _set_geometry(caption, image_x, 5.84, 8.15, 0.72)
    return False


def _motivation_content(payload: dict[str, Any], value: Any, *, target: bool) -> str:
    heading = _localized(payload, "目标状态" if target else "当前状态", "TARGET STATE" if target else "CURRENT STATE")
    if isinstance(value, dict):
        label = _display(value.get("label") or value.get("title"))
        metric = _display(value.get("value") or value.get("metric") or value.get("number"))
        detail = _display(value.get("points") or value.get("detail") or value.get("description") or value.get("content"))
        return "\n".join(part for part in [heading, label, metric, detail] if part)
    content = _display(value)
    return "\n".join(part for part in [heading, content] if part)


def _bind_motivation(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 8)
    kicker = _shape_named(slide, "tpl-kicker-label")
    if kicker is not None:
        _set_geometry(kicker, 0.76, 1.18, 2.15, 0.38)
        kicker.fill.background()
        kicker.line.fill.background()
        kicker.text_frame.margin_left = 0
        kicker.text_frame.margin_right = 0
        kicker.text_frame.margin_top = 0
        kicker.text_frame.margin_bottom = 0
        kicker.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set_shape_text(kicker, "WHY NOW?", max_size=15, bold=True)
        for paragraph in kicker.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
    current_value = payload.get("current")
    target_value = payload.get("target")
    current = _motivation_content(payload, current_value or "\n".join(points[:3]), target=False)
    target = _motivation_content(payload, target_value or "\n".join(points[3:6] or points[:3]), target=True)
    current_shape = _shape_named(slide, "motivation-current")
    target_shape = _shape_named(slide, "motivation-target")
    gap_shape = _shape_named(slide, "motivation-gap")
    _set_shape_text(current_shape, current, max_size=16)
    _set_shape_text(target_shape, target, max_size=16)
    gap_value = _display(payload.get("gap")) or (points[6] if len(points) > 6 else _localized(payload, "尚未闭合的关键差距", "The critical gap to close"))
    _set_shape_text(
        gap_shape,
        "\n".join([_localized(payload, "关键差距", "CRITICAL GAP"), gap_value]),
        max_size=15,
        bold=True,
    )
    row = _structured_list(payload, "row")
    card_height = 3.45 if row else 4.78
    _set_geometry(current_shape, 0.75, 1.72, 4.10, card_height)
    _set_geometry(target_shape, 8.48, 1.72, 4.10, card_height)
    _set_geometry(gap_shape, 5.24, 2.55 if not row else 2.30, 2.85, 1.85)
    for shape in (current_shape, target_shape):
        if shape is not None:
            shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
            shape.text_frame.margin_left = Inches(0.30)
            shape.text_frame.margin_right = Inches(0.26)
            shape.text_frame.margin_top = Inches(0.26)
        _set_text_spacing(shape, after=12, line_spacing=1.12)
        _style_first_paragraph(shape, color=SEMANTIC_DARK)
    _set_text_spacing(gap_shape, after=5, line_spacing=1.05, middle=True)
    _set_shape_colors(current_shape, fill=SEMANTIC_MUTED_LIGHT, line=RGBColor(218, 227, 233))
    _set_shape_colors(target_shape, fill=SEMANTIC_PRIMARY_LIGHT, line=RGBColor(192, 221, 222))
    _set_shape_colors(gap_shape, fill=SEMANTIC_WARM_LIGHT, line=SEMANTIC_WARM)
    arrow = _shape_named(slide, "motivation-arrow")
    if arrow is not None and current_shape is not None and target_shape is not None:
        center_y = float(current_shape.top + current_shape.height / 2)
        _set_line_endpoints(
            arrow,
            (float(current_shape.left + current_shape.width), center_y),
            (float(target_shape.left), center_y),
        )
        try:
            arrow.line.color.rgb = SEMANTIC_PRIMARY
            arrow.line.width = Pt(1.4)
        except Exception:
            pass
    if row:
        for col, value in enumerate(row, 1):
            _set_named(slide, f"motivation-mini-table-r2c{col}", _display(value), max_size=11)
            _set_geometry(_shape_named(slide, f"motivation-mini-table-r1c{col}"), 0.75 + (col - 1) * 2.91, 5.55, 2.91, 0.45)
            _set_geometry(_shape_named(slide, f"motivation-mini-table-r2c{col}"), 0.75 + (col - 1) * 2.91, 6.00, 2.91, 0.56)
        _style_minimal_grid(slide, "motivation-mini-table", columns=4, data_rows=1)
    else:
        _remove_prefix(slide, "motivation-mini-table-")


def _bind_challenge(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 6)
    structured = _structured_list(payload, "nodes") or _structured_list(payload, "risks")
    cards = [_card_text(item) for item in structured if _card_text(item)] or points[:4]
    cards = cards[:4]
    card_shapes = _shapes_named(slide, "challenge-card")
    for index, shape in enumerate(card_shapes[:len(cards)]):
        _set_shape_text(shape, cards[index], max_size=15)
        _set_text_spacing(shape, after=8, line_spacing=1.08)
    _remove_occurrences_after(slide, "challenge-card", len(cards))
    _remove_occurrences_after(slide, "challenge-link", len(cards))
    core_shape = _shape_named(slide, "challenge-core")
    if len(cards) == 1 and card_shapes:
        _set_geometry(card_shapes[0], 1.05, 1.62, 7.45, 4.82)
        _set_geometry(core_shape, 8.88, 2.42, 3.35, 2.05)
        _set_geometry(_shape_named(slide, "challenge-note"), 8.88, 4.78, 3.35, 0.85)
    elif len(cards) == 2:
        _set_geometry(card_shapes[0], 0.75, 1.72, 4.65, 4.65)
        _set_geometry(card_shapes[1], 7.93, 1.72, 4.65, 4.65)
        _set_geometry(core_shape, 5.61, 2.75, 2.11, 2.18)
    elif len(cards) == 3:
        positions = ((0.75, 1.68), (4.87, 1.68), (8.99, 1.68))
        for shape, (x, y) in zip(card_shapes[:3], positions):
            _set_geometry(shape, x, y, 3.59, 3.72)
        _set_geometry(core_shape, 2.15, 5.72, 9.03, 0.78)
    elif len(cards) >= 4:
        positions = ((0.75, 1.50), (8.08, 1.50), (0.75, 4.82), (8.08, 4.82))
        for shape, (x, y) in zip(card_shapes[:4], positions):
            _set_geometry(shape, x, y, 4.50, 1.68)
        _set_geometry(core_shape, 4.77, 2.93, 3.79, 1.62)
    core = _display(payload.get("core") or payload.get("objective") or payload.get("conclusion"))
    _set_named(slide, "challenge-core", core or (points[0] if points else _localized(payload, "核心挑战", "Core challenge")), max_size=15)
    note = _display(payload.get("note"))
    if note:
        _set_named(slide, "challenge-note", note, max_size=11)
    else:
        _remove_named(slide, "challenge-note")
    # The template connectors originally run from center to center and sit
    # above the corner cards, so their inner halves can cross visible text.
    # Re-route after all adaptive geometry changes and stop at both borders.
    _clip_lines_to_shape_borders(
        _shapes_named(slide, "challenge-link"),
        _shapes_named(slide, "challenge-card"),
        _shape_named(slide, "challenge-core"),
    )


def _bind_pipeline(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 10)
    # Pipeline nodes should remain scannable labels. Long explanations belong
    # in the supporting row/notes; placing full paragraphs inside five narrow
    # nodes forces PowerPoint to shrink the visible font dramatically.
    steps = [
        _item_parts(item)[0] or _card_body(item)
        for item in _structured_list(payload, "steps")
        if _item_parts(item)[0] or _card_body(item)
    ] or points[1:6]
    steps = steps[:5]
    input_value = payload.get("input") or (points[0] if points else "输入")
    _set_named(slide, "pipeline-input", _compact_pipeline_endpoint(payload, input_value, output=False), max_size=15)
    _remove_crosses(slide, "pipeline-input")
    row = _structured_list(payload, "row")
    has_support_row = bool(row)
    count = max(1, len(steps))
    left, right, gap = 2.85, 10.48, 0.24
    step_width = min(2.25, (right - left - gap * (count - 1)) / count)
    total_width = step_width * count + gap * (count - 1)
    start = left + (right - left - total_width) / 2
    step_y = 1.68 if has_support_row else 1.62
    step_height = 3.18 if has_support_row else 4.55
    for index in range(1, 6):
        value = steps[index - 1] if index - 1 < len(steps) else ""
        step_shape = _shape_named(slide, f"pipeline-step-{index}")
        if value and step_shape is not None:
            _set_geometry(step_shape, start + (index - 1) * (step_width + gap), step_y, step_width, step_height)
            _set_shape_text(step_shape, f"{index:02d}\n{value}", max_size=15)
        elif step_shape is not None:
            _remove_shape(step_shape)
        arrow = _shape_named(slide, f"pipeline-arrow-{index}")
        if arrow is not None and index < count:
            arrow.left = Inches(start + index * step_width + (index - 1) * gap + 0.04)
            arrow.top = Inches(step_y + step_height / 2)
            arrow.width = Inches(max(0.10, gap - 0.08))
            arrow.height = Inches(0.02)
        elif arrow is not None:
            _remove_shape(arrow)
    output_shape = _shape_named(slide, "pipeline-output")
    if output_shape is not None:
        _set_geometry(output_shape, 10.78, 2.55 if has_support_row else 2.72, 1.82, 1.52)
        _style_pipeline_endpoint(output_shape)
    input_shape = _shape_named(slide, "pipeline-input")
    if input_shape is not None:
        _set_geometry(input_shape, 0.73, 2.55 if has_support_row else 2.72, 1.82, 1.52)
        _style_pipeline_endpoint(input_shape)
    output_value = payload.get("output") or (points[-1] if points else "输出")
    _set_named(slide, "pipeline-output", _compact_pipeline_endpoint(payload, output_value, output=True), max_size=15)
    _remove_crosses(slide, "pipeline-output")
    if has_support_row:
        headers = (
            ["模块", "输入", "操作", "输出", "检查"]
            if _is_cjk_payload(payload)
            else ["Module", "Input", "Operation", "Output", "Check"]
        )
        for col, value in enumerate(headers, 1):
            _set_named(slide, f"pipeline-grid-r1c{col}", value, max_size=11)
            _set_named(slide, f"pipeline-grid-r2c{col}", _display(row[col - 1]) if col - 1 < len(row) else "", max_size=11)
            _set_geometry(_shape_named(slide, f"pipeline-grid-r1c{col}"), 2.85 + (col - 1) * 1.526, 5.18, 1.526, 0.55)
            _set_geometry(_shape_named(slide, f"pipeline-grid-r2c{col}"), 2.85 + (col - 1) * 1.526, 5.73, 1.526, 0.75)
    else:
        _remove_prefix(slide, "pipeline-grid-")


def _bind_loop(slide, payload: dict[str, Any]) -> None:
    points = _payload_points(payload, 6)
    center = _shape_named(slide, "loop-center")
    _set_geometry(center, 4.86, 2.79, 3.61, 1.52)
    _set_shape_text(center, _display(payload.get("objective")) or (points[0] if points else "目标"), max_size=18, bold=True)
    _set_text_spacing(center, after=2, line_spacing=1.0, middle=True)
    nodes = [
        _item_parts(item)[0] or _card_text(item)
        for item in _structured_list(payload, "steps")
        if _item_parts(item)[0] or _card_text(item)
    ] or points[1:5]
    nodes = nodes[:4]
    names = ["Plan", "Act", "Reflect", "Observe"]
    if len(nodes) <= 3:
        positions = ((5.37, 1.45), (8.73, 4.43), (2.20, 4.43))
    else:
        positions = ((5.37, 1.43), (9.42, 3.02), (5.37, 4.75), (1.32, 3.02))
    active_shapes: list[Any] = []
    for index, name in enumerate(names):
        shape = _shape_named(slide, f"loop-node-{name}")
        if index < len(nodes):
            x, y = positions[index]
            _set_geometry(shape, x, y, 2.59, 1.00)
            _set_shape_text(shape, nodes[index], max_size=15, bold=True)
            _set_text_spacing(shape, after=0, line_spacing=1.0, middle=True)
            active_shapes.append(shape)
        elif shape is not None:
            _remove_shape(shape)
    arrows = [_shape_named(slide, f"loop-a{index}") for index in range(1, 5)]
    for index, arrow in enumerate(arrows):
        if index >= len(active_shapes):
            _remove_shape(arrow)
            continue
        source = active_shapes[index]
        target = active_shapes[(index + 1) % len(active_shapes)]
        _set_line_endpoints(
            arrow,
            _rectangle_edge_point(source, target),
            _rectangle_edge_point(target, source),
        )
    guardrail = _shape_named(slide, "loop-guardrail")
    _set_geometry(guardrail, 0.85, 5.97, 11.63, 0.62)
    _set_shape_text(
        guardrail,
        _display(payload.get("guardrail")) or (points[-1] if points else "评价与停止条件"),
        max_size=13,
    )
    _set_text_spacing(guardrail, after=2, line_spacing=1.0, middle=True)


def _kpis(payload: dict[str, Any], count: int = 3) -> list[dict[str, Any]]:
    raw = _structured_list(payload, "kpis")
    result: list[dict[str, Any]] = []
    for item in raw:
        if isinstance(item, dict):
            result.append({
                "value": _display(item.get("value") or item.get("group") or item.get("title")),
                "label": _display(item.get("label") or item.get("metrics")),
                "note": _display(item.get("note") or item.get("description")),
                "primary": item.get("primary") or item.get("highlight") or item.get("best") or item.get("is_primary"),
                "status": _display(item.get("status") or item.get("state")),
                "direction": _display(item.get("direction")),
            })
        else:
            result.append({"value": _display(item), "label": "关键指标", "note": "", "primary": False})
    numbers = _payload_numbers(payload, count)
    points = _payload_points(payload, count * 2)
    while len(result) < count:
        idx = len(result)
        result.append({
            "value": numbers[idx] if idx < len(numbers) else str(idx + 1),
            "label": points[idx] if idx < len(points) else "关键指标",
            "note": points[idx + count] if idx + count < len(points) else "",
            "primary": False,
        })
    return result[:count]


def _fill_kpis(slide, prefix: str, payload: dict[str, Any]) -> None:
    items = _kpis(payload)
    explicit_count = min(3, len(_structured_list(payload, "kpis")))
    count = explicit_count or len(items)
    count = max(1, min(3, count))
    for index, item in enumerate(items[:count], 1):
        value_size = 20 if not re.search(r"\d", item["value"]) else (32 if prefix == "big" else 28)
        _set_named(slide, f"{prefix}-kpi{index}-value", item["value"], max_size=value_size, bold=True)
        _set_named(slide, f"{prefix}-kpi{index}-label", item["label"], max_size=14)
        _set_named(slide, f"{prefix}-kpi{index}-note", item["note"], max_size=12)
    for index in range(count + 1, 4):
        for suffix in ("", "-accent", "-value", "-label", "-note"):
            _remove_named(slide, f"{prefix}-kpi{index}{suffix}")
    bases = [_shape_named(slide, f"{prefix}-kpi{index}") for index in range(1, count + 1)]
    bases = [shape for shape in bases if shape is not None]
    if bases:
        gap = 0.28
        available_left = 0.78
        available_right = 12.25 if prefix == "big" else 8.55
        max_width = 3.68 if prefix == "big" else 3.4
        width = min(max_width, (available_right - available_left - gap * (count - 1)) / count)
        total = width * count + gap * (count - 1)
        start = available_left + (available_right - available_left - total) / 2
        for index in range(1, count + 1):
            x = start + (index - 1) * (width + gap)
            for suffix in ("", "-accent"):
                shape = _shape_named(slide, f"{prefix}-kpi{index}{suffix}")
                if shape is not None:
                    shape.left = Inches(x)
                    if suffix == "":
                        shape.width = Inches(width)
            for suffix in ("-value", "-label", "-note"):
                shape = _shape_named(slide, f"{prefix}-kpi{index}{suffix}")
                if shape is not None:
                    shape.left = Inches(x + 0.2)
                    shape.width = Inches(max(0.7, width - 0.32))
            value_shape = _shape_named(slide, f"{prefix}-kpi{index}-value")
            label_shape = _shape_named(slide, f"{prefix}-kpi{index}-label")
            note_shape = _shape_named(slide, f"{prefix}-kpi{index}-note")
            if value_shape is not None:
                value_size = 20 if not re.search(r"\d", items[index - 1]["value"]) else (32 if prefix == "big" else 28)
                _refit_shape_text(value_shape, max_size=value_size, min_size=14)
            if label_shape is not None:
                _refit_shape_text(label_shape, max_size=15, min_size=12)
            if note_shape is not None:
                _refit_shape_text(note_shape, max_size=12, min_size=10)


def _rows(payload: dict[str, Any]) -> list[list[str]]:
    table = payload.get("table") if isinstance(payload.get("table"), dict) else {}
    raw = table.get("rows") if isinstance(table.get("rows"), list) else payload.get("rows")
    result: list[list[str]] = []
    if isinstance(raw, list):
        for row in raw:
            if isinstance(row, dict):
                result.append([_display(value) for value in row.values()])
            elif isinstance(row, list):
                result.append([_display(value) for value in row])
            else:
                result.append([_display(row)])
    return result


def _target_row(payload: dict[str, Any]) -> list[str]:
    raw = payload.get("row")
    if isinstance(raw, dict):
        return [
            _display(raw.get("output") or raw.get("deliverable")),
            _display(raw.get("indicator") or raw.get("metric")),
            _display(raw.get("acceptance") or raw.get("criterion")),
            _display(raw.get("owner") or raw.get("responsible")),
        ]
    if isinstance(raw, list):
        return [_display(item) for item in raw[:4]]
    return []


def _evaluation_status_table(payload: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    raw = payload.get("status_rows")
    if not isinstance(raw, list):
        return [], []
    if raw and all(isinstance(item, dict) for item in raw):
        records: list[dict[str, str]] = []
        for item in raw:
            records.append({
                "gate": _display(item.get("item") or item.get("label") or item.get("dimension") or item.get("gate") or item.get("title")),
                "status": _display(item.get("status") or item.get("state")),
                "evidence": _display(item.get("evidence") or item.get("detail") or item.get("signals") or item.get("requirements")),
                "purpose": _display(item.get("purpose") or item.get("meaning") or item.get("outcome") or item.get("use")),
            })
        has_status = any(record["status"] for record in records)
        if has_status:
            headers = (
                ["证据项", "状态", "所需证据"]
                if _is_cjk_payload(payload)
                else ["Evidence gate", "Status", "Required evidence"]
            )
            rows = [
                [record["gate"], record["status"], record["evidence"] or record["purpose"]]
                for record in records
            ]
        else:
            headers = (
                ["证据项", "所需信号", "评测用途"]
                if _is_cjk_payload(payload)
                else ["Evidence gate", "Required signals", "Purpose"]
            )
            rows = [
                [record["gate"], record["evidence"], record["purpose"]]
                for record in records
            ]
        return headers, rows
    return _headers(payload), _rows({"rows": raw})


def _headers(payload: dict[str, Any]) -> list[str]:
    table = payload.get("table") if isinstance(payload.get("table"), dict) else {}
    raw = table.get("headers") if isinstance(table.get("headers"), list) else payload.get("headers")
    return [_display(item) for item in raw] if isinstance(raw, list) else []


def _fill_grid(slide, stem: str, payload: dict[str, Any], columns: int, max_rows: int, default_headers: list[str]) -> None:
    headers = _headers(payload) or default_headers
    rows = _rows(payload)
    if not rows:
        points = _payload_points(payload, max_rows)
        numbers = _payload_numbers(payload, max_rows)
        rows = [[point, numbers[index] if index < len(numbers) else "", "", "", "", ""] for index, point in enumerate(points[:max_rows])]
    for col in range(1, columns + 1):
        _set_named(slide, f"{stem}-r1c{col}", headers[col - 1] if col - 1 < len(headers) else "", max_size=11)
    for row_index in range(2, max_rows + 2):
        row = rows[row_index - 2] if row_index - 2 < len(rows) else []
        for col in range(1, columns + 1):
            name = f"{stem}-r{row_index}c{col}"
            if row:
                value = row[col - 1] if col - 1 < len(row) else ""
                _set_named(slide, name, value, max_size=11)
            else:
                _remove_named(slide, name)


def _style_minimal_grid(
    slide: Any,
    stem: str,
    *,
    columns: int,
    data_rows: int,
    highlight_row: int | None = None,
) -> None:
    """Turn card-like cell boxes into a quiet editable matrix."""
    for row_index in range(1, data_rows + 2):
        if row_index == 1:
            fill = RGBColor(18, 61, 82)
        elif highlight_row is not None and row_index == highlight_row:
            fill = RGBColor(255, 244, 194)
        elif row_index % 2 == 0:
            fill = RGBColor(247, 250, 252)
        else:
            fill = RGBColor(255, 255, 255)
        for col in range(1, columns + 1):
            shape = _shape_named(slide, f"{stem}-r{row_index}c{col}")
            if shape is None:
                continue
            try:
                shape.fill.solid()
                shape.fill.fore_color.rgb = fill
                shape.line.fill.background()
            except Exception:
                pass


def _highlight_row_for_values(rows: list[list[str]], keywords: tuple[str, ...], *, default: int | None = 2) -> int | None:
    for index, row in enumerate(rows, start=2):
        text = " ".join(_display(value) for value in row).lower()
        if any(keyword.lower() in text for keyword in keywords):
            return index
    return default if rows else None


def _evidence_positions(count: int) -> list[tuple[float, float, float, float, float]]:
    if count <= 0:
        return []
    if count == 1:
        return [(2.0, 1.58, 9.3, 3.62, 1.12)]
    if count == 2:
        return [(0.95, 1.62, 5.55, 3.35, 1.16), (6.83, 1.62, 5.55, 3.35, 1.16)]
    if count == 3:
        return [(0.81 + index * 3.81, 1.62, 3.53, 3.35, 1.16) for index in range(3)]
    if count == 4:
        return [
            (1.05, 1.48, 5.35, 1.65, 0.78),
            (6.93, 1.48, 5.35, 1.65, 0.78),
            (1.05, 4.08, 5.35, 1.65, 0.78),
            (6.93, 4.08, 5.35, 1.65, 0.78),
        ]
    positions: list[tuple[float, float, float, float, float]] = []
    for index in range(count):
        row, col = divmod(index, 3)
        positions.append((0.81 + col * 3.81, 1.46 + row * 2.58, 3.53, 1.52, 0.72))
    return positions


def _evidence_text_card_positions(count: int) -> list[tuple[float, float, float, float]]:
    if count == 1:
        return [(1.5, 1.48, 10.33, 5.02)]
    if count == 2:
        return [(0.75, 1.48, 5.72, 5.02), (6.86, 1.48, 5.72, 5.02)]
    if count == 3:
        return [(0.75 + index * 4.03, 1.48, 3.72, 5.02) for index in range(3)]
    positions: list[tuple[float, float, float, float]] = []
    columns = 2 if count == 4 else 3
    width = 5.72 if columns == 2 else 3.72
    x_gap = 6.11 if columns == 2 else 4.03
    for index in range(count):
        row, col = divmod(index, columns)
        positions.append((0.75 + col * x_gap, 1.48 + row * 2.58, width, 2.28))
    return positions


def _reflow_case_thumbnails(slide, count: int, caption_flags: list[bool] | None = None) -> None:
    if count <= 0:
        return
    if count == 1:
        positions = [(8.23, 1.43, 4.35, 4.15, 0.86)]
    elif count == 2:
        positions = [(8.23, 1.43, 4.35, 1.78, 0.72), (8.23, 4.08, 4.35, 1.72, 0.68)]
    else:
        positions = [
            (8.23, 1.43, 1.90, 1.55, 0.72),
            (10.40, 1.43, 1.90, 1.55, 0.72),
            (8.23, 4.08, 1.90, 1.55, 0.72),
            (10.40, 4.08, 1.90, 1.55, 0.72),
        ]
    for index, (x, y, width, image_h, caption_h) in enumerate(positions[:count]):
        slot = _shape_named(slide, f"case-thumb-{index}")
        caption = _shape_named(slide, f"case-thumb-caption-{index}")
        has_caption = caption_flags[index] if caption_flags and index < len(caption_flags) else caption is not None
        if not has_caption:
            image_h += caption_h + 0.11
        _set_geometry(slot, x, y, width, image_h)
        if has_caption:
            _set_geometry(caption, x, y + image_h + 0.11, width, caption_h)


def _bind_benchmark(slide, payload: dict[str, Any]) -> None:
    _fill_kpis(slide, "bench", payload)
    points = _payload_points(payload, 8)
    _set_named(slide, "bench-protocol", _display(payload.get("protocol")) or "\n".join(points[:4]), max_size=12)
    if payload.get("_benchmark_rows_derived"):
        _remove_prefix(slide, "bench-table-")
        count = max(1, min(3, len(_structured_list(payload, "kpis")) or 3))
        gap = 0.28
        available_left, available_right = 0.75, 12.58
        width = min(5.6, (available_right - available_left - gap * (count - 1)) / count)
        total = width * count + gap * (count - 1)
        start = available_left + (available_right - available_left - total) / 2
        for index in range(1, count + 1):
            x = start + (index - 1) * (width + gap)
            _set_geometry(_shape_named(slide, f"bench-kpi{index}"), x, 1.45, width, 3.72)
            _set_geometry(_shape_named(slide, f"bench-kpi{index}-accent"), x, 1.45, 0.07, 3.72)
            _set_geometry(_shape_named(slide, f"bench-kpi{index}-value"), x + 0.22, 1.64, width - 0.42, 1.08)
            _set_geometry(_shape_named(slide, f"bench-kpi{index}-label"), x + 0.22, 2.82, width - 0.42, 0.56)
            _set_geometry(_shape_named(slide, f"bench-kpi{index}-note"), x + 0.22, 3.50, width - 0.42, 1.10)
            _refit_shape_text(_shape_named(slide, f"bench-kpi{index}-value"), max_size=28, min_size=16)
            _refit_shape_text(_shape_named(slide, f"bench-kpi{index}-label"), max_size=15, min_size=12)
            _refit_shape_text(_shape_named(slide, f"bench-kpi{index}-note"), max_size=12, min_size=10)
            _set_text_spacing(_shape_named(slide, f"bench-kpi{index}-note"), after=8, line_spacing=1.1)
        _set_geometry(_shape_named(slide, "bench-protocol"), 0.75, 5.38, 11.83, 1.17)
        _set_text_spacing(_shape_named(slide, "bench-protocol"), after=8, line_spacing=1.08, middle=True)
        return
    headers = (
        ["指标", "定义", "方向", "数据集", "说明"]
        if _is_cjk_payload(payload)
        else ["Metric group", "Metrics", "Direction", "Dataset", "Notes"]
    )
    _fill_grid(slide, "bench-table", payload, 5, 4, headers)


def _bind_big_numbers(slide, payload: dict[str, Any]) -> None:
    _fill_kpis(slide, "big", payload)
    items = _kpis(payload)
    count = max(1, min(3, len(_structured_list(payload, "kpis")) or len(items)))
    primary = _semantic_primary_index(items[:count], payload)
    secondary = [index for index in range(count) if index != primary]
    positions: dict[int, tuple[float, float, float, float, bool]] = {}
    if count == 1:
        positions[primary] = (1.55, 1.48, 10.23, 3.92, False)
    elif count == 2:
        positions[primary] = (0.85, 1.48, 7.10, 3.92, False)
        positions[secondary[0]] = (8.28, 1.48, 4.20, 3.92, False)
    else:
        positions[primary] = (0.85, 1.48, 5.28, 3.92, False)
        positions[secondary[0]] = (6.45, 1.48, 6.03, 1.80, True)
        positions[secondary[1]] = (6.45, 3.60, 6.03, 1.80, True)
    for item_index in range(count):
        x, y, width, height, compact = positions[item_index]
        shape_index = item_index + 1
        base = _shape_named(slide, f"big-kpi{shape_index}")
        accent = _shape_named(slide, f"big-kpi{shape_index}-accent")
        value = _shape_named(slide, f"big-kpi{shape_index}-value")
        label = _shape_named(slide, f"big-kpi{shape_index}-label")
        note = _shape_named(slide, f"big-kpi{shape_index}-note")
        _set_geometry(base, x, y, width, height)
        _set_geometry(accent, x, y, 0.08, height)
        if compact:
            _set_geometry(value, x + 0.25, y + 0.22, 1.72, 0.60)
            _set_geometry(label, x + 2.08, y + 0.23, width - 2.34, 0.58)
            _set_geometry(note, x + 0.25, y + 1.02, width - 0.52, 0.52)
            _refit_shape_text(value, max_size=24, min_size=16)
        else:
            _set_geometry(value, x + 0.28, y + 0.34, width - 0.56, 0.94)
            _set_geometry(label, x + 0.28, y + 1.50, width - 0.56, 0.55)
            _set_geometry(note, x + 0.28, y + 2.25, width - 0.56, 1.08)
            _refit_shape_text(value, max_size=36 if item_index == primary else 30, min_size=18)
        _refit_shape_text(label, max_size=15, min_size=12)
        _refit_shape_text(note, max_size=12, min_size=10)
        for inner_shape in (value, label, note):
            _set_shape_colors(inner_shape, fill=None, line=None)
        is_primary = item_index == primary
        _set_shape_colors(base, fill=SEMANTIC_PRIMARY_LIGHT if is_primary else RGBColor(255, 255, 255), line=RGBColor(207, 221, 228))
        _set_shape_colors(accent, fill=SEMANTIC_PRIMARY if is_primary else SEMANTIC_MUTED, line=None)
        _set_shape_text_color(value, SEMANTIC_PRIMARY if is_primary else SEMANTIC_DARK)
    points = _payload_points(payload, 6)
    primary_item = items[primary] if primary < len(items) else {"value": "", "label": "", "note": ""}
    interpretation = _display(payload.get("interpretation") or payload.get("conclusion")) or (points[-1] if points else "")
    if not interpretation:
        interpretation = "：".join(part for part in [primary_item.get("label", ""), primary_item.get("value", "")] if part)
        interpretation = "；".join(part for part in [interpretation, primary_item.get("note", "")] if part)
    interpretation_shape = _shape_named(slide, "big-interpretation")
    _set_geometry(interpretation_shape, 0.85, 5.70, 11.63, 0.88)
    _set_shape_text(
        interpretation_shape,
        f"{_localized(payload, '关键结论', 'KEY FINDING')}  {interpretation}",
        max_size=15,
        bold=True,
    )
    _set_shape_colors(interpretation_shape, fill=SEMANTIC_WARM_LIGHT, line=None)
    _set_text_spacing(interpretation_shape, after=2, line_spacing=1.0, middle=True)


def _bars(payload: dict[str, Any], count: int) -> list[dict[str, Any]]:
    raw = _structured_list(payload, "bars")
    result: list[dict[str, Any]] = []
    for item in raw:
        if isinstance(item, dict):
            result.append({
                "label": _display(item.get("label")),
                "value": _display(item.get("value")),
                "primary": item.get("primary") or item.get("highlight") or item.get("best") or item.get("is_primary"),
                "status": _display(item.get("status") or item.get("state")),
            })
    points = _payload_points(payload, count + 3)
    numbers = _payload_numbers(payload, count)
    while len(result) < count:
        index = len(result)
        result.append({
            "label": points[index] if index < len(points) else f"指标 {index + 1}",
            "value": numbers[index] if index < len(numbers) else str(35 + index * 15),
            "primary": False,
        })
    return result[:count]


def _numeric_ratio(value: str) -> float:
    match = re.search(r"[+\-]?([0-9]+(?:\.[0-9]+)?)", str(value or ""))
    if not match:
        return 0.55
    number = float(match.group(1))
    return max(0.08, min(1.0, number / 100 if number <= 100 else number / max(number, 100)))


def _resize_fill(slide, fill_name: str, track_name: str, value: str) -> None:
    fill = _shape_named(slide, fill_name)
    track = _shape_named(slide, track_name)
    if fill is not None and track is not None:
        fill.left = track.left
        fill.width = max(Inches(0.08), int(track.width * _numeric_ratio(value)))


def _bind_results_bars(slide, payload: dict[str, Any]) -> None:
    bars = _bars(payload, 4)
    explicit_count = min(4, len(_structured_list(payload, "bars")))
    count = explicit_count or len(bars)
    active_bars = bars[:count]
    primary = _semantic_primary_index(active_bars, payload)
    _set_geometry(_shape_named(slide, "bars-chart-panel"), 0.75, 1.48, 8.62, 5.10)
    _set_shape_colors(_shape_named(slide, "bars-chart-panel"), fill=RGBColor(255, 255, 255), line=RGBColor(218, 227, 233))
    _set_geometry(_shape_named(slide, "bars-title"), 1.12, 1.80, 4.20, 0.42)
    _set_named(slide, "bars-title", _display(payload.get("chart_title")) or _localized(payload, "指标对比", "Metric comparison"), max_size=17, bold=True)
    row_gap = 0.92 if count >= 4 else 1.02
    start_y = 2.48 + max(0.0, (4 - count) * 0.18)
    for index, item in enumerate(active_bars, 1):
        item_index = index - 1
        is_primary = item_index == primary
        y = start_y + item_index * row_gap
        bar_height = 0.46 if is_primary else 0.30
        _set_geometry(_shape_named(slide, f"bar{index}-label"), 1.12, y, 1.45, 0.36)
        _set_geometry(_shape_named(slide, f"bar{index}-track"), 2.82, y + (0.36 - bar_height) / 2, 4.75, bar_height)
        _set_geometry(_shape_named(slide, f"bar{index}-fill"), 2.82, y + (0.36 - bar_height) / 2, 4.75, bar_height)
        _set_geometry(_shape_named(slide, f"bar{index}-value"), 7.76, y, 0.76, 0.36)
        _set_named(slide, f"bar{index}-label", item["label"], max_size=13)
        _set_named(slide, f"bar{index}-value", item["value"], max_size=15, bold=True)
        _resize_fill(slide, f"bar{index}-fill", f"bar{index}-track", item["value"])
        track = _shape_named(slide, f"bar{index}-track")
        fill = _shape_named(slide, f"bar{index}-fill")
        _set_shape_colors(track, fill=RGBColor(233, 241, 245), line=None)
        _set_shape_colors(fill, fill=SEMANTIC_PRIMARY if is_primary else SEMANTIC_MUTED, line=None)
        if is_primary:
            _set_shape_text_color(_shape_named(slide, f"bar{index}-label"), SEMANTIC_DARK)
            _set_shape_text_color(_shape_named(slide, f"bar{index}-value"), SEMANTIC_PRIMARY)
            _style_first_paragraph(_shape_named(slide, f"bar{index}-label"), bold=True)
    for index in range(count + 1, 5):
        _remove_prefix(slide, f"bar{index}-")
    points = _payload_points(payload, 8)
    conclusion = _display(payload.get("conclusion") or payload.get("interpretation"))
    if not conclusion and active_bars:
        baseline_index = next((index for index, item in enumerate(active_bars) if re.search(r"baseline|基线|对照", item["label"], re.IGNORECASE)), None)
        if baseline_index is not None and baseline_index != primary:
            primary_value = _numeric_value(active_bars[primary]["value"])
            baseline_value = _numeric_value(active_bars[baseline_index]["value"])
            if primary_value is not None and baseline_value is not None:
                delta = primary_value - baseline_value
                unit = _localized(payload, " 个百分点", " pts") if "%" in active_bars[primary]["value"] else ""
                conclusion = _localized(
                    payload,
                    f"{active_bars[primary]['label']} 相比 {active_bars[baseline_index]['label']} 提升 {delta:+g}{unit}",
                    f"{active_bars[primary]['label']} improves on {active_bars[baseline_index]['label']} by {delta:+g}{unit}",
                )
        if not conclusion:
            conclusion = f"{active_bars[primary]['label']}  {active_bars[primary]['value']}"
    note_lines = [_localized(payload, "关键发现", "KEY FINDING"), conclusion] if conclusion else []
    note_lines.extend(f"• {point}" for point in points[:3] if point and point != conclusion)
    note = "\n".join(note_lines[:4]) or _localized(payload, "突出最重要的比较结论", "State the most important comparison")
    _set_geometry(_shape_named(slide, "bars-note1"), 9.72, 1.48, 2.86, 5.10)
    _set_named(slide, "bars-note1", note, max_size=14)
    note_shape = _shape_named(slide, "bars-note1")
    _set_shape_colors(note_shape, fill=SEMANTIC_PRIMARY_LIGHT, line=None)
    _set_text_spacing(note_shape, after=10, line_spacing=1.10, middle=True)
    _style_first_paragraph(note_shape, color=SEMANTIC_PRIMARY)
    if note_shape is not None and len(note_shape.text_frame.paragraphs) > 1:
        for run in note_shape.text_frame.paragraphs[1].runs:
            run.font.bold = True
    _remove_named(slide, "bars-note2")
    _remove_named(slide, "bars-note3")


def _bind_leaderboard(slide, payload: dict[str, Any]) -> None:
    _fill_grid(slide, "leaderboard-grid", payload, 5, 6, ["排名", "方法", "得分", "成本", "说明"])
    points = _payload_points(payload, 8)
    rows = _rows(payload)
    active_rows = min(6, len(rows) or len(points) or 1)
    has_notes = bool(points)
    left = 0.75
    widths = (0.95, 2.25, 1.45, 1.35, 2.55) if has_notes else (1.05, 3.10, 1.80, 1.60, 4.28)
    row_height = min(0.82, 5.02 / (active_rows + 1))
    for row_index in range(1, active_rows + 2):
        x = left
        for col, width in enumerate(widths, start=1):
            shape = _shape_named(slide, f"leaderboard-grid-r{row_index}c{col}")
            if shape is not None:
                _set_geometry(shape, x, 1.52 + (row_index - 1) * row_height, width, row_height)
                _refit_shape_text(shape, max_size=13 if row_index == 1 else 12, min_size=11)
            x += width
    if has_notes:
        note = "\n".join(f"• {point}" for point in points[:4])
        _set_geometry(_shape_named(slide, "leader-note1"), 9.85, 1.52, 2.73, 5.02)
        _set_named(slide, "leader-note1", note, max_size=14)
        _set_text_spacing(_shape_named(slide, "leader-note1"), after=12, line_spacing=1.12, middle=True)
    else:
        _remove_named(slide, "leader-note1")
    _remove_named(slide, "leader-note2")
    _remove_named(slide, "leader-note3")


def _bind_ablation(slide, payload: dict[str, Any]) -> None:
    rows = _rows(payload)
    headers = _headers(payload)
    explicit_columns = max([len(headers), *(len(row) for row in rows)], default=0)
    column_count = max(3, min(6, explicit_columns or 6))
    _fill_grid(
        slide,
        "ablation-grid",
        payload,
        column_count,
        5,
        ["方案", "模块A", "模块B", "模块C", "得分", "变化"][:column_count],
    )
    for row_index in range(1, 7):
        for col in range(column_count + 1, 7):
            _remove_named(slide, f"ablation-grid-r{row_index}c{col}")
    active_rows = min(5, len(rows) or len(_payload_points(payload, 5)) or 1)
    table_width = 8.25
    cell_width = table_width / column_count
    row_height = min(0.90, 4.98 / (active_rows + 1))
    for row_index in range(1, active_rows + 2):
        for col in range(1, column_count + 1):
            shape = _shape_named(slide, f"ablation-grid-r{row_index}c{col}")
            if shape is not None:
                _set_geometry(shape, 0.75 + (col - 1) * cell_width, 1.48 + (row_index - 1) * row_height, cell_width, row_height)
                _refit_shape_text(shape, max_size=13 if row_index == 1 else 12, min_size=11)
    highlight_row = _highlight_row_for_values(rows, ("full", "ours", "proposed", "完整", "当前方案", "全部"))
    _style_minimal_grid(slide, "ablation-grid", columns=column_count, data_rows=active_rows, highlight_row=highlight_row)
    points = _payload_points(payload, 8)
    _set_named(slide, "ablation-callout", _display(payload.get("conclusion")) or (points[0] if points else ""), max_size=15)
    guide = _display(payload.get("guide"))
    if guide:
        _set_geometry(_shape_named(slide, "ablation-callout"), 9.35, 1.48, 3.23, 2.32)
        _set_geometry(_shape_named(slide, "ablation-guide"), 9.35, 4.10, 3.23, 2.36)
        _set_named(slide, "ablation-guide", guide, max_size=12)
    else:
        _remove_named(slide, "ablation-guide")
        _set_geometry(_shape_named(slide, "ablation-callout"), 9.35, 1.48, 3.23, 4.98)
        _set_text_spacing(_shape_named(slide, "ablation-callout"), after=12, line_spacing=1.12, middle=True)


def _bind_evidence(slide, payload: dict[str, Any]) -> None:
    cards = _structured_list(payload, "cards")
    caption_items = [item for item in _structured_list(payload, "captions") if _display(item)] or _payload_points(payload, 6)
    image_count = max(0, min(6, int(payload.get("_image_count") or 0)))
    slots = [(row, col) for row in range(2) for col in range(3)]
    if image_count:
        positions = _evidence_positions(image_count)
        for index, (row, col) in enumerate(slots):
            slot_name = f"evidence-img-{row}-{col}"
            caption_name = f"evidence-caption-{row}-{col}"
            if index >= image_count:
                _remove_slot(slide, slot_name)
                _remove_named(slide, caption_name)
                continue
            _set_named(slide, slot_name, "", max_size=14)
            _remove_crosses(slide, slot_name)
            x, y, width, image_h, caption_h = positions[index]
            caption_item = caption_items[index] if index < len(caption_items) else ""
            caption = _figure_caption(payload, caption_item, index) if caption_item or _source_for_index(payload, index) else ""
            if caption:
                _set_named(slide, caption_name, caption, max_size=11)
                caption_shape = _shape_named(slide, caption_name)
                _set_geometry(caption_shape, x, y + image_h + 0.1, width, caption_h)
                _style_figure_caption(caption_shape)
            else:
                image_h += caption_h + 0.1
                _remove_named(slide, caption_name)
            _set_geometry(_shape_named(slide, slot_name), x, y, width, image_h)
        return
    items = cards or caption_items
    items = items[:6]
    positions = _evidence_positions(len(items))
    text_card_positions = _evidence_text_card_positions(len(items)) if cards else []
    for index, (row, col) in enumerate(slots):
        slot_name = f"evidence-img-{row}-{col}"
        caption_name = f"evidence-caption-{row}-{col}"
        if index >= len(items):
            _remove_slot(slide, slot_name)
            _remove_named(slide, caption_name)
            continue
        title, body = _item_parts(items[index])
        if cards:
            _set_named(slide, slot_name, _card_text(items[index]), max_size=15)
            _remove_crosses(slide, slot_name)
            _style_text_card(_shape_named(slide, slot_name))
            _set_text_spacing(_shape_named(slide, slot_name), after=10, line_spacing=1.1, middle=True)
            _remove_named(slide, caption_name)
            if index < len(text_card_positions):
                _set_geometry(_shape_named(slide, slot_name), *text_card_positions[index])
            continue
        if not cards:
            title, body = title or _localized(payload, "证据", "Evidence"), body
        _set_named(slide, slot_name, title, max_size=14)
        _remove_crosses(slide, slot_name)
        _set_named(slide, caption_name, body, max_size=11)
        if index < len(positions):
            x, y, width, image_h, caption_h = positions[index]
            _set_geometry(_shape_named(slide, slot_name), x, y, width, image_h)
            _set_geometry(_shape_named(slide, caption_name), x, y + image_h + 0.1, width, caption_h)


def _bind_case(slide, payload: dict[str, Any]) -> None:
    cards = _structured_list(payload, "cards")
    caption_items = [item for item in _structured_list(payload, "captions") if _display(item)] or _payload_points(payload, 5)
    image_count = max(0, min(5, int(payload.get("_image_count") or 0)))
    if image_count:
        _set_named(slide, "case-main", "", max_size=15)
        _remove_crosses(slide, "case-main")
        main_item = caption_items[0] if caption_items else ""
        main_caption = _figure_caption(payload, main_item, 0) if main_item or _source_for_index(payload, 0) else ""
        if main_caption:
            _set_geometry(_shape_named(slide, "case-main"), 0.75, 1.43, 7.18, 4.25)
            _set_named(slide, "case-caption-main", main_caption, max_size=11)
            _set_geometry(_shape_named(slide, "case-caption-main"), 0.75, 5.79, 7.18, 0.72)
            _style_figure_caption(_shape_named(slide, "case-caption-main"))
        else:
            _remove_named(slide, "case-caption-main")
            _set_geometry(_shape_named(slide, "case-main"), 0.75, 1.43, 7.18, 5.12)
        caption_flags: list[bool] = []
        for index in range(4):
            slot_name = f"case-thumb-{index}"
            caption_name = f"case-thumb-caption-{index}"
            if index + 1 >= image_count:
                _remove_slot(slide, slot_name)
                _remove_named(slide, caption_name)
                continue
            _set_named(slide, slot_name, "", max_size=13)
            _remove_crosses(slide, slot_name)
            item_index = index + 1
            caption_item = caption_items[item_index] if item_index < len(caption_items) else ""
            caption = _figure_caption(payload, caption_item, item_index) if caption_item or _source_for_index(payload, item_index) else ""
            if caption:
                _set_named(slide, caption_name, caption, max_size=11)
                _style_figure_caption(_shape_named(slide, caption_name))
                caption_flags.append(True)
            else:
                _remove_named(slide, caption_name)
                caption_flags.append(False)
        _reflow_case_thumbnails(slide, image_count - 1, caption_flags)
        return
    items = (cards or caption_items)[:5]
    main_item = items[0] if items else ""
    main_title, main_body = _item_parts(main_item)
    _set_named(slide, "case-caption-main", main_body, max_size=11)
    _set_named(slide, "case-main", _display(payload.get("main_label")) or main_title, max_size=15)
    _remove_crosses(slide, "case-main")
    for index in range(4):
        slot_name = f"case-thumb-{index}"
        caption_name = f"case-thumb-caption-{index}"
        item_index = index + 1
        if item_index >= len(items):
            _remove_slot(slide, slot_name)
            _remove_named(slide, caption_name)
            continue
        title, body = _item_parts(items[item_index])
        _set_named(slide, slot_name, title, max_size=13)
        _remove_crosses(slide, slot_name)
        _set_named(slide, caption_name, body, max_size=11)
    _reflow_case_thumbnails(slide, max(0, len(items) - 1))


def _bind_summary(slide, payload: dict[str, Any]) -> None:
    findings_raw = _structured_list(payload, "findings") or _structured_list(payload, "cards")
    findings: list[str] = []
    for item in findings_raw:
        value = _card_text(item) if isinstance(item, dict) else _display(item)
        if value:
            findings.append(value)
    findings = findings or _payload_points(payload, 4)
    conclusion = _display(payload.get("conclusion"))
    if conclusion and not findings:
        findings.insert(0, conclusion)
    limitation_value = payload.get("limitations") or payload.get("limitation") or payload.get("constraints")
    limitation = _display(limitation_value)
    entries: list[tuple[str, str]] = []
    finding_limit = 2 if limitation else 3
    for index, finding in enumerate(findings[:finding_limit]):
        entries.append(("primary" if index == 0 else "support", finding))
    if limitation and len(entries) < 3:
        entries.append(("limitation", limitation))
    entries = entries[:3]
    role_labels = {
        "primary": _localized(payload, "主要发现", "PRIMARY FINDING"),
        "support": _localized(payload, "支撑发现", "SUPPORTING FINDING"),
        "limitation": _localized(payload, "限制条件", "LIMITATION"),
    }
    for index in range(1, 4):
        role, value = entries[index - 1] if index - 1 < len(entries) else ("", "")
        if value:
            shape = _shape_named(slide, f"summary-card{index}")
            _set_shape_text(shape, f"{role_labels[role]}\n{value}", max_size=18 if role == "primary" else 15)
            if shape is not None:
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                shape.text_frame.margin_left = Inches(0.32)
                shape.text_frame.margin_right = Inches(0.28)
            _set_text_spacing(shape, after=10, line_spacing=1.1, middle=True)
            if role == "primary":
                _set_shape_colors(shape, fill=SEMANTIC_PRIMARY_LIGHT, line=None)
                _style_first_paragraph(shape, color=SEMANTIC_PRIMARY)
            elif role == "limitation":
                _set_shape_colors(shape, fill=SEMANTIC_WARM_LIGHT, line=None)
                _style_first_paragraph(shape, color=SEMANTIC_WARM)
            else:
                _set_shape_colors(shape, fill=RGBColor(255, 255, 255), line=RGBColor(218, 227, 233))
                _style_first_paragraph(shape, color=SEMANTIC_DARK)
        else:
            _remove_named(slide, f"summary-card{index}")
    active_shapes = [_shape_named(slide, f"summary-card{index}") for index in range(1, len(entries) + 1)]
    active_shapes = [shape for shape in active_shapes if shape is not None]
    if len(active_shapes) == 1:
        _set_geometry(active_shapes[0], 1.72, 1.58, 9.89, 3.85)
    elif len(active_shapes) == 2:
        _set_geometry(active_shapes[0], 0.85, 1.58, 5.55, 3.85)
        _set_geometry(active_shapes[1], 6.93, 1.58, 5.55, 3.85)
    elif len(active_shapes) >= 3:
        _set_geometry(active_shapes[0], 0.85, 1.58, 5.42, 3.85)
        _set_geometry(active_shapes[1], 6.67, 1.58, 5.81, 1.67)
        _set_geometry(active_shapes[2], 6.67, 3.76, 5.81, 1.67)
    next_step = _display(payload.get("next_step") or payload.get("recommendation") or payload.get("decision_request"))
    summary_next = _shape_named(slide, "summary-next")
    if summary_next is not None:
        summary_next.text_frame.margin_left = Inches(0.03)
        summary_next.text_frame.margin_right = Inches(0.03)
        summary_next.text_frame.margin_top = Inches(0.02)
        summary_next.text_frame.margin_bottom = Inches(0.02)
        _set_geometry(summary_next, 0.85, 5.73, 11.63, 0.85)
    closing = next_step or conclusion or _localized(payload, "明确下一步行动与责任人", "Define the next action and owner")
    _set_shape_text(
        summary_next,
        f"{_localized(payload, '下一步', 'NEXT STEP')}  {closing}",
        max_size=15,
        bold=True,
    )
    _set_shape_colors(summary_next, fill=SEMANTIC_DARK, line=None)
    _set_shape_text_color(summary_next, RGBColor(255, 255, 255))
    _set_text_spacing(summary_next, after=2, line_spacing=1.0, middle=True)



__all__ = [
    "_bind_basic_content",
    "_configure_basic_content_layout",
    "_motivation_content",
    "_bind_motivation",
    "_bind_challenge",
    "_bind_pipeline",
    "_bind_loop",
    "_kpis",
    "_fill_kpis",
    "_rows",
    "_target_row",
    "_evaluation_status_table",
    "_headers",
    "_fill_grid",
    "_style_minimal_grid",
    "_highlight_row_for_values",
    "_evidence_positions",
    "_evidence_text_card_positions",
    "_reflow_case_thumbnails",
    "_bind_benchmark",
    "_bind_big_numbers",
    "_bars",
    "_numeric_ratio",
    "_resize_fill",
    "_bind_results_bars",
    "_bind_leaderboard",
    "_bind_ablation",
    "_bind_evidence",
    "_bind_case",
    "_bind_summary",
]
